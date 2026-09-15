import openpyxl
from backend.utils.template_fill_binding import bind_template_fill, remap_local_formula


def test_bind_actual_headers_without_touching_merged_heading():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells('AK2:AS2')
    ws['AK2'] = 'Employer'
    ws['A3'], ws['N3'], ws['AL3'], ws['BK3'] = 'StaffNumber', '天数', '医疗保险', 'One-childfee'
    ws['A4'], ws['A5'], ws['A6'] = '001', '002', 'Total'
    ws['AL4'] = 100
    expected = [{'letter': 'A', 'name': 'StaffNumber'}, {'letter': 'N', 'name': '天数'},
                {'letter': 'AF', 'name': 'One-childfee'}]
    rows, mapping = bind_template_fill(ws, expected, {'AL': '智算_实际在职工作日'})
    assert rows == [4, 5]
    assert mapping['AF'] == 63
    assert mapping['AL'] > 63
    assert ws['AL4'].value == 100
    assert str(next(iter(ws.merged_cells.ranges))) == 'AK2:AS2'
    formula = '=IF(AL4>0,AF4,0)+SUM(\'源_名单\'!AL4:AL5)+IF(A4="AL4",1,0)'
    converted = remap_local_formula(formula, mapping)
    assert 'BL4>0,BK4' in converted
    assert "'源_名单'!AL4:AL5" in converted
    assert '"AL4"' in converted
