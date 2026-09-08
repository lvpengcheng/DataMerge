"""Regressions for train/compute parity and false-positive training validation."""
import ast
import logging
import os
import shutil
import tempfile
import time
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock
from typing import Dict, Any

import openpyxl
import pandas as pd
import pytest

from backend.utils.excel_comparator import _compare_dataframes_core, require_complete_comparison
from backend.utils.source_selector import find_source_sheet
from backend.utils.script_entry import invoke_script_main


def test_output_selection_never_uses_mtime_to_guess_business_result(tmp_path):
    from backend.utils.result_selection import pick_result_output
    template = tmp_path / 'template.xlsx'
    template.write_bytes(b'template')
    copy = tmp_path / 'copy.xlsx'
    copy.write_bytes(template.read_bytes())
    result = tmp_path / 'result.xlsx'
    result.write_bytes(b'filled workbook')
    diff = tmp_path / '_diff.xlsx'
    diff.write_bytes(b'diff')
    assert pick_result_output([copy, diff, result], template) == result
    second = tmp_path / 'newer.xlsx'
    second.write_bytes(b'another real result')
    with pytest.raises(ValueError, match='多个结果工作簿'):
        pick_result_output([copy, result, second], template)
    assert pick_result_output([diff]) is None


@pytest.mark.parametrize('declaration', ['工号，月份', '["工号", "月份"]', '`工号` + `月份`（联合唯一）'])
def test_composite_primary_key_declaration_preserves_all_columns(declaration):
    from backend.utils.excel_comparator import extract_primary_keys_from_rules, _resolve_primary_keys
    keys = extract_primary_keys_from_rules('- 主键：' + declaration)
    assert keys == ['工号', '月份']
    expected = pd.DataFrame(columns=['工号', '月份', '金额'])
    result = pd.DataFrame(columns=['工号', '金额'])
    with pytest.raises(ValueError, match='不能降级'):
        _resolve_primary_keys(expected, result, keys)
    with pytest.raises(ValueError, match='缺失或不唯一'):
        extract_primary_keys_from_rules('- 主键：工号', ['上月工号', '本月工号'])


def test_manual_selection_preserves_measured_accuracy(tmp_path, monkeypatch):
    import sys
    path = Path(__file__).parent / 'api' / 'training_chat.py'
    node = next(n for n in ast.parse(path.read_text(encoding='utf-8')).body
                if isinstance(n, ast.FunctionDef) and n.name == 'set_as_best')
    node.decorator_list = []
    for argument in node.args.args:
        argument.annotation = None
    node.args.defaults = [ast.Constant(None) for _ in node.args.defaults]
    session_model, iteration_model = object(), object()
    session = SimpleNamespace(config={}, tenant_id='tenant', mode='formula', source_structure={})
    iteration = SimpleNamespace(accuracy=0.75, generated_code='print(1)', iteration_num=3)
    queries = {session_model: MagicMock(), iteration_model: MagicMock()}
    queries[session_model].filter_by.return_value.first.return_value = session
    queries[iteration_model].filter_by.return_value.first.return_value = iteration
    db = MagicMock()
    db.query.side_effect = lambda model: queries[model]
    storage = MagicMock()
    storage.save_script.return_value = {'script_id': 'script1'}
    persistence = MagicMock()
    persistence.save_script.return_value = SimpleNamespace(id=3, version=2)
    monkeypatch.setitem(sys.modules, 'backend.storage.storage_manager', SimpleNamespace(StorageManager=lambda: storage))
    monkeypatch.setitem(sys.modules, 'backend.api.training_persistence', SimpleNamespace(TrainingPersistence=lambda db: persistence))
    from datetime import datetime
    namespace = {'__package__': 'backend.api', 'TrainingSession': session_model,
                 'TrainingIteration': iteration_model, 'datetime': datetime,
                 'logger': logging.getLogger('chain-test'), '_add_message': MagicMock()}
    exec(compile(ast.fix_missing_locations(ast.Module(body=[node], type_ignores=[])), str(path), 'exec'), namespace)
    result = namespace[node.name](7, SimpleNamespace(iteration_id=10), db, SimpleNamespace(id=1))
    assert result['accuracy'] == 0.75 and result['manual_approved']
    assert iteration.accuracy == 0.75
    assert storage.save_script.call_args.args[2]['best_score'] == 0.75
    assert persistence.save_script.call_args.kwargs['accuracy'] == 0.75
    queries[iteration_model].filter_by.assert_called_once_with(id=10, session_id=7)


@pytest.mark.parametrize('kind', ['formula', 'template'])
def test_generated_fallback_loads_formula_values_and_rejects_partial_input(tmp_path, kind):
    from backend.ai_engine.formula_code_generator import FormulaCodeGenerator
    from backend.ai_engine.template_code_generator import TemplateCodeGenerator
    from backend.test_pipeline_accuracy import function_namespace
    from backend.utils.data_helpers import assign_sheet_keys, apply_dataframe_column_schemas, region_schemas_by_name
    from excel_parser import IntelligentExcelParser
    source = tmp_path / 'source'
    source.mkdir()
    workbook(source / 'sample.XLSX', {'data': [('001', '=2*5'), ('002', '=3*5')],
                                      '忽略页': [('003', 99)]})
    calls = []
    class Parser(IntelligentExcelParser):
        def parse_excel_file(self, path, **kwargs):
            calls.append(kwargs)
            return super().parse_excel_file(path, **kwargs)
    if kind == 'formula':
        code = FormulaCodeGenerator(ai_provider=object())._build_complete_code('def fill_result_sheets(*args):\n    pass')
    else:
        code = TemplateCodeGenerator(ai_provider=object())._build_complete_code(
            'def fill_template(*args):\n    pass', str(tmp_path / 'template.xlsx'), str(source),
            {'sheets': {}}, {}, False)
    namespace = function_namespace(code, {'os': os, 'Path': Path, 'input_folder': str(source),
        'IntelligentExcelParser': Parser, 'assign_sheet_keys': assign_sheet_keys,
        'apply_dataframe_column_schemas': apply_dataframe_column_schemas,
        'region_schemas_by_name': region_schemas_by_name, '_COL_MAP': {'DATA': {}},
        '_RESULT_SHEET_NAMES': ['DATA']})
    # Template imports the parser locally, so replace the same class at the module boundary.
    from unittest.mock import patch
    with patch('excel_parser.IntelligentExcelParser', Parser):
        run = lambda: namespace['load_source_data'](str(source), {}) if kind == 'formula' else namespace['load_source_data']()
        loaded = run()
        assert list(loaded) == ['sample_data']  # Case-insensitive reserved result name.
        assert loaded['sample_data']['df']['金额'].tolist() == [10, 15]
        assert loaded['sample_data']['df']['工号'].tolist() == ['001', '002']
        assert calls[-1]['active_sheet_only'] is True
        assert calls[-1]['calculate_formulas'] and calls[-1]['normalize_source'] and calls[-1]['raise_errors']
        (source / 'broken.xlsx').write_bytes(b'not an excel archive')
        with pytest.raises(ValueError, match='broken.xlsx'):
            run()


@pytest.mark.parametrize('failure', ['open'])
def test_comparison_releases_native_workbooks_on_sample_failure(monkeypatch, failure):
    import aspose_init
    aspose_init.ensure_license()
    import sys
    from backend.utils.excel_comparator import _open_comparison_workbooks
    allocated = []
    def constructor(path):
        if path == 'expected' and failure == 'open':
            raise OSError('bad sample')
        book = MagicMock()
        allocated.append(book)
        if path == 'expected':
            book.CalculateFormula.side_effect = OSError('bad sample')
        return book
    monkeypatch.setitem(sys.modules, 'Aspose.Cells', SimpleNamespace(Workbook=constructor))
    with pytest.raises(OSError, match='bad sample'):
        with _open_comparison_workbooks('result', 'expected'):
            pytest.fail('invalid sample must not be compared')
    assert allocated
    for book in allocated:
        book.Dispose.assert_called_once()


@pytest.mark.parametrize('killed', [False, True])
def test_failed_full_source_loading_never_falls_back_to_another_parse(monkeypatch, killed):
    import backend.utils.subprocess_runner as runner
    from backend.utils.training_validation import TrainingResourceFailure
    monkeypatch.setattr(runner, 'run_in_subprocess', lambda *a, **kw: SimpleNamespace(
        success=False, result=None, timed_out=killed, killed_by_memory=False,
        killed=killed, termination_failed=False, error='parse failed'))
    path = Path(__file__).parent / 'api' / 'training_chat.py'
    node = next(n for n in ast.parse(path.read_text(encoding='utf-8')).body
                if isinstance(n, ast.FunctionDef) and n.name == '_load_full_source_data_subproc')
    namespace = {'logger': logging.getLogger('chain-test')}
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(path), 'exec'), namespace)
    with pytest.raises(TrainingResourceFailure if killed else ValueError, match='加载失败'):
        namespace[node.name]('source')


@pytest.mark.parametrize('expected,result,keys', [
    ({'工号': ['001'], '金额': [10]}, {'工号': [1], '金额': [10]}, ['工号']),
    ({'工号': ['a_b'], '月份': ['c'], '金额': [10]},
     {'工号': ['a'], '月份': ['b_c'], '金额': [10]}, ['工号', '月份']),
    ({'工号': [None], '金额': [10]}, {'工号': [None], '金额': [10]}, ['工号']),
    ({'工号': ['1'], '金额': [10]}, {'工号': ['1'], '金额': [10.01]}, ['工号']),
    ({'工号': ['1'], '金额': ['123456789012345678']},
     {'工号': ['1'], '金额': ['123456789012345679']}, ['工号']),
    ({'工号': ['1'], '金额': [None]}, {'工号': ['1'], '金额': [0]}, ['工号']),
])
def test_comparison_never_hides_identity_or_value_difference(tmp_path, expected, result, keys):
    comparison = _compare_dataframes_core(pd.DataFrame(result), pd.DataFrame(expected), keys,
                                          str(tmp_path / 'diff.xlsx'))
    assert comparison['success'] is False
    assert comparison['total_differences'] > 0
    assert comparison['match_rate'] < 1


def test_extra_rows_reduce_accuracy(tmp_path):
    expected = pd.DataFrame({'工号': ['1'], '金额': [10]})
    result = pd.DataFrame({'工号': ['1', '2'], '金额': [10, 10]})
    comparison = _compare_dataframes_core(result, expected, ['工号'], str(tmp_path / 'diff.xlsx'))
    assert comparison['unmatched_result'] == 1
    assert comparison['match_rate'] == 0.5


def test_empty_comparison_cannot_pass(tmp_path):
    comparison = _compare_dataframes_core(pd.DataFrame(columns=['工号', '金额']),
        pd.DataFrame(columns=['工号', '金额']), ['工号'], str(tmp_path / 'diff.xlsx'))
    assert comparison['success'] is False
    with pytest.raises(ValueError, match='没有可校验'):
        require_complete_comparison(comparison)


def workbook(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(['工号', '金额'])
        for row in rows:
            sheet.append(row)
    book.save(path)


def test_multi_sheet_read_failure_invalidates_entire_score(tmp_path, monkeypatch):
    from backend.utils import excel_comparator as compare
    expected, result = tmp_path / 'expected.xlsx', tmp_path / 'result.xlsx'
    sheets = {'正常表': [('001', 10)], '坏表': [('001', 20)]}
    workbook(expected, sheets)
    workbook(result, sheets)
    real = compare._aspose_read_sheet_df
    def fail_bad_sheet(ws, *args, **kwargs):
        if ws.Name == '坏表':
            raise OSError('simulated read failure')
        return real(ws, *args, **kwargs)
    monkeypatch.setattr(compare, '_aspose_read_sheet_df', fail_bad_sheet)
    comparison = compare._compare_excel_files_multi_sheet_impl(
        str(result), str(expected), str(tmp_path / 'diff.xlsx'), ['工号'], result_calculated=True)
    assert comparison['success'] is False and comparison['comparison_complete'] is False
    with pytest.raises(ValueError, match='坏表'):
        require_complete_comparison(comparison)


def test_expected_formula_saved_values_are_preserved_and_headers_reuse_workbooks(tmp_path, monkeypatch):
    import excel_parser
    from backend.utils.excel_comparator import _compare_excel_files_multi_sheet_impl
    expected, result = tmp_path / 'expected.xlsx', tmp_path / 'result.xlsx'
    workbook(expected, {'01_订单': [('001', '=10+5')], '02_明细': [('001', 5)]})
    # 保存值特意与重算值不同，验证参考答案不会被对比过程改成 15。
    import zipfile
    with zipfile.ZipFile(expected) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(
        b'<f>10+5</f><v></v>', b'<f>10+5</f><v>12</v>')
    with zipfile.ZipFile(expected, 'w') as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    workbook(result, {'01_订单': [('001', 12)], '02_明细': [('001', 5)]})
    original = expected.read_bytes()
    monkeypatch.setattr(excel_parser, '_licensed_workbook', lambda *a, **k: pytest.fail('header detection reopened workbook'))
    comparison = _compare_excel_files_multi_sheet_impl(
        str(result), str(expected), str(tmp_path / 'diff.xlsx'), ['工号'], result_calculated=True)
    assert comparison['success'], comparison
    assert len(comparison['per_sheet']) == 2
    assert expected.read_bytes() == original


@pytest.mark.parametrize('result_sheet', ['考勤', '源_工资', '工资备份'])
def test_wrong_sheet_name_is_not_matched_by_index(tmp_path, result_sheet):
    from backend.utils.excel_comparator import _compare_excel_files_multi_sheet_impl
    expected, result = tmp_path / 'expected.xlsx', tmp_path / 'result.xlsx'
    workbook(expected, {'工资': [('001', 15)]})
    workbook(result, {result_sheet: [('001', 15)]})
    comparison = _compare_excel_files_multi_sheet_impl(
        str(result), str(expected), str(tmp_path / 'diff.xlsx'), ['工号'], result_calculated=True)
    assert comparison['success'] is False and comparison['missing_sheets'] == ['工资']


def test_source_selection_rejects_ambiguity_and_respects_month_boundaries():
    sources = {name: {'df': pd.DataFrame({'工号': ['1']})} for name in ['工资_1月', '工资_11月']}
    assert find_source_sheet(sources, sheet_name_hint='工资', salary_year=2026, salary_month=1) == '工资_1月'
    with pytest.raises(ValueError, match='不唯一'):
        find_source_sheet(sources, ['工号'])
    with pytest.raises(KeyError, match='无匹配'):
        find_source_sheet({'唯一表': sources['工资_1月']}, sheet_name_hint='不存在')


def test_shared_entry_passes_selected_period_and_keyword_only_parameters():
    def main(input_folder, output_folder, *, salary_year=2020, salary_month=1, standard_hours=174):
        return input_folder, output_folder, salary_year, salary_month, standard_hours
    result = invoke_script_main(main, {'input_folder': 'in', 'output_folder': 'out',
                                      'salary_year': 2026, 'salary_month': 9, 'monthly_standard_hours': 168})
    assert result == ('in', 'out', 2026, 9, 168)


def load_chat_iteration():
    path = Path(__file__).parent / 'api' / 'training_chat.py'
    node = next(n for n in ast.parse(path.read_text(encoding='utf-8')).body
                if isinstance(n, ast.FunctionDef) and n.name == '_run_single_iteration')
    namespace = {'__package__': 'backend.api', '__file__': str(path), 'Path': Path,
                 'Dict': Dict, 'Any': Any, 'tempfile': tempfile, 'shutil': shutil,
                 'os': os, 'time': time, 'logger': logging.getLogger('chain-test')}
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(path), 'exec'), namespace)
    return namespace[node.name]


def test_chat_training_chain_uses_selected_parameters_and_cached_formulas(tmp_path, monkeypatch):
    from backend.sandbox.code_sandbox import CodeSandbox
    # Run sandbox core directly; native worker isolation has separate real-process tests.
    import io
    def execute(self, code, environment):
        return self._run_script_core(code, environment, io.StringIO(), io.StringIO())
    monkeypatch.setattr(CodeSandbox, 'execute_script', execute)
    monkeypatch.setenv('_IN_SUBPROCESS_WORKER', '1')
    source = tmp_path / 'source'
    source.mkdir()
    expected = tmp_path / 'expected.xlsx'
    workbook(expected, {'结果': [('001', '=9*10')]})
    original = expected.read_bytes()
    code = '''import openpyxl
import os
salary_month = 1
def main(input_folder, output_folder, salary_month=1):
    wb = openpyxl.Workbook()
    wb.active.title = "结果"
    wb.active.append(["工号", "金额"])
    wb.active.append(["001", "=" + str(salary_month) + "*10"])
    wb.save(os.path.join(output_folder, "result.xlsx"))
'''
    result = load_chat_iteration()(1, code, 'test', str(source), str(expected), 1,
        salary_year=2026, salary_month=9, rules_content='主键：工号',
        expected_structure={'sheets': {'结果': {'headers': {'工号': 'A', '金额': 'B'}}}})
    try:
        assert result['success'] and result['accuracy'] == 1, result
        assert expected.read_bytes() == original
        out = openpyxl.load_workbook(Path(result['output_dir']) / 'result.xlsx', data_only=True)
        assert out['结果']['B2'].value == 90
        out.close()
    finally:
        shutil.rmtree(Path(result['output_dir']).parent, ignore_errors=True)


def test_header_mapping_cannot_silently_drop_columns():
    from excel_parser import ExcelRegion
    from backend.utils.fast_header_matcher import FastHeaderMatcher
    region = ExcelRegion(head_data={'工号': 'A', '原工号': 'B'}, data=[{'A': '1', 'B': '2'}])
    mapping = {'source.xlsx': {'file_path': 'source.xlsx', 'expected_file': 'staff.xlsx',
        'sheet_mapping': {'数据': '数据'}, 'needs_rewrite': True, 'header_mapping': {'原工号': '工号'}}}
    with pytest.raises(ValueError, match='多个源列'):
        FastHeaderMatcher()._build_pre_loaded_from_memory(mapping,
            {('source.xlsx', '数据'): SimpleNamespace(regions=[region])})


def test_empty_region_retains_schema_and_columns():
    from excel_parser import ExcelRegion
    from backend.utils.data_helpers import convert_region_to_dataframe
    region = ExcelRegion(head_data={'工号': 'A', '金额': 'B'}, data=[])
    result = convert_region_to_dataframe(region)
    assert result.empty and list(result.columns) == ['工号', '金额']


@pytest.mark.parametrize('resource_failure', [False, True])
def test_formula_training_revalidates_history_and_isolates_sample(tmp_path, monkeypatch, resource_failure):
    from backend.ai_engine.training_engine import TrainingEngine
    import backend.ai_engine.training_engine as engine_module
    import backend.ai_engine.formula_code_generator as generator_module
    from backend.utils import training_validation
    expected = tmp_path / '薪资汇总表.xlsx'
    workbook(expected, {'结果': [('001', 20)]})
    original = expected.read_bytes()
    source_dir = tmp_path / 'source'
    source_dir.mkdir()
    source = source_dir / 'staff.xlsx'
    workbook(source, {'数据': [('001', 10)]})
    log_dir = tmp_path / 'logs'
    log_dir.mkdir()
    builder = MagicMock()
    builder.load_source_data.return_value = {'sheets': {}}
    builder.source_sheets = {}
    generator = MagicMock()
    generator.formula_builder = builder
    generator.generate_code.return_value = ('generated code', '')
    monkeypatch.setattr(generator_module, 'FormulaCodeGenerator', lambda **kw: generator)
    monkeypatch.setattr(training_validation, 'finalize_training_output', lambda *args: {'calculated': True})
    comparison = {'comparison_complete': True, 'success': False, 'total_cells': 1,
                  'matched_cells': 0, 'total_differences': 1, 'field_diff_samples': {}}
    monkeypatch.setattr(engine_module, 'compare_excel_files', lambda **kw: pytest.fail('训练必须按 Sheet 名选择结果'))
    monkeypatch.setattr(engine_module, 'compare_excel_files_multi_sheet', lambda **kw: comparison)
    seen = []
    def execute(code, environment):
        folder = Path(environment['output_folder'])
        seen.append(folder)
        assert folder != expected.parent
        if resource_failure:
            return {'success': False, 'resource_failure': True, 'error': '内存超限'}
        workbook(folder / expected.name, {'结果': [('001', 10)]})
        return {'success': True}
    logger = MagicMock()
    logger.log_dir = log_dir
    def save_output(path, **kwargs):
        saved = log_dir / 'saved.xlsx'
        shutil.copy(path, saved)
        return str(saved)
    logger.save_output_excel.side_effect = save_output
    engine = SimpleNamespace(
        training_logger=logger, logger=logging.getLogger('formula-test'),
        ai_provider=object(), stream_callback=None, max_iterations=3 if resource_failure else 1,
        training_perfect_threshold=1.0, training_success_threshold=0.95,
        file_passwords={}, salary_year=2026, salary_month=9, monthly_standard_hours=168,
        tenant_id='test', sandbox=SimpleNamespace(execute_script=execute),
        _load_historical_best=lambda tenant: {'best_score': 1.0, 'best_code': 'old code'},
        _save_historical_best=MagicMock(), _format_detailed_diff=lambda *a, **k: '金额有差异',
        _db_record_iteration=MagicMock(), _db_complete_session=MagicMock())
    args = (engine, [str(source)], str(expected), '主键：工号', {},
            {'sheets': {'结果': {'headers': {'工号': 'A', '金额': 'B'}}}})
    if resource_failure:
        with pytest.raises(training_validation.TrainingResourceFailure, match='内存超限'):
            TrainingEngine._train_formula_mode(*args)
    else:
        result = TrainingEngine._train_formula_mode(*args)
        assert result['success'] is False and result['current_score'] == 0
        assert result['best_code'] != 'old code'
    assert generator.generate_code.call_count == 1
    assert len(seen) == 1 and not seen[0].exists()
    assert expected.read_bytes() == original
