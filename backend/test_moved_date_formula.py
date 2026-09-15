import pytest
import openpyxl
from datetime import datetime
from backend.utils.output_postprocess import finalize_output_workbook


@pytest.mark.parametrize('destination', [4, 10])
def test_today_keeps_template_date_format_when_row_moves(tmp_path, destination):
    template, output = tmp_path / 'template.xlsx', tmp_path / 'output.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = '人员'
    for row in range(2, 7):
        ws.cell(row, 1, str(row))
    ws['B7'] = '=TODAY()'
    ws['B7'].number_format = 'yyyy-mm-dd'
    wb.save(template)
    ws['B7'] = None
    ws.cell(destination, 2, '=TODAY()').number_format = 'General'
    wb.save(output)
    wb.close()
    finalize_output_workbook(str(output), str(template))
    formulas = openpyxl.load_workbook(output)
    values = openpyxl.load_workbook(output, data_only=True)
    try:
        assert formulas.active.cell(destination, 2).value == '=TODAY()'
        assert formulas.active.cell(destination, 2).number_format == 'yyyy-mm-dd'
        assert isinstance(values.active.cell(destination, 2).value, datetime)
    finally:
        formulas.close()
        values.close()
