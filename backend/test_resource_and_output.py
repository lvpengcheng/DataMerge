"""Resource lifecycle and final workbook correctness, with no AI or tenant data."""
import asyncio
import io
import os
import subprocess
import threading
import time
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

from backend.utils import resource_guard as guard
from backend.utils import subprocess_runner as runner
from backend.utils.upload_stream import ExcelWorkGate
from backend.utils.compute_preload_cache import save_preload, load_preload


@pytest.fixture(autouse=True)
def isolated_resources(monkeypatch):
    monkeypatch.setattr(guard, "_fault", None)
    monkeypatch.setattr(runner, "_semaphore", threading.Semaphore(1))
    monkeypatch.setenv("SUBPROCESS_POOL_SIZE", "0")
    monkeypatch.setenv("SUBPROCESS_QUEUE_TIMEOUT", "2")
    monkeypatch.delenv("_IN_SUBPROCESS_WORKER", raising=False)


def test_memory_budget_honors_host_and_explicit_caps(monkeypatch):
    monkeypatch.setattr(guard, "memory_snapshot_mb", lambda: (2048, 1200))
    assert 800 < guard.memory_limit_mb(0) < 1000
    assert guard.memory_limit_mb(512) == 512
    monkeypatch.setattr(guard, "memory_snapshot_mb", lambda: (8192, 6400))
    assert guard.memory_limit_mb(0) > 4096
    assert guard.memory_limit_mb(2048) == 2048


def test_memory_admission_does_not_start_when_vm_is_full(monkeypatch):
    monkeypatch.setattr(guard, "memory_snapshot_mb", lambda: (2048, 200))
    with pytest.raises(RuntimeError, match="内存不足"):
        guard.wait_for_memory(0)


def test_unkillable_child_trips_circuit_with_bounded_wait(monkeypatch):
    waits = []
    def wait(timeout):
        waits.append(timeout)
        raise subprocess.TimeoutExpired("blocked", timeout)
    proc = SimpleNamespace(pid=987654, poll=lambda: None, wait=wait)
    monkeypatch.setattr(guard, "kill_tree", lambda pid: None)
    assert guard.terminate_process(proc, grace_seconds=0.01) is False
    assert waits == [0.01]
    with pytest.raises(RuntimeError, match="暂停派发"):
        guard.ensure_healthy()
    result = runner.run_in_fresh_subprocess("builtins:len", ([1],))
    assert not result.success and "暂停派发" in result.error


def test_gate_fifo_and_cancelled_waiter_releases_slot(monkeypatch):
    monkeypatch.setattr(guard, "admission_ready", lambda: True)
    async def scenario():
        gate = ExcelWorkGate(1)
        await gate.acquire()
        order = []
        async def work(i):
            async with gate:
                order.append(i)
        tasks = [asyncio.create_task(work(i)) for i in range(4)]
        await asyncio.sleep(0.01)
        tasks[1].cancel()
        gate.release()
        await asyncio.gather(*tasks, return_exceptions=True)
        assert order == [0, 2, 3]
        assert gate.pending == 0 and gate._value == 1
    asyncio.run(scenario())


def test_gate_memory_wait_cancel_and_queue_limit(monkeypatch):
    monkeypatch.setattr(guard, "admission_ready", lambda: False)
    monkeypatch.setenv("EXCEL_QUEUE_MAX_PENDING", "1")
    async def scenario():
        gate = ExcelWorkGate(1)
        task = asyncio.create_task(gate.acquire())
        await asyncio.sleep(0.02)
        with pytest.raises(RuntimeError, match="队列已满"):
            await gate.acquire()
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)
        assert gate._value == 1 and gate.pending == 0
    asyncio.run(scenario())


def test_gate_wait_has_deadline(monkeypatch):
    monkeypatch.setenv("EXCEL_QUEUE_MAX_WAIT", "1")
    monkeypatch.setattr(guard, "admission_ready", lambda: False)
    async def scenario():
        gate = ExcelWorkGate(1)
        with pytest.raises(TimeoutError):
            await gate.acquire()
        assert gate._value == 1 and gate.pending == 0
    asyncio.run(scenario())


def test_old_pool_reader_does_not_invalidate_replacement():
    old = SimpleNamespace(stdout=io.BytesIO(b""))
    current = object()
    slot = runner._WorkerSlot.__new__(runner._WorkerSlot)
    slot.proc, slot.lock, slot.dead, slot.pending = current, threading.Lock(), False, {"new": {}}
    slot._read_loop(old)
    assert slot.dead is False and "new" in slot.pending


def test_pool_does_not_restart_before_reaping(monkeypatch):
    slot = runner._WorkerSlot.__new__(runner._WorkerSlot)
    slot.proc = SimpleNamespace(pid=987655, poll=lambda: None)
    slot.lock, slot.dead, slot.pending = threading.Lock(), False, {}
    monkeypatch.setattr(runner, "terminate_process", lambda proc: False)
    monkeypatch.setattr(slot, "restart", lambda: pytest.fail("must not start replacement"))
    assert slot.kill_and_restart() is False
    assert slot.dead is True


def test_real_worker_success_and_error():
    result = runner.run_in_fresh_subprocess("builtins:len", ([1, 2, 3],), timeout=20, max_memory_mb=1024)
    assert result.success and result.result == 3, result.error
    result = runner.run_in_fresh_subprocess("builtins:int", ("bad",), timeout=20, max_memory_mb=1024)
    assert not result.success and "ValueError" in result.error


def test_real_worker_timeout_is_reaped():
    result = runner.run_in_fresh_subprocess("time:sleep", (30,), timeout=0.4, max_memory_mb=1024)
    assert result.timed_out and not result.success and not result.termination_failed
    assert result.duration < 12


def test_real_async_cancellation_releases_actual_process_slot():
    async def scenario():
        task = asyncio.create_task(runner.run_in_fresh_subprocess_async(
            "time:sleep", (30,), timeout=40, max_memory_mb=1024))
        await asyncio.sleep(0.3)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        deadline = time.monotonic() + 12
        while time.monotonic() < deadline:
            if runner._semaphore.acquire(blocking=False):
                runner._semaphore.release()
                return
            await asyncio.sleep(0.1)
        pytest.fail("child still owns execution slot")
    asyncio.run(scenario())


def test_precheck_cache_invalidates_on_source_or_mapping_change(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    file = source / "staff.xlsx"
    file.write_bytes(b"original")
    context = ({"files": {"staff.xlsx": {}}}, None, {"sheets": {}})
    data = {"staff": {"df": pd.DataFrame({"工号": ["001"]})}}
    save_preload(source, data, {"staff.xlsx": {}}, context)
    mapping, loaded = load_preload(source, context)
    assert loaded["staff"]["df"].iloc[0, 0] == "001"
    assert mapping == {"staff.xlsx": {}}
    assert load_preload(source, ({}, None, {})) is None
    file.write_bytes(b"changed-data")
    assert load_preload(source, context) is None


def test_finalize_single_open_save_and_cached_cross_sheet_formula(tmp_path, monkeypatch):
    from backend.utils import output_postprocess as post
    path = tmp_path / "result.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "结果"
    wb.active.append(["工号", "金额", "入职日期"])
    wb.active.append(["001", "=SUM('源_工资'!B2:B3)", "20260906"])
    ws = wb.create_sheet("源_工资")
    ws.append(["工号", "金额"])
    ws.append(["001", 60.74])
    ws.append(["002", 80.25])
    wb.save(path)
    structure = {"sheets": {"结果": {"column_schemas": {
        "工号": {"field_type": "text", "number_format": "@"},
        "金额": {"field_type": "decimal", "number_format": "0.00"},
        "入职日期": {"field_type": "date", "number_format": "yyyy-mm-dd"},
    }}}}
    opened, saves = [], []
    original_open = post._OutputSession.open
    class CountedBook:
        def __init__(self, real): self.real = real
        def __getattr__(self, name): return getattr(self.real, name)
        def Save(self, dest):
            saves.append(str(dest))
            return self.real.Save(dest)
    def counted_open(session, name):
        key = os.path.normcase(os.path.abspath(str(name)))
        new = key not in session.books
        handle = original_open(session, name)
        if new:
            opened.append(key)
            handle.workbook = CountedBook(handle.workbook)
        return handle
    monkeypatch.setattr(post._OutputSession, "open", counted_open)
    result = post.finalize_output_workbook(path, expected_structure=structure,
                                           sheet_name_map={"源_工资": "源_O'Brien"})
    assert result["calculated"] and len(opened) == len(saves) == 1
    formula = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    assert values["结果"]["B2"].value == pytest.approx(140.99)
    assert values["结果"]["C2"].value == datetime(2026, 9, 6)
    assert "O''Brien" in formula["结果"]["B2"].value
    assert values["源_O'Brien"]["B2"].value == 60.74
    assert values["结果"]["A2"].value == "001"
    formula.close()
    values.close()


def test_finalize_failure_preserves_original_file(tmp_path, monkeypatch):
    from backend.utils import output_postprocess as post
    path = tmp_path / "result.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=1+2"
    wb.save(path)
    original = path.read_bytes()
    monkeypatch.setattr(post, "normalize_key_columns_to_text", lambda p: (_ for _ in ()).throw(RuntimeError("disk failed")))
    with pytest.raises(RuntimeError, match="disk failed"):
        post.finalize_output_workbook(path)
    assert path.read_bytes() == original


def test_single_source_open_preserves_original_and_repairs_numeric_date_style(tmp_path, monkeypatch):
    import excel_parser
    from backend.utils.fast_header_matcher import FastHeaderMatcher
    path = tmp_path / "staff.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "数据"
    wb.active.append(["工号", "金额", "入职日期"])
    for i in range(1, 10):
        wb.active.append([f"{i:03}", 60.74, datetime(2026, 9, 6)])
        wb.active.cell(i + 1, 2).number_format = "yyyy-mm-dd"
    wb.save(path)
    original = path.read_bytes()
    opens = []
    real = excel_parser._licensed_workbook
    def record(*args, **kwargs):
        opens.append(str(args[0]))
        return real(*args, **kwargs)
    monkeypatch.setattr(excel_parser, "_licensed_workbook", record)
    structure = {"files": {"staff.xlsx": {"sheets": {"数据": {
        "headers": {"工号": "A", "金额": "B", "入职日期": "C"}}}}}}
    ok, error, mapping, loaded = FastHeaderMatcher().match_parse_and_prepare(
        structure, [str(path)], manual_headers={"staff.xlsx": {"数据": ["A1", "C1"]}})
    assert ok, error
    assert len(opens) == 1
    assert path.read_bytes() == original
    assert loaded["数据"]["df"].iloc[0]["金额"] == pytest.approx(60.74)
    assert loaded["数据"]["df"].iloc[0]["工号"] == "001"


def test_finalize_1904_date_serial_is_not_shifted(tmp_path):
    from backend.utils.output_postprocess import finalize_output_workbook
    from openpyxl.utils.datetime import MAC_EPOCH
    path = tmp_path / "mac.xlsx"
    wb = openpyxl.Workbook()
    wb.epoch = MAC_EPOCH
    wb.active.title = "结果"
    wb.active.append(["入职日期"])
    wb.active.append([44809])  # 2026-09-06 in the 1904 epoch
    wb.save(path)
    finalize_output_workbook(path, expected_structure={"sheets": {"结果": {"column_schemas": {
        "入职日期": {"field_type": "date", "number_format": "yyyy-mm-dd"}}}}})
    out = openpyxl.load_workbook(path, data_only=True)
    assert out.active["A2"].value == datetime(2026, 9, 6)
    out.close()


def test_legacy_rule_excel_keeps_formula_evidence(tmp_path):
    import aspose_init
    aspose_init.ensure_license()
    from Aspose.Cells import Workbook
    from backend.utils.formula_evidence import collect_formula_evidence
    from backend.ai_engine.document_parser import DocumentParser
    path = tmp_path / "sample.xls"
    wb = Workbook()
    try:
        wb.Worksheets[0].Cells["A1"].PutValue("金额")
        wb.Worksheets[0].Cells["A2"].PutValue(60.74)
        wb.Worksheets[0].Cells["B2"].Formula = "=ROUND(A2*0.5%,2)"
        wb.Save(str(path))
    finally:
        wb.Dispose()
    evidence = collect_formula_evidence(path)
    assert evidence["Sheet1"]["formulas"]["B2"] == "=ROUND(A2*0.5%,2)"
    assert "B2: =ROUND(A2*0.5%,2)" in DocumentParser().parse_document(str(path))


def test_encoding_repair_never_replaces_formula_or_number_format(tmp_path):
    from backend.utils.source_normalizer import _restore_decoded_text
    target, decoded = tmp_path / "target.xlsx", tmp_path / "decoded.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["garbled", "金额"])
    wb.active.append([1, "=A2*60.74"])
    wb.active["B2"].number_format = "0.00"
    wb.save(target)
    wb.active["A1"] = "工号"
    wb.active["B2"] = 60.74
    wb.active["B2"].number_format = "General"
    wb.save(decoded)
    _restore_decoded_text(str(decoded), str(target))
    out = openpyxl.load_workbook(target, data_only=False)
    assert out.active["A1"].value == "工号"
    assert out.active["B2"].value == "=A2*60.74"
    assert out.active["B2"].number_format == "0.00"
    out.close()


def test_finalize_template_format_repairs_share_live_workbook(tmp_path):
    from backend.utils.output_postprocess import finalize_output_workbook
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "结果"
    wb.active.append(["工号", "金额"])
    wb.active.append(["001", None])
    wb.active["B2"].number_format = "0.00"
    wb.save(template)
    wb.active["B2"] = "=SUM('源_数据'!B2:B3)"
    wb.active["B2"].number_format = "General"
    source = wb.create_sheet("源_数据")
    source.append(["工号", "金额"])
    source.append([1, 10])
    source.append([2, 20])
    wb.save(output)
    finalize_output_workbook(output, template_path=str(template))
    result = openpyxl.load_workbook(output, data_only=True)
    assert result["结果"]["B2"].value == 30
    assert result["结果"]["B2"].number_format == "0.00"
    assert result["源_数据"]["A2"].value == "1"
    result.close()
