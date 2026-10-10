"""用户批量导入与模板；不修改已有账户。"""
import io
import re
from copy import copy
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill
from sqlalchemy.exc import IntegrityError

from ..database.models import User, Role, Organization
from ..auth.utils import get_password_hash

HEADERS = ["用户名", "初始密码", "显示名", "邮箱", "电话", "角色", "组织名称", "启用状态"]


def is_privileged_role(role):
    permissions = role.permissions if isinstance(role.permissions, dict) else {}
    return role.name.lower() == "admin" or permissions.get("admin") is True


def build_user_import_template(db):
    book = load_workbook(Path(__file__).parent / "templates" / "user_import_template.xlsx")
    sheet = book["用户导入"]
    references = book["填写选项"]
    for column, header in enumerate(HEADERS, 1):
        sheet.cell(1, column, header)
    references.cell(1, 4, "组织名称")
    references.column_dimensions["D"].width = references.column_dimensions["E"].width
    references.delete_cols(5)
    instructions = book["填写说明"]
    for cells in instructions.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("组织ID", "组织名称")
    # 每次下载重新填充选项，避免静态模板残留已删除的组织或角色。
    for cells in references.iter_rows(min_row=2):
        for cell in cells:
            cell.value = None
    sheet.data_validations.dataValidation = [
        validation for validation in sheet.data_validations.dataValidation
        if str(validation.sqref) == "H2:H501"
    ]
    roles = [role for role in db.query(Role).order_by(Role.name).all() if not is_privileged_role(role)]
    organizations = db.query(Organization).filter_by(is_active=True).order_by(Organization.name, Organization.id).all()
    for row, role in enumerate(roles, 2):
        references.cell(row, 1, role.name).data_type = "s"
        references.cell(row, 2, role.description or "").data_type = "s"
    for row, organization in enumerate(organizations, 2):
        references.cell(row, 4, organization.name).data_type = "s"
    for column, reference_column, count in [("F", "A", len(roles)), ("G", "D", len(organizations))]:
        if count:
            validation = DataValidation(type="list", formula1=f"'填写选项'!${reference_column}$2:${reference_column}${count + 1}", allow_blank=column == "G")
            validation.errorTitle = "请选择有效选项"
            validation.error = "请从填写选项中选择，角色不能为系统管理员。"
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}501")
    sheet.freeze_panes = "A2"
    for worksheet in book:
        worksheet.sheet_view.showGridLines = True
        for cell in worksheet[1]:
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            cell.alignment = alignment
        for cells in worksheet.iter_rows(min_row=2):
            for cell in cells:
                cell.fill = PatternFill()
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


def import_users(content, filename, db):
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="请上传填写后的 .xlsx 模板")
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件不能为空且不能超过 5MB")
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception:
        raise HTTPException(status_code=400, detail="无法读取 Excel，请重新下载模板并填写")
    try:
        sheet = book["用户导入"] if "用户导入" in book.sheetnames else book.active
        rows = sheet.iter_rows()
        headers = [str(cell.value or "").strip() for cell in next(rows, [])]
        if headers != HEADERS:
            raise HTTPException(status_code=400, detail="表头不符合模板，请保留模板中的八列及顺序")
        roles = {role.name: role for role in db.query(Role).all()}
        organizations = {}
        for org in db.query(Organization).filter_by(is_active=True).all():
            organizations.setdefault(org.name.strip(), []).append(org)
        existing = {username for (username,) in db.query(User.username).all()}
        seen = set()
        prepared, errors = [], []
        nonempty = 0
        for row_number, cells in enumerate(rows, 2):
            if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                continue
            nonempty += 1
            if nonempty > 500 or row_number > 5001:
                raise HTTPException(status_code=400, detail="每次最多导入 500 个用户，请分批导入")
            issues = []
            if any(cell.data_type == "f" for cell in cells):
                issues.append("请填写实际值，不能使用公式")
            values = [str(cell.value) if cell.value is not None else "" for cell in cells]
            values += [""] * (8 - len(values))
            username, password, display, email, phone, role_name, org_text, active = values[:8]
            username, display, email, phone, role_name, org_text, active = [v.strip() for v in (username, display, email, phone, role_name, org_text, active)]
            if not username or len(username) > 50:
                issues.append("用户名必填且不超过 50 字符")
            if username in existing or username in seen:
                issues.append("用户名已存在或在文件中重复")
            seen.add(username)
            if not password.strip() or len(password) < 6 or len(password.encode("utf-8")) > 72:
                issues.append("初始密码必填，至少 6 字符且不超过 72 字节")
            if len(display) > 100 or len(email) > 100 or len(phone) > 20:
                issues.append("显示名/邮箱不超过 100 字符，电话不超过 20 字符")
            if email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
                issues.append("邮箱格式无效")
            role = roles.get(role_name)
            if not role:
                issues.append("角色必填且必须是现有角色名称")
            elif is_privileged_role(role):
                issues.append("批量导入不能使用系统管理员或具有超级管理员权限的角色")
            org_id = None
            if org_text:
                matches = organizations.get(org_text, [])
                if not matches:
                    issues.append("组织名称不存在或组织已停用")
                elif len(matches) > 1:
                    issues.append("组织名称对应多个有效组织，请先调整同名组织后再导入")
                else:
                    org_id = matches[0].id
            if active not in ("", "启用", "禁用"):
                issues.append("启用状态只能填写启用或禁用，留空默认启用")
            if issues:
                errors.append({"row": row_number, "username": username, "message": "；".join(issues)})
            else:
                prepared.append(dict(username=username, password=password, display_name=display or username,
                                     email=email, phone=phone, role_id=role.id, org_id=org_id, is_active=active != "禁用"))
        if not nonempty:
            raise HTTPException(status_code=400, detail="模板中没有待导入的用户，请从第 2 行开始填写")
        if errors:
            return {"created": 0, "total": nonempty, "errors": errors}
        try:
            for values in prepared:
                password = values.pop("password")
                db.add(User(password_hash=get_password_hash(password), **values))
            db.commit()
        except IntegrityError:
            db.rollback()
            raise HTTPException(status_code=409, detail="用户名已被其他操作创建，本次未导入任何用户，请检查后重试")
        except Exception:
            db.rollback()
            raise
        return {"created": len(prepared), "total": nonempty, "errors": []}
    finally:
        book.close()
