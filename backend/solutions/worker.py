"""Frozen worker bootstrap. Runs as a file under Python -I before application imports."""
import json
import os
import pickle
import shutil
import sys
import tempfile
from pathlib import Path


def main():
    package, mode, params_file, host = map(str, sys.argv[1:5])
    package, host = Path(package), Path(host)
    # The package is never writable by ordinary execution: use a private copy
    # because the existing app initializes log files relative to its source root.
    work = Path(tempfile.mkdtemp(prefix='solution-worker-'))
    try:
        shutil.copytree(package, work / 'package')
        package = work / 'package'
        runtime = package / 'runtime'
        sys.dont_write_bytecode = True
        sys.path.insert(0, str(runtime))
        from backend.solutions.package import verify
        manifest = verify(package)
        if manifest.get('environment') is not None:
            from backend.solutions.environment import fingerprint
            if fingerprint(runtime) != manifest['environment']:
                raise ValueError('隔离工作进程的实际运行环境与方案不一致，已停止计算')
        from backend.solutions.store import materialize
        settings = manifest['settings']
        from backend.solutions.environment import calculation_settings
        for key in calculation_settings():
            os.environ.pop(key, None)
        os.environ.update(settings)
        # Never load current deployment calculation defaults into an old version.
        import dotenv
        dotenv.load_dotenv = lambda *a, **k: False
        os.environ['_IN_SUBPROCESS_WORKER'] = '1'
        os.environ['COMPUTE_DELETE_RUN_FILES_AFTER_FINISH'] = 'false'
        os.chdir(host)
        # License stays deployment-local and is not part of exportable artifacts.
        import aspose_init
        aspose_init._lic_path = str(host / 'libs/Aspose.Total.NET.lic')
        aspose_init.ensure_license()
        from backend.utils import template_resolver
        def resolve(*args, **kwargs):
            import hashlib
            code = kwargs.get('script_code') or ''
            override = kwargs.get('uploaded_override')
            if override:
                return override
            key = hashlib.sha256(code.encode()).hexdigest()
            relative = manifest['templates'].get(key)
            if relative:
                return str(package / relative)
            name, _, baked = template_resolver.extract_template_ref(code)
            if name or baked:
                raise ValueError('冻结方案未登记脚本所需模板')
            return None
        template_resolver.resolve_template_path = resolve
        import backend.app.main as app
        metadata = materialize(manifest['metadata'], package)
        app._load_script_info_for_precheck = lambda *a, **k: metadata
        from backend.storage.storage_manager import StorageManager
        app.storage_manager = StorageManager(str(host / 'tenants'))
        if mode == 'compute':
            from backend.compute import compute_worker
            params = json.loads(Path(params_file).read_text(encoding='utf-8'))
            if params.get('_solution', {}).get('digest') != manifest['digest']:
                raise ValueError('计算参数与方案版本不一致')
            # 旧封存运行时尚未接收显示名参数；仅在本次隔离进程固定交付名称，兼容已有方案包。
            if params.get('script_display_name'):
                app._resolve_script_display_name = lambda *a, **k: params['script_display_name']
            frozen_inputs = params['_solution']
            if frozen_inputs.get('inputs_dir'):
                from backend.solutions.inputs import install
                install(frozen_inputs['inputs_dir'], params['tenant_id'], frozen_inputs['inputs_digest'])
            params['script_content'] = (package / 'business/script.py').read_bytes().decode('utf-8')
            pinned_params = work / 'compute.json'
            pinned_params.write_text(json.dumps(params, ensure_ascii=False), encoding='utf-8')
            # compute_worker must not reload environment (dotenv is disabled).
            sys.argv = ['compute_worker', str(pinned_params)]
            compute_worker.main()
        else:
            from backend.utils.subprocess_worker import _run_task
            with open(params_file, 'rb') as stream:
                payload = pickle.load(stream)
            if payload['entry'] == 'backend.solutions.worker:validate':
                payload['args'][0]['package'] = str(package)
                payload['args'][0]['_verified_manifest'] = manifest
            if payload.get('args') and isinstance(payload['args'][0], dict):
                data = payload['args'][0]
                if data.get('_solution') and data['_solution']['digest'] != manifest['digest']:
                    raise ValueError('预检参数与方案版本不一致')
                frozen_inputs = data.get('_solution') or {}
                if frozen_inputs.get('inputs_dir'):
                    from backend.solutions.inputs import install
                    install(frozen_inputs['inputs_dir'], frozen_inputs['tenant_id'], frozen_inputs['inputs_digest'])
                if payload['entry'] in ('backend.app.main:_compute_upload_precheck_subprocess', 'backend.solutions.worker:confirm'):
                    data['script_content'] = (package / 'business/script.py').read_bytes().decode('utf-8')
                if payload['entry'] == 'backend.app.main:_compute_upload_precheck_subprocess':
                    for key in ('source_structure', 'manual_headers', 'expected_structure', 'use_history'):
                        value = metadata.get(key)
                        if isinstance(value, str) and key != 'use_history':
                            value = json.loads(value)
                        data[key] = value
            _run_task(payload)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def precheck(payload):
    import backend.app.main as app
    return app._compute_upload_precheck_subprocess(payload)


def confirm(payload):
    from backend.utils.compute_ingest import read_meta, resolve_with_confirmations
    meta = read_meta(payload['session_dir'])
    if meta is None:
        raise ValueError('计算会话解析数据不存在')
    kwargs = {k: v for k, v in payload.items() if k not in ('session_dir', '_solution')}
    return resolve_with_confirmations(meta, **kwargs)


def save_confirmation(payload):
    from backend.utils.confirmed_source_mapping import save_confirmation_state
    return save_confirmation_state(payload['session_dir'], payload['body'], payload['params'])


def validate(payload):
    """Execute the frozen business script and production output finalization."""
    import hashlib
    import importlib.util
    from backend.solutions.store import materialize
    from backend.solutions.package import verify
    from backend.utils.script_entry import invoke_script_main
    from backend.utils.multi_workbook import workbook_definitions, validate_workbooks
    from backend.utils.training_validation import finalize_training_output
    from backend.utils.excel_comparator import compare_excel_files_multi_sheet, require_complete_comparison
    from backend.utils.result_selection import pick_result_output
    # package path is explicit: store.ROOT here points at the private runtime.
    package = Path(payload['package'])
    manifest = payload.get('_verified_manifest') or verify(package)
    records = materialize(json.loads((package / 'archive/records.json').read_text(encoding='utf-8')), package)
    from backend.solutions.inputs import install
    install(package / 'archive/context', payload['tenant_id'])
    session = records.get('session') or {}
    cfg = session.get('config') or {}
    source = cfg.get('source_dir')
    if not source or not Path(source).is_dir():
        raise ValueError('缺少原始训练源文件，无法完成方案验收')
    structure = manifest['metadata'].get('expected_structure') or {}
    structure = materialize(structure, package)
    if isinstance(structure, str):
        structure = json.loads(structure)
    code = (package / 'business/script.py').read_bytes().decode('utf-8')
    with tempfile.TemporaryDirectory(prefix='solution-validation-') as temp:
        source_copy = Path(temp) / 'source'
        shutil.copytree(source, source_copy)
        output = Path(temp) / 'output'
        output.mkdir()
        metadata = materialize(manifest['metadata'], package)
        source_structure = metadata.get('source_structure')
        if isinstance(source_structure, str):
            source_structure = json.loads(source_structure)
        headers = metadata.get('manual_headers') or {}
        if isinstance(headers, str):
            headers = json.loads(headers)
        preload = None
        execution_source = str(source_copy)
        if source_structure:
            # Replay the same frozen parsing, mapping and execution-copy builders
            # used by compute; do not validate only the generated script body.
            from backend.utils.compute_precheck import precheck_compute, _extract_colmap
            from backend.utils.compute_ingest import read_meta, read_sources, build_preload
            from backend.utils.confirmed_source_mapping import write_execution_sources
            from backend.database.connection import SessionLocal
            with SessionLocal() as db:
                check = precheck_compute(str(source_copy), source_structure, headers, code,
                    payload['tenant_id'], cfg.get('salary_year', session.get('salary_year')),
                    cfg.get('salary_month', session.get('salary_month')), db,
                    use_history=metadata.get('use_history'), expected_structure=structure,
                    # 验收回放使用训练原模板，目标关系固定为训练表自身；
                    # 智算上传的本月模板仍必须经过完整人工确认。
                    confirmed_target_map={key: key for key in (_extract_colmap(code) or {})},
                    in_worker=True, session_dir=temp, defer_base_fill=True)
            if not check.ok:
                raise ValueError('冻结版本的源文件预检未通过，请补齐原始训练资料后重新转方案')
            if check.file_mapping:
                preload = build_preload(read_meta(temp), read_sources(temp), check.file_mapping)
                execution_source = write_execution_sources(Path(temp) / 'confirmed_source',
                    check.file_mapping, preload, source_structure, structure)
        spec = importlib.util.spec_from_file_location('solution_business', package / 'business/script.py')
        module = importlib.util.module_from_spec(spec)
        import backend.app.main as app
        app._inject_sandbox_helpers(module)
        env = {'input_folder': str(execution_source), 'output_folder': str(output),
               'source_files': [p.name for p in source_copy.iterdir() if p.is_file()],
               'manual_headers': headers,
               'file_passwords': {},
               'salary_year': cfg.get('salary_year', session.get('salary_year')),
               'salary_month': cfg.get('salary_month', session.get('salary_month')),
               'monthly_standard_hours': cfg.get('monthly_standard_hours'),
               'tenant_id': payload['tenant_id']}
        module.__dict__.update(env)
        import importlib
        for name in ('math', 're', 'json', 'datetime', 'collections', 'itertools', 'functools',
                     'decimal', 'statistics', 'random', 'string', 'hashlib', 'pandas', 'openpyxl'):
            setattr(module, name, importlib.import_module(name))
        from openpyxl.utils import column_index_from_string, get_column_letter
        module.column_index_from_string, module.get_column_letter = column_index_from_string, get_column_letter
        from backend.utils.historical_data import HistoricalDataProvider
        module.history_provider = HistoricalDataProvider(payload['tenant_id'])
        spec.loader.exec_module(module)
        module.__dict__.update(env)
        if preload:
            module._pre_loaded_source_data = preload
            if hasattr(module, 'load_source_data'):
                module.load_source_data = lambda *a, **k: preload
        tpl = manifest['templates'].get(hashlib.sha256(code.encode()).hexdigest())
        template = str(package / tpl) if tpl else None
        if template:
            module.TEMPLATE_PATH = template
            module._template_override_path = template
        execution = invoke_script_main(module.main, env)
        if isinstance(execution, dict) and execution.get('success') is False:
            raise ValueError(execution.get('error') or '方案脚本报告执行失败')
        if len(workbook_definitions(structure)) > 1:
            return validate_workbooks(output, code, structure, cfg.get('rules_content'), payload['tenant_id'])
        result = pick_result_output(output.glob('*.xlsx'), template)
        expected = cfg.get('expected_file')
        if not result or not expected or not Path(expected).is_file():
            raise ValueError('缺少可比较的结果或训练期望文件')
        finalized = finalize_training_output(str(result), code, structure, template)
        from backend.utils.excel_comparator import extract_primary_keys_from_rules
        rules = manifest['metadata'].get('rules_content')
        comparison = compare_excel_files_multi_sheet(str(result), expected, str(Path(temp) / 'diff.xlsx'),
                                                     primary_keys=extract_primary_keys_from_rules(rules) if rules else None,
                                                     result_calculated=True, calculate_expected_formulas=True)
        require_complete_comparison(comparison)
        comparison['formula_warning'] = bool(finalized.get('formula_warning'))
        return comparison


if __name__ == '__main__':
    main()
