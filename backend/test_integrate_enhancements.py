"""多表整合对比增强功能的回归测试。"""

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from backend.api.tools import (
    _match_scheme_config,
    _required_cols_by_role,
    _resolved_scheme_config,
    _suggest_column_map,
)
from backend.utils.integrate_engine import (
    CLEAR_CELL,
    build_source_indexes,
    collect_source_keys,
    compute_diffs,
    eval_source_expr_cross,
    resolve_overwrites,
)
from backend.utils.integrate_writer import _apply_integration_impl


def test_collect_source_keys_unions_all_source_tables_in_order():
    parsed = {
        "主表.xlsx": {"df": pd.DataFrame(columns=["工号"])},
        "A.xlsx": {"df": pd.DataFrame({"工号": ["A1", "A2", "A3"]})},
        "B.xlsx": {"df": pd.DataFrame({"工号": ["B1", "B2", "B3", "B4"]})},
    }
    keys = collect_source_keys(
        parsed, {"主表.xlsx": "工号", "A.xlsx": "工号", "B.xlsx": "工号"}, "主表.xlsx")
    assert [item["value"] for item in keys] == ["A1", "A2", "A3", "B1", "B2", "B3", "B4"]


def test_collect_source_keys_deduplicates_normalized_keys_across_tables():
    parsed = {
        "主表.xlsx": {"df": pd.DataFrame(columns=["工号"])},
        "A.xlsx": {"df": pd.DataFrame({"工号": [" 001 ", "002", "002"]})},
        "B.xlsx": {"df": pd.DataFrame({"工号": ["001", "003", None]})},
    }
    keys = collect_source_keys(
        parsed, {"A.xlsx": "工号", "B.xlsx": "工号"}, "主表.xlsx")
    assert [item["normalized"] for item in keys] == ["001", "002", "003"]
    assert keys[0]["value"] == " 001 "


def test_union_index_includes_template_values_in_formula_and_preserves_long_ids():
    parsed = {
        "main.xlsx": {"df": pd.DataFrame([{"ID": "001", "Amount": 5, "Card": "110101199001011234", "Name": "模板姓名"}])},
        "a.xlsx": {"df": pd.DataFrame([{"ID": "001", "Amount": 3, "Card": "110101199001011234", "Name": "来源姓名"}])},
        "b.xlsx": {"df": pd.DataFrame([{"ID": "001", "Amount": 7, "Card": "110101199001011234", "Name": "来源姓名"}])},
    }
    indexes = build_source_indexes(
        parsed, {name: "ID" for name in parsed}, "main.xlsx", include_main=True)
    assert eval_source_expr_cross(
        "main.xlsx.Amount+Amount+b.xlsx.Amount", "a.xlsx", indexes, "001") == 15
    assert eval_source_expr_cross(
        "main.xlsx.Card+Card+b.xlsx.Card", "a.xlsx", indexes, "001") == "110101199001011234"
    assert eval_source_expr_cross(
        "main.xlsx.Name+Name+b.xlsx.Name", "a.xlsx", indexes, "001") == "模板姓名"


def test_union_writer_adds_only_missing_keys_and_keeps_row_formulas(tmp_path):
    source = tmp_path / "main.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Main"
    sheet.append(["ID", "Amount", "Calculated", "Note"])
    sheet.append(["E0", 1, "=B2*2", "existing value"])
    sheet.append(["E1", 4, "=B3*2", "second existing value"])
    sheet["B3"].number_format = "0.00"
    sheet["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.append(["Total", "=SUM(B2:B3)", None, None])
    book.save(source)

    indexes = {
        "main.xlsx": {
            "cols": ["Amount"],
            "rows": {"e0": [{"Amount": 1}], "e1": [{"Amount": 4}]},
        },
        "source.xlsx": {
            "cols": ["Amount"],
            "rows": {
                "e0": [{"Amount": 1}],
                "a1": [{"Amount": 2}],
                "b1": [{"Amount": 3}],
            },
        },
    }
    stat = _apply_integration_impl(
        str(source), str(output), "Main",
        {"ID": "A", "Amount": "B", "Calculated": "C", "Note": "D"},
        "ID", 2, 3,
        [{"a_col": "Amount", "source_file": "source.xlsx",
          "source_expr": "main.xlsx.Amount+Amount"}],
        indexes,
        seed_keys=[
            {"normalized": "e0", "value": "E0"},
            {"normalized": "a1", "value": "A1"},
            {"normalized": "b1", "value": "B1"},
        ],
    )

    assert stat["added_rows"] == 2
    assert stat["matched_rows"] == 4
    assert stat["overwritten_cells"] == 4
    result = load_workbook(output, data_only=False)["Main"]
    assert [result.cell(row, 1).value for row in range(2, 6)] == ["E0", "A1", "B1", "E1"]
    assert [result.cell(row, 2).value for row in range(2, 6)] == [2, 2, 3, 4]
    assert result["C3"].value == "=B3*2"
    assert result["C4"].value == "=B4*2"
    assert result["D3"].value is None and result["D4"].value is None
    assert result["D2"].value == "existing value"
    assert result["D5"].value == "second existing value"
    assert result["B6"].value == "=SUM(B2:B5)"
    assert result["B3"].number_format == "0.00"
    assert result["B3"].fill.fgColor.rgb.endswith("FFF2CC")


def test_union_writer_populates_an_empty_main_table(tmp_path):
    source = tmp_path / "empty-main.xlsx"
    output = tmp_path / "empty-result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Main"
    sheet.append(["ID", "Amount"])
    book.save(source)

    keys = [f"A{i}" for i in range(1, 4)] + [f"B{i}" for i in range(1, 5)]
    indexes = {"source.xlsx": {
        "cols": ["Amount"],
        "rows": {key.lower(): [{"Amount": i}] for i, key in enumerate(keys, start=1)},
    }}
    stat = _apply_integration_impl(
        str(source), str(output), "Main", {"ID": "A", "Amount": "B"},
        "ID", 2, 0,
        [{"a_col": "Amount", "source_file": "source.xlsx", "source_expr": "Amount"}],
        indexes,
        seed_keys=[{"normalized": key.lower(), "value": key} for key in keys],
    )

    assert stat["added_rows"] == 7
    assert stat["matched_rows"] == 7
    result = load_workbook(output, data_only=True)["Main"]
    assert [result.cell(row, 1).value for row in range(2, 9)] == keys
    assert [result.cell(row, 2).value for row in range(2, 9)] == list(range(1, 8))


def test_writer_keeps_identity_number_exact_and_formats_it_as_text(tmp_path):
    source = tmp_path / "identity-main.xlsx"
    output = tmp_path / "identity-result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Main"
    sheet.append(["工号", "身份证号码"])
    sheet.append(["A1", None])
    book.save(source)

    identity = "110101199001011234"
    stat = _apply_integration_impl(
        str(source), str(output), "Main", {"工号": "A", "身份证号码": "B"},
        "工号", 2, 2,
        [{"a_col": "身份证号码", "source_file": "source.xlsx", "source_expr": "身份证号码"}],
        {"source.xlsx": {"cols": ["身份证号码"],
                         "rows": {"a1": [{"身份证号码": identity}]}}},
    )
    assert stat["overwritten_cells"] == 1
    result = load_workbook(output, data_only=True)["Main"]
    assert result["B2"].value == identity
    assert result["B2"].number_format == "@"
    assert result["B2"].data_type == "s"


def _scheme_config():
    return {
        "main_fp": "saved-main",
        "source_fps": ["saved-source"],
        "roles": [
            {"fp": "saved-main", "file": "主表.xlsx"},
            {"fp": "saved-source", "file": "5月社保.xlsx"},
        ],
        "cols_by_fp": {
            "saved-main": ["工号", "姓名", "个人养老", "已删除的无关列"],
            "saved-source": ["工号", "2026年5月社保个人养老", "已删除的无关列"],
        },
        "key_map_by_role": {"0": "工号", "1": "工号"},
        "overwrite_pairs": [{
            "a_col": "个人养老",
            "source_fp": "saved-source",
            "source_role": 1,
            "source_expr": "ROUND(2026年5月社保个人养老,2)",
            "source_col": "ROUND(2026年5月社保个人养老,2)",
        }],
        "compare_pairs": [],
        "name_col": "姓名",
        "output_mode": 1,
    }


def test_scheme_ignores_unreferenced_added_and_removed_columns():
    cfg = _scheme_config()
    required = _required_cols_by_role(cfg)
    assert "已删除的无关列" not in required[0]
    assert "已删除的无关列" not in required[1]

    files = [
        {"name": "新主表.xlsx", "fingerprint": "new-main",
         "columns": ["新增列", "工号", "姓名", "个人养老"]},
        {"name": "6月社保.xlsx", "fingerprint": "new-source",
         "columns": ["工号", "2026年6月社保个人养老", "新增列"]},
    ]
    match = _match_scheme_config(cfg, files)
    assert match is not None
    assert match["role_files"] == ["新主表.xlsx", "6月社保.xlsx"]


def test_month_header_change_produces_confirmable_mapping_and_resolved_formula():
    result = _suggest_column_map(
        ["2026年5月社保个人养老"], ["2026年6月社保个人养老"])
    assert result["missing"] == []
    assert result["suggestions"][0]["method"] == "period"

    cfg = _scheme_config()
    files = [
        {"name": "新主表.xlsx", "fingerprint": "new-main",
         "columns": ["工号", "姓名", "个人养老"]},
        {"name": "6月社保.xlsx", "fingerprint": "new-source",
         "columns": ["工号", "2026年6月社保个人养老"]},
    ]
    match = _match_scheme_config(cfg, files)
    resolved = _resolved_scheme_config(cfg, match)
    assert resolved["overwrite_pairs"][0]["source_expr"] == "ROUND(2026年6月社保个人养老,2)"


def test_excel_round_and_if_formula_subset():
    indexes = {
        "6月社保.xlsx": {
            "cols": ["养老", "调整"],
            "rows": {"001": [{"养老": 10.235, "调整": 2}, {"养老": 1, "调整": 3}]},
        }
    }
    assert eval_source_expr_cross(
        "ROUND(养老-调整,2)", "6月社保.xlsx", indexes, "001") == 6.24
    assert eval_source_expr_cross(
        "IF(养老>10,ROUND(养老,2),0)", "6月社保.xlsx", indexes, "001") == 11.24


def test_cross_table_formula_uses_zero_for_partially_missing_keys():
    indexes = {
        "a.xlsx": {"cols": ["金额"], "rows": {"001": [{"金额": 20}]}},
        "b.xlsx": {"cols": ["金额"], "rows": {}},
        "c.xlsx": {"cols": ["金额"], "rows": {"001": [{"金额": 10}]}},
    }
    expr = "金额+b.xlsx.金额+c.xlsx.金额"
    assert eval_source_expr_cross(expr, "a.xlsx", indexes, "001") == 30
    assert resolve_overwrites("001", [{
        "a_col": "合计", "source_file": "a.xlsx", "source_expr": expr,
    }], indexes) == {"合计": 30}
    diffs = compute_diffs(
        pd.DataFrame([{"工号": "001", "合计": 0, "姓名": "测试"}]),
        indexes, "工号",
        [{"a_col": "合计", "source_file": "a.xlsx", "source_expr": expr}],
        "姓名", None,
    )
    assert len(diffs) == 1 and "30.00" in diffs[0]["差异类型"]

    # 默认来源表和 b 表都缺失，c 表存在：两个缺失项按 0，结果仍为 10。
    indexes["a.xlsx"]["rows"] = {}
    assert eval_source_expr_cross(expr, "a.xlsx", indexes, "001") == 10
    assert resolve_overwrites("001", [{
        "a_col": "合计", "source_file": "a.xlsx", "source_expr": expr,
    }], indexes) == {"合计": 10}

    # 所有引用均不存在时，不把纯 0 当成业务结果，保持为空。
    indexes["c.xlsx"]["rows"] = {}
    assert eval_source_expr_cross(expr, "a.xlsx", indexes, "001") is None
    cleared = resolve_overwrites("001", [{
        "a_col": "合计", "source_file": "a.xlsx", "source_expr": expr,
    }], indexes)
    assert cleared.get("合计") is CLEAR_CELL

    # 普通单列映射仍沿用旧语义：源行不存在时不覆盖主表旧值。
    assert resolve_overwrites("001", [{
        "a_col": "合计", "source_file": "a.xlsx", "source_expr": "金额",
    }], indexes) == {}
    assert compute_diffs(
        pd.DataFrame([{"工号": "001", "合计": 0}]), indexes, "工号",
        [{"a_col": "合计", "source_file": "a.xlsx", "source_expr": expr}],
        None, None,
    ) == []
