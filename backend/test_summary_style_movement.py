import openpyxl
from openpyxl.styles import PatternFill, Font
from backend.utils.output_postprocess import finalize_output_workbook


def test_moved_total_does_not_leave_total_style_on_data_row(tmp_path):
    template, result = tmp_path / 'template.xlsx', tmp_path / 'result.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append(['工号', 'Net total 实发合计'])
    for row in range(2, 7):
        ws.append([str(row), 10])
        ws.row_dimensions[row].height = 18
    ws.append(['合计', '=SUM(B2:B6)'])
    ws.row_dimensions[7].height = 32
    for cell in ws[7]:
        cell.fill = PatternFill('solid', fgColor='FFFF00')
        cell.font = Font(bold=True)
    wb.save(template)
    ws['A7'] = 'new'
    ws['B7'] = 10
    ws.append(['new2', 10])
    ws.append(['合计', '=SUM(B2:B8)'])
    wb.save(result)
    wb.close()
    code = "_COL_MAP = {'Summary': {'regions': [{'data_start_row': 2}]}}"
    finalize_output_workbook(str(result), str(template), script_code=code)
    book = openpyxl.load_workbook(result)
    try:
        ws = book['Summary']
        assert ws['A7'].value == 'new'
        assert ws['B9'].value == '=SUM(B2:B8)'
        assert not ws['A7'].font.bold
        assert ws['A7'].fill.patternType != 'solid'
        assert ws['A9'].font.bold
        assert ws['A9'].fill.fgColor.rgb == 'FFFFFF00'
        assert ws.row_dimensions[7].height == 18
        assert ws.row_dimensions[9].height == 32
    finally:
        book.close()
