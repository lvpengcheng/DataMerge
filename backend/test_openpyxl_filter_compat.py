from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from backend.utils.openpyxl_compat import ensure_custom_filter_compat
from backend.utils.script_entry import invoke_script_main

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def filtered_book(value, operator='equal'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['部门', '金额'])
    ws.append(['研发', 12])
    ws.append(['销售', '=B2*2'])
    ws['B2'].number_format = '0.00'
    ws.auto_filter.ref = 'A1:B3'
    raw = BytesIO()
    wb.save(raw)
    wb.close()
    result = BytesIO()
    with ZipFile(raw) as src, ZipFile(result, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(data)
                filters = root.find(NS + 'autoFilter')
                col = ET.SubElement(filters, NS + 'filterColumn', colId='0')
                custom = ET.SubElement(col, NS + 'customFilters')
                ET.SubElement(custom, NS + 'customFilter', operator=operator, val=value)
                data = ET.tostring(root)
            dst.writestr(info, data)
    result.seek(0)
    return result


@pytest.mark.parametrize('value', ['研发', '', '0012', '-3.5', '*研发*', '研?部', '~*'])
def test_custom_filter_roundtrip_preserves_semantics(value):
    ensure_custom_filter_compat()
    wb = openpyxl.load_workbook(filtered_book(value, 'notEqual'))
    cell_filter = wb.active.auto_filter.filterColumn[0].customFilters.customFilter[0]
    assert cell_filter.val == value
    assert cell_filter.operator == 'notEqual'
    saved = BytesIO()
    wb.save(saved)
    wb.close()
    saved.seek(0)
    out = openpyxl.load_workbook(saved)
    assert out.active.auto_filter.filterColumn[0].customFilters.customFilter[0].val == value
    assert out.active['B3'].value == '=B2*2'
    assert out.active['B2'].number_format == '0.00'
    assert out.active['A2'].value == '研发'
    out.close()


def test_old_script_load_via_shared_entry():
    # Same entry used by training and compute, including scripts importing load_workbook directly.
    def main():
        from openpyxl import load_workbook
        with_book = load_workbook(filtered_book('研发'))
        value = with_book.active['B2'].value
        with_book.close()
        return value
    assert invoke_script_main(main, {}) == 12
    assert ensure_custom_filter_compat() is False  # repeat calls do not repatch
