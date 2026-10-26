"""Keep explicit script number formats across legacy defaults and delivery.

The manifest travels inside the xlsx, so renamed/copied/multi-target results
retain their own overrides without global state or filename sidecars.
"""
import inspect
import ast
import copy
import json
import os
import zipfile
from contextlib import contextmanager
from contextvars import ContextVar
from threading import Lock

_state = ContextVar('explicit_script_formats', default=None)
_preserve_template = ContextVar('preserve_template_formats_on_write', default=False)
_lock = Lock()
_installed = False
_MANIFEST = 'datamerge/explicit_number_formats.json'
_DEFAULT_HELPERS = {
    '_restore_number_formats', 'apply_target_column_schemas', '_normalize_key_columns',
    'write_source_sheets', '_append_source_sheets', '_snapshot_number_formats', 'write_source_dataframe',
}


def prefer_script_column_schemas(structure, script_code):
    """Honor edits to the generated schema constant as well as cell assignments."""
    if not script_code:
        return structure
    schemas = None
    for node in ast.parse(script_code).body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == '_TARGET_COLUMN_SCHEMAS'
                                               for t in node.targets):
            try:
                schemas = ast.literal_eval(node.value)
            except (ValueError, TypeError):
                continue
    if not isinstance(schemas, dict):
        return structure
    result = copy.deepcopy(json.loads(structure) if isinstance(structure, str) else (structure or {}))
    for name, columns in schemas.items():
        if isinstance(columns, dict):
            sheet = result.setdefault('sheets', {}).setdefault(name, {})
            sheet.setdefault('column_schemas', {}).update(columns)
    return result


def _explicit_cells(state, workbook=None):
    from openpyxl.utils.cell import column_index_from_string
    for obj, code in list(state.values()):
        ws = obj.parent
        if workbook is not None and ws.parent is not workbook:
            continue
        if ws not in ws.parent.worksheets:
            continue
        if hasattr(obj, 'coordinate'):
            if ws._cells.get((obj.row, obj.column)) is obj:
                yield obj, code
        elif isinstance(obj.index, str):
            first = obj.min or column_index_from_string(obj.index)
            last = obj.max or first
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=first, max_col=last):
                for cell in row:
                    yield cell, code
        else:
            for cell in ws[obj.index]:
                yield cell, code


def protect_explicit_values(function):
    """Old script schema helpers may also coerce values before saving."""
    from functools import wraps
    @wraps(function)
    def wrapped(*args, **kwargs):
        saved = [(cell, cell.value, cell.data_type) for cell, _ in _explicit_cells(_state.get() or {})]
        result = function(*args, **kwargs)
        for cell, value, data_type in saved:
            cell._value, cell.data_type = value, data_type
        return result
    return wrapped


def read_explicit_formats(path):
    with zipfile.ZipFile(path) as archive:
        if _MANIFEST not in archive.namelist():
            return {}
        data = json.loads(archive.read(_MANIFEST))
    if not isinstance(data, dict):
        raise ValueError('脚本格式记录无效')
    return data


def write_explicit_formats(path, formats):
    # Called after an Excel writer replaces the archive; never append duplicates.
    if formats:
        with zipfile.ZipFile(path) as archive:
            if _MANIFEST in archive.namelist():
                if json.loads(archive.read(_MANIFEST)) == formats:
                    return
                # Some writers preserve unknown archive members. Replace the old
                # manifest rather than adding a duplicate (notably after rename).
                import tempfile
                fd, staged = tempfile.mkstemp(suffix='.xlsx', dir=os.path.dirname(os.path.abspath(path)))
                os.close(fd)
                try:
                    with zipfile.ZipFile(staged, 'w') as replacement:
                        for info in archive.infolist():
                            if info.filename != _MANIFEST:
                                replacement.writestr(info, archive.read(info.filename))
                        replacement.writestr(_MANIFEST, json.dumps(formats, ensure_ascii=False),
                                             compress_type=zipfile.ZIP_DEFLATED)
                    archive.close()
                    os.replace(staged, path)
                finally:
                    if os.path.exists(staged):
                        os.remove(staged)
                return
        with zipfile.ZipFile(path, 'a', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(_MANIFEST, json.dumps(formats, ensure_ascii=False))


def apply_explicit_formats(wb, formats, sheet_name_map=None):
    """Apply to Aspose after default repairs and before final calculation/save."""
    from openpyxl.utils.cell import coordinate_to_tuple
    for name, cells in formats.items():
        name = (sheet_name_map or {}).get(name, name)
        ws = wb.Worksheets[name]
        if ws is None:
            raise ValueError(f'脚本指定格式的工作表不存在: {name}')
        for address, code in cells.items():
            row, col = coordinate_to_tuple(address)
            cell = ws.Cells[row - 1, col - 1]
            style = cell.GetStyle()
            style.Custom = code
            cell.SetStyle(style)


def _install_tracking():
    global _installed
    with _lock:
        if _installed:
            return
        from openpyxl.styles.styleable import StyleableObject
        from openpyxl.workbook.workbook import Workbook
        descriptor = vars(StyleableObject)['number_format']
        save = Workbook.save

        class TrackedNumberFormat:
            def __get__(self, instance, owner=None):
                if instance is None:
                    return self
                return descriptor.__get__(instance, owner)

            def __set__(self, instance, value):
                frame = inspect.currentframe().f_back
                try:
                    # v2 retains the existing template format at assignment time.
                    # Explicit number_format writes still pass through unchanged.
                    if (_preserve_template.get() and frame.f_code.co_name == '_bind_value'
                            and frame.f_globals.get('__name__', '').startswith('openpyxl.')):
                        return
                    descriptor.__set__(instance, value)
                    state = _state.get()
                    if state is None or not (hasattr(instance, 'coordinate') or hasattr(instance, 'index')):
                        return
                    # Library defaults (including datetime's implicit formatting)
                    # and generated compatibility helpers are not user overrides.
                    cursor = frame
                    while cursor:
                        module = cursor.f_globals.get('__name__', '')
                        if cursor.f_code.co_name in _DEFAULT_HELPERS:
                            return
                        if cursor is frame and module.startswith('openpyxl.'):
                            return
                        if cursor.f_code.co_name == 'invoke_script_main':
                            break
                        cursor = cursor.f_back
                    state.pop(id(instance), None)
                    state[id(instance)] = (instance, value)
                finally:
                    del frame

        def tracked_save(wb, filename):
            state = _state.get()
            formats = {}
            if state is not None:
                for cell, code in _explicit_cells(state, wb):
                    ws = cell.parent
                    descriptor.__set__(cell, code)
                    formats.setdefault(ws.title, {})[cell.coordinate] = code
            result = save(wb, filename)
            if formats:
                if not isinstance(filename, (str, os.PathLike)):
                    raise ValueError('包含脚本格式设置的工作簿须保存到文件路径')
                write_explicit_formats(filename, formats)
            return result

        StyleableObject.number_format = TrackedNumberFormat()
        Workbook.save = tracked_save
        _installed = True


@contextmanager
def track_script_formats():
    _install_tracking()
    token = _state.set({})
    try:
        yield
    finally:
        _state.reset(token)


@contextmanager
def preserve_template_formats():
    """v2: no format snapshots or restoration; retain defaults while writing."""
    _install_tracking()
    token = _preserve_template.set(True)
    legacy = _state.set(None)
    try:
        yield
    finally:
        _state.reset(legacy)
        _preserve_template.reset(token)
