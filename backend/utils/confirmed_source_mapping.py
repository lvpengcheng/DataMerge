"""人工源映射的校验、合并与执行文件生成。所有坐标均指向真实上传文件。"""
from copy import deepcopy
from pathlib import Path


def merge_confirmation_state(previous, incoming):
    """会话累积确认；按训练目标列替换旧来源，省略的项目保留，空选择保留。"""
    result = {key: deepcopy(value) for key, value in (previous or {}).items() if key in (
        'confirmed_mapping', 'confirmed_renames', 'confirmed_target_map', 'skipped_missing_files', 'skip_history_check')}
    incoming = incoming or {}
    def flatten(mapping):
        rows = {}
        for filename, info in (mapping or {}).get('file_mapping', mapping or {}).items():
            for sheet, target in (info.get('sheet_mapping') or {}).items():
                columns = (info.get('header_mapping_by_sheet') or {}).get(sheet, info.get('header_mapping') or {})
                for column, target_column in columns.items():
                    rows[(info.get('expected_file', filename), target, target_column)] = (filename, sheet, column)
        return rows
    if incoming.get('confirmed_mapping') is not None:
        rows = flatten(result.get('confirmed_mapping'))
        updated = flatten(incoming['confirmed_mapping'])
        skipped = {tuple(item) for item in (result.get('confirmed_mapping') or {}).get('unmatched_columns', [])}
        skipped.update(tuple(item) for item in incoming['confirmed_mapping'].get('unmatched_columns', []))
        skipped.difference_update(updated)
        rows.update(updated)
        for key in skipped:
            rows.pop(key, None)
        files = {}
        for (target_file, target_sheet, target_column), (filename, sheet, column) in rows.items():
            info = files.setdefault(filename, {'expected_file': target_file, 'sheet_mapping': {},
                                               'header_mapping_by_sheet': {}})
            if info['expected_file'] != target_file or info['sheet_mapping'].get(sheet, target_sheet) != target_sheet:
                raise ValueError(f'同一来源被分配给不同训练表: {filename}/{sheet}')
            info['sheet_mapping'][sheet] = target_sheet
            columns = info['header_mapping_by_sheet'].setdefault(sheet, {})
            if column in columns and columns[column] != target_column:
                raise ValueError(f'来源列被重复选择: {filename}/{sheet}/{column}')
            columns[column] = target_column
        result['confirmed_mapping'] = {'file_mapping': files, 'unmatched_columns': [list(item) for item in sorted(skipped)]}
    for name in ('confirmed_renames', 'confirmed_target_map'):
        result[name] = {**(result.get(name) or {}), **(incoming.get(name) or {})}
    result['skipped_missing_files'] = list(dict.fromkeys(
        (result.get('skipped_missing_files') or []) + (incoming.get('skipped_missing_files') or [])))
    result['skip_history_check'] = bool(result.get('skip_history_check') or incoming.get('skip_history_check'))
    return result


def save_confirmation_state(session_dir, incoming, initial=None):
    """只读写服务端会话目录；确认阶段即原子落盘，不等计算任务开始。"""
    import json
    import os
    import tempfile
    path = Path(session_dir) / '_confirmations.json'
    previous = json.loads(path.read_text(encoding='utf-8')) if path.exists() else initial
    state = merge_confirmation_state(previous, incoming)
    fd, staged = tempfile.mkstemp(dir=str(path.parent), suffix='.tmp')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            json.dump(state, stream, ensure_ascii=False)
        os.replace(staged, path)
    finally:
        Path(staged).unlink(missing_ok=True)
    return state


def apply_confirmed_mapping(meta, automatic, confirmed):
    skipped = {tuple(item) for item in (confirmed or {}).get('unmatched_columns', [])}
    confirmed = (confirmed or {}).get('file_mapping', confirmed or {})
    result = deepcopy(automatic or {})
    actual = {(s['file_name'], s['sheet_name']): s for s in meta.input_sheets}
    expected = {(s['file_name'], s['sheet_name']): s for s in meta.train_sheets}
    destinations = set()
    for filename, info in confirmed.items():
        target_file = info.get('expected_file') or filename
        for sheet, target_sheet in (info.get('sheet_mapping') or {}).items():
            source = actual.get((filename, sheet))
            target = expected.get((target_file, target_sheet))
            if source is None or target is None:
                raise ValueError(f'人工映射的文件或 Sheet 不存在: {filename}/{sheet} → {target_file}/{target_sheet}')
            destination = (target_file, target_sheet)
            if destination in destinations:
                raise ValueError(f'多个来源指向同一训练表: {target_file}/{target_sheet}')
            destinations.add(destination)
            headers = dict((info.get('header_mapping_by_sheet') or {}).get(
                sheet, info.get('header_mapping') or {}))
            for col, target_col in headers.items():
                if col not in source['headers'] or target_col not in target['headers']:
                    raise ValueError(f'人工映射的列不存在: {filename}/{sheet}/{col} → {target_col}')
            # 人工选择优先：同名旧列不能抢占已经被选择的目标列。
            claimed = set(headers.values())
            headers.update({c: c for c in source['headers'] if c in target['headers']
                            and c not in headers and c not in claimed and (target_file, target_sheet, c) not in skipped})
            if len(set(headers.values())) != len(headers):
                raise ValueError(f'多个来源列指向同一训练列: {filename}/{sheet}')
            missing = {c for c in target['headers'] if c not in headers.values()
                       and (target_file, target_sheet, c) not in skipped}
            if missing:
                raise ValueError(f'请完成 {target_file}/{target_sheet} 的列匹配: {", ".join(sorted(missing))}')
            # 清除自动匹配占用的目标，不能让自动推断覆盖人工选择。
            for old_name, old in list(result.items()):
                for old_sheet, old_target in list((old.get('sheet_mapping') or {}).items()):
                    if ((old.get('expected_file'), old_target) == destination or
                            (old_name == filename and old_sheet == sheet)):
                        del old['sheet_mapping'][old_sheet]
                        (old.get('header_mapping_by_sheet') or {}).pop(old_sheet, None)
                if not old.get('sheet_mapping'):
                    del result[old_name]
            entry = result.setdefault(filename, {'expected_file': target_file, 'sheet_mapping': {},
                'header_mapping': {}, 'header_mapping_by_sheet': {}, 'file_path': source['file_path']})
            if entry['expected_file'] != target_file:
                raise ValueError(f'同一上传文件不能映射到多个训练文件: {filename}')
            entry['sheet_mapping'][sheet] = target_sheet
            entry.setdefault('header_mapping_by_sheet', {})[sheet] = headers
            entry['confirmed'] = True
            entry['needs_rewrite'] = True
            # 未选择的源列不进入训练表，避免同名旧列覆盖选中的新列。
            entry.setdefault('selected_columns_by_sheet', {})[sheet] = list(headers)
    return result


def fully_unmatched_sheets(source_structure, unmatched_columns):
    """整张训练表的列均已选择无匹配时，不再强制要求该表出现在输入中。"""
    skipped = {tuple(item) for item in unmatched_columns or []}
    return {(filename, sheet) for filename, file_info in (source_structure or {}).get('files', {}).items()
            for sheet, info in file_info.get('sheets', {}).items()
            if info.get('headers') and all((filename, sheet, col) in skipped for col in info['headers'])}


def write_execution_sources(output_dir, file_mapping, source_data, source_structure, expected_structure=None):
    """从最终预加载数据生成脚本输入文件，不重读上传 Excel。"""
    import openpyxl
    from backend.utils.data_helpers import assign_sheet_keys
    from backend.utils.fast_header_matcher import FastHeaderMatcher
    from backend.utils.source_sheet_writer import write_source_dataframe

    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    pairs = {(Path(info['expected_file']).stem, sheet)
             for info in file_mapping.values() for sheet in info['sheet_mapping'].values()}
    keys = assign_sheet_keys(sorted(pairs | set(FastHeaderMatcher.expected_pairs(source_structure or {}))),
                             reserved_names=set((expected_structure or {}).get('sheets', {})))
    books = {}
    try:
        for info in file_mapping.values():
            filename = info['expected_file']
            if Path(filename).name != filename or '/' in filename or '\\' in filename:
                raise ValueError('映射目标必须是文件名')
            if filename not in books:
                wb = openpyxl.Workbook()
                wb.remove(wb.active)
                books[filename] = wb
            wb = books[filename]
            for sheet in info['sheet_mapping'].values():
                entry = source_data[keys[(Path(filename).stem, sheet)]]
                if sheet in wb.sheetnames:
                    raise ValueError(f'重复的执行表: {filename}/{sheet}')
                write_source_dataframe(wb.create_sheet(sheet), entry['df'],
                                       entry.get('column_schemas'), entry.get('column_formats'))
        for filename, wb in books.items():
            path = output / filename
            if path.suffix.lower() == '.xls':
                import aspose_init
                aspose_init.ensure_license()
                from Aspose.Cells import Workbook, SaveFormat
                staged = path.with_suffix('.staged.xlsx')
                wb.save(staged)
                native = Workbook(str(staged))
                try:
                    native.Save(str(path), SaveFormat.Excel97To2003)
                finally:
                    native.Dispose()
                    staged.unlink(missing_ok=True)
            else:
                wb.save(path)
    finally:
        for wb in books.values():
            wb.close()
    return str(output)
