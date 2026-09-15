"""Bind legacy fill rules to actual template headers, never to merged headings."""
import re
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter
from backend.utils.template_row_planner import _norm_key, _is_summary_key


def bind_template_fill(ws, expected_columns, helper_columns):
    def norm(value):
        return ''.join(str(value or '').split()).casefold()
    expected = {c['letter']: c['name'] for c in expected_columns}
    signatures = {norm(v) for v in expected.values()}
    candidates = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        headers = {}
        for cell in row:
            if cell.value is not None:
                headers.setdefault(norm(cell.value), []).append(cell.column)
        score = len(signatures & headers.keys())
        candidates.append((score, row[0].row, headers))
    candidates.sort(key=lambda entry: entry[0], reverse=True)
    if not candidates or candidates[0][0] < max(2, len(signatures) * 0.7):
        raise ValueError('目标模板列头无法可靠定位，不能按旧行号填充')
    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
        raise ValueError('目标模板存在多个相同列头区域，请指定填充区域')
    _, header_row, headers = candidates[0]
    mapping = {}
    for letter, label in expected.items():
        found = headers.get(norm(label), [])
        if len(found) != 1:
            raise ValueError(f'目标模板列 {label!r} 缺失或不唯一')
        mapping[letter] = found[0]
    key_col = mapping['A']
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        key = _norm_key(ws.cell(r, key_col).value)
        if not key or _is_summary_key(key):
            break
        rows.append(r)
    if not rows:
        raise ValueError('目标模板员工数据区为空，需先执行人员行清洗')
    # Append helper columns after all existing content/merges, never over insurance columns.
    next_col = max([c for (r, c), cell in ws._cells.items() if cell.value is not None] +
                   [m.max_col for m in ws.merged_cells.ranges] + [0]) + 1
    for letter, label in helper_columns.items():
        found = headers.get(norm(label), [])
        if len(found) > 1:
            raise ValueError(f'模板中间项 {label!r} 不唯一')
        mapping[letter] = found[0] if found else next_col
        if not found:
            ws.cell(header_row, next_col, label)
            next_col += 1
    return rows, mapping


def remap_local_formula(formula, mapping):
    """Only remap unqualified A1 references; quoted text/source-sheet ranges stay intact."""
    if not isinstance(formula, str) or not formula.startswith('='):
        return formula
    tokens = Tokenizer(formula).items
    for token in tokens:
        if token.type != 'OPERAND' or token.subtype != 'RANGE' or '!' in token.value:
            continue
        def replace(match):
            dollar, letter, row = match.groups()
            return dollar + get_column_letter(mapping.get(letter, column_index_from_string(letter))) + row
        token.value = re.sub(r'(?<![A-Za-z0-9_])(\$?)([A-Z]{1,3})(\$?\d+)(?![A-Za-z0-9_])', replace, token.value)
    return '=' + ''.join(t.value for t in tokens)
