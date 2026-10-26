"""Versioned template execution: retain the template and copy original sources.

Only new template scripts opt in. Legacy scripts keep their normalized-source
coordinates and compatibility repairs.
"""
import ast
import copy
import os
import re
import tempfile
from pathlib import Path


def is_template_pipeline_v2(code):
    if not code:
        return False
    for node in ast.parse(code).body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'TEMPLATE_PIPELINE_VERSION'
                                               for t in node.targets):
            return isinstance(node.value, ast.Constant) and node.value.value == 2
    return False


def is_source_pipeline_v2(code):
    """Formula scripts also opt into native source sheets, independently of templates."""
    if not code:
        return False
    return any(isinstance(n, ast.Assign)
               and any(isinstance(t, ast.Name) and t.id == 'SOURCE_PIPELINE_VERSION' for t in n.targets)
               and isinstance(n.value, ast.Constant) and n.value.value == 2
               for n in ast.parse(code).body)


def uses_original_template_sources(code):
    if is_template_pipeline_v2(code) or is_source_pipeline_v2(code):
        mode = 'copy'
        for node in ast.parse(code).body:
            if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'SOURCE_SHEET_MODE' for t in node.targets):
                mode = ast.literal_eval(node.value)
        return mode == 'copy'
    from backend.utils.multi_workbook import script_targets
    return any(uses_original_template_sources(target.get('code', '')) for target in script_targets(code).values())


def source_mode_from_rules(rules):
    modes = re.findall(r'^[ \t]*(?:SOURCE_SHEET_MODE|源表处理方式)[ \t]*[:：=][ \t]*([^\n]+)',
                       str(rules or ''), re.M)
    resolved = []
    for mode in modes:
        mode = mode.strip().strip('"\'`').lower()
        mode = {'复制': 'copy', '原样复制': 'copy', '重建': 'rebuild'}.get(mode, mode)
        if mode not in ('copy', 'rebuild'):
            raise ValueError('源表处理方式必须为 copy（原表复制）或 rebuild（清洗合并后重建）')
        resolved.append(mode)
    if len(set(resolved)) > 1:
        raise ValueError('当前目标有冲突的源表处理方式')
    return resolved[0] if resolved else 'copy'


def sync_source_mode(code, rules):
    if not (is_template_pipeline_v2(code) or is_source_pipeline_v2(code)) or not re.search(r'^[ \t]*(?:SOURCE_SHEET_MODE|源表处理方式)[ \t]*[:：=]', str(rules or ''), re.M):
        return code
    mode = source_mode_from_rules(rules)
    lines = code.splitlines(keepends=True)
    assignments = [n for n in ast.parse(code).body if isinstance(n, ast.Assign)
                   and any(isinstance(t, ast.Name) and t.id == 'SOURCE_SHEET_MODE' for t in n.targets)]
    for node in reversed(assignments):
        lines[node.lineno - 1:node.end_lineno] = ['SOURCE_SHEET_MODE = ' + repr(mode) + '\n']
    return ''.join(lines)


def region_layout(region, header_mapping=None):
    """Capture real coordinates and schemas without sampling/rebuilding cells."""
    columns = []
    schemas = getattr(region, 'column_schemas', None) or {}
    for name, letter in (region.head_data or {}).items():
        schema = dict(schemas.get(letter) or {})
        columns.append({'name': (header_mapping or {}).get(name, name), 'letter': letter,
                        **schema})
    header_end = getattr(region, 'head_row_end', None)
    start = getattr(region, 'data_row_start', None) or ((header_end + 1) if header_end else None)
    end = getattr(region, 'data_row_end', None)
    return {'header_row': getattr(region, 'head_row_start', None), 'header_row_end': header_end,
            'data_start_row': start, 'data_end_row': end if end and start and end >= start else None,
            'columns': columns, 'formulas': dict(getattr(region, 'formula', None) or {})}


def source_origin(path, sheet, regions, header_mapping=None):
    return {'file_path': str(Path(path).resolve()), 'sheet_name': sheet,
            'regions': [region_layout(r, header_mapping) for r in regions]}


def write_template_cell(ws, row, column, value, overwrite=False, number_format=None):
    """Default to filling blanks; explicit overwrite authorizes replacing values/formulas."""
    cell = ws.cell(row, column)
    if not overwrite and cell.value is not None and cell.value != '':
        return False
    cell.value = value
    if number_format is not None:
        cell.number_format = number_format
    return True


def clear_template_cells(ws, address, include_formulas=False):
    """Explicit clearing only, retaining styles, dimensions and default formulas."""
    from openpyxl.utils.cell import range_boundaries
    first_col, first_row, last_col, last_row = range_boundaries(address)
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        for cell in row:
            if cell.data_type != 'f' or include_formulas:
                cell.value = None


def plan_original_sources(source_data, source_map, key_to_sheet, input_folder, reserved_names=()):
    """Resolve each logical source to its original file and actual worksheet.

    Matched upload names may differ from training names; provenance wins over
    training literals. Runtime layouts replace the training coordinate hints.
    """
    from backend.utils.data_helpers import build_prefixed_sheet_names
    plan, runtime_map, used = [], {}, set(reserved_names)
    seen = set()
    for key, entry in (source_data or {}).items():
        if id(entry) in seen:
            continue
        seen.add(id(entry))
        name = key_to_sheet.get(key)
        if name is None:
            name = build_prefixed_sheet_names([key], reserved=used)[key]
            key_to_sheet[key] = name
        if name in used:
            raise ValueError(f'原始源表名称与模板或另一源表冲突: {name}')
        used.add(name)
        layout = copy.deepcopy(source_map.get(name) or {})
        origins = entry.get('source_origins') or []
        if len(origins) > 1:
            raise ValueError(f'{key} 合并了多个原始工作表，不能按一张原表复制；请明确使用 SOURCE_SHEET_MODE: rebuild')
        if origins:
            origin = copy.deepcopy(origins[0])
        else:
            filename = entry.get('source_file') or layout.get('source_file')
            sheet = entry.get('source_sheet') or layout.get('source_sheet')
            if not filename or not sheet:
                raise ValueError(f'{key} 缺少原始文件/工作表信息，请重新加载源文件')
            matches = [p for p in Path(input_folder).iterdir() if p.is_file()
                       and p.suffix.lower() in ('.xlsx', '.xlsm', '.xls')
                       and (p.name.casefold() == str(filename).casefold() or p.stem.casefold() == str(filename).casefold())]
            if len(matches) != 1:
                raise ValueError(f'{key} 无法唯一定位原始源文件: {filename}')
            origin = {'file_path': str(matches[0].resolve()), 'sheet_name': sheet,
                      'regions': layout.get('regions') or [layout]}
        if not Path(origin['file_path']).is_file():
            raise ValueError(f'原始源文件不可用: {origin["file_path"]}')
        regions = origin.get('regions') or []
        if not regions or not regions[0].get('columns') or not regions[0].get('data_start_row'):
            raise ValueError(f'{key} 缺少真实行列位置，请重新解析源文件后执行')
        layout.update(copy.deepcopy(regions[0]))
        layout.update(regions=regions, source_file=entry.get('source_file') or layout.get('source_file'),
                      source_sheet=origin['sheet_name'])
        runtime_map[name] = layout
        plan.append({**origin, 'destination': name})
    return plan, runtime_map


def rebuilt_source_layouts(source_data, key_to_sheet):
    from openpyxl.utils import get_column_letter
    result = {}
    for key, entry in source_data.items():
        df = entry.get('df')
        if df is None:
            continue
        name = key_to_sheet[key]
        schemas = entry.get('column_schemas') or {}
        result[name] = {'header_row': 1, 'data_start_row': 2, 'data_end_row': len(df) + 1,
                        'columns': [{'name': column, 'letter': get_column_letter(i + 1),
                                     **(schemas.get(column) or {})} for i, column in enumerate(df.columns)]}
    return result


def prepare_formula_source_sheets(wb, source_data, input_folder):
    """Expose data and real coordinates; reserve empty sheets for native copying."""
    from backend.utils.data_helpers import make_unique_sheet_key
    reserved = set(wb.sheetnames) | {'参数', '历史数据'}
    used, keys, primary = set(reserved), {}, {}
    for key, entry in source_data.items():
        if id(entry) not in primary:
            primary[id(entry)] = make_unique_sheet_key(key, used)
        keys[key] = primary[id(entry)]
    plan, layouts = plan_original_sources(source_data, {}, keys, input_folder, reserved)
    sources, by_entry = {}, {}
    for key, entry in source_data.items():
        identity = id(entry)
        if identity in by_entry:
            sources[key] = by_entry[identity]
            continue
        layout = layouts[keys[key]]
        item = {**entry, 'ws': wb.create_sheet(keys[key]), 'layout': layout,
                'column_letters': {col['name']: col['letter'] for col in layout['columns']},
                'data_start_row': layout['data_start_row'], 'data_end_row': layout.get('data_end_row')}
        sources[key] = by_entry[identity] = item
    for spec in plan:
        spec['placeholder'] = True
    return sources, plan


def _sheet_references(expression):
    """Chart series are not formula cells, but their source sheets matter too."""
    from openpyxl.formula import Tokenizer
    expression = str(expression or '')
    if not expression:
        return []
    names = []
    for token in Tokenizer(expression if expression.startswith('=') else '=' + expression).items:
        if token.type != 'OPERAND' or token.subtype != 'RANGE' or '!' not in token.value:
            continue
        sheet = token.value.rsplit('!', 1)[0]
        if '[' in sheet or ']' in sheet:
            continue  # an original external chart dependency remains external
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        names.append(sheet)
    return names


def copy_original_source_sheets(output_path, plan, workbook_factory=None):
    """Native whole-worksheet copy AFTER openpyxl saves the target.

    Sources never pass through openpyxl load/save or DataFrame row writing.
    Include same-file formula dependencies and precreate destinations, so the
    native copier resolves references inside the result instead of new links.
    """
    if not plan:
        return []
    import aspose_init
    aspose_init.ensure_license()
    from Aspose.Cells import Workbook, CopyOptions, CellValueType
    from backend.utils.data_helpers import build_prefixed_sheet_names
    from backend.utils.output_postprocess import _silent_load_options
    def open_native(path):
        options = _silent_load_options()
        return Workbook(path, options) if options is not None else Workbook(path)
    factory = workbook_factory or open_native
    books, output, staged = {}, None, None
    try:
        output = factory(str(output_path))
        reserved = {str(output.Worksheets[i].Name) for i in range(output.Worksheets.Count)}
        groups = {}
        for spec in plan:
            path = str(Path(spec['file_path']).resolve())
            groups.setdefault(path, []).append(spec)
        operations = []
        for path, specs in groups.items():
            source = books[path] = factory(path)
            worksheets = {str(source.Worksheets[i].Name): source.Worksheets[i]
                          for i in range(source.Worksheets.Count)}
            names = {}
            for spec in specs:
                sheet, destination = spec['sheet_name'], spec['destination']
                if sheet not in worksheets:
                    raise ValueError(f'原始源工作表不存在: {Path(path).name}/{sheet}')
                placeholder = (spec.get('placeholder') and destination in reserved
                               and output.Worksheets[destination].Cells.MaxDataRow < 0)
                if (destination in reserved and not placeholder) or sheet in names:
                    raise ValueError(f'源工作表重复或名称冲突: {destination}')
                reserved.add(destination)
                names[sheet] = destination
            pending = list(names)
            while pending:
                original = pending.pop()
                dependencies = set()
                iterator = worksheets[original].Cells.GetEnumerator()
                while iterator.MoveNext():
                    cell = iterator.Current
                    if not cell.IsFormula:
                        continue
                    for area in cell.GetPrecedents():
                        dependency = str(area.SheetName or '')
                        if area.IsExternalLink or not dependency or dependency == original:
                            continue  # original external links remain original links
                        dependencies.add(dependency)
                charts = worksheets[original].Charts
                for ci in range(charts.Count):
                    series = charts[ci].NSeries
                    for si in range(series.Count):
                        for expression in (series[si].Values, series[si].XValues, series[si].Name):
                            dependencies.update(_sheet_references(expression))
                for dependency in sorted(dependencies):
                    if dependency not in worksheets:
                        raise ValueError(f'源公式依赖无法定位: {Path(path).name}/{dependency}')
                    if dependency in names:
                        continue
                    key = Path(path).stem + '_' + dependency
                    destination = build_prefixed_sheet_names([key], reserved=reserved)[key]
                    reserved.add(destination)
                    names[dependency] = destination
                    pending.append(dependency)
            # Rename in the source's in-memory copy; native reference updates
            # preserve ranges, named references and cross-sheet formulas.
            import uuid
            for ws in worksheets.values():
                ws.Name = 'tmp_' + uuid.uuid4().hex[:20]
            for original, destination in names.items():
                worksheets[original].Name = destination
                operations.append((worksheets[original], destination))
        for _, destination in operations:
            if output.Worksheets[destination] is None:
                output.Worksheets.Add(destination)
        for source, destination in operations:
            # CopyOptions retains source-workbook string/style index mappings.
            # Reusing it across files maps later sources to the first file's
            # strings and styles (including amounts turning into dates).
            # Each worksheet copy must start with its own mapping context.
            options = CopyOptions()
            options.CopyNames = True
            options.ReferToSheetWithSameName = True
            options.ReferToDestinationSheet = True
            copied = output.Worksheets[destination]
            copied.Copy(source, options)
            if source.Workbook.Settings.Date1904 != output.Settings.Date1904:
                # One result has one date system. Native worksheet copy retains
                # serial numbers, so adjust date constants to retain their dates.
                iterator = source.Cells.GetEnumerator()
                while iterator.MoveNext():
                    cell = iterator.Current
                    if not cell.IsFormula and cell.Type == CellValueType.IsDateTime:
                        copied.Cells[cell.Row, cell.Column].PutValue(cell.DateTimeValue)
        fd, staged = tempfile.mkstemp(prefix='.sources_', suffix=Path(output_path).suffix,
                                      dir=Path(output_path).resolve().parent)
        os.close(fd)
        output.Save(staged)
        output.Dispose()
        output = None
        os.replace(staged, output_path)
        staged = None
        return [destination for _, destination in operations]
    finally:
        if output is not None:
            output.Dispose()
        for book in books.values():
            book.Dispose()
        if staged and os.path.exists(staged):
            os.remove(staged)
