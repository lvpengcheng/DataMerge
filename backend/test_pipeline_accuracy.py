"""规则证据、日期、多表关联及生成脚本的离线回归，不调用 AI 服务。"""
import ast
from datetime import datetime
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
import pytest
from openpyxl.utils.datetime import to_excel, MAC_EPOCH

from backend.utils.source_sheet_writer import coerce_source_date, write_source_dataframe
from backend.utils.table_merge import merge_source_tables
from backend.utils.formula_evidence import collect_formula_evidence
from backend.utils.data_helpers import assign_sheet_keys, make_unique_sheet_key
from backend.ai_engine.rule_organizer import RuleOrganizer
from backend.ai_engine.document_parser import DocumentParser
from backend.ai_engine.table_analyzer import TableAnalyzer
from backend.ai_engine.prompt_generator import PromptGenerator


@pytest.mark.parametrize('value', [datetime(2026, 9, 6), '2026-09-06', '2026年9月6日',
                                  20260906, '20260906', to_excel(datetime(2026, 9, 6)),
                                  str(to_excel(datetime(2026, 9, 6))), np.int64(46271)])
def test_excel_dates_have_no_unix_epoch_conversion(value):
    assert coerce_source_date(value) == datetime(2026, 9, 6)


@pytest.mark.parametrize('value', [None, pd.NaT, pd.NA, float('nan'), '', '  ', 0])
def test_blank_date_stays_blank(value):
    assert coerce_source_date(value) is None


def test_date_epoch_invalid_and_boolean():
    assert coerce_source_date(1) == datetime(1900, 1, 1)
    assert coerce_source_date(1, epoch=MAC_EPOCH) == datetime(1904, 1, 2)
    assert coerce_source_date('未入职') == '未入职'
    assert coerce_source_date(20260230) == 20260230
    assert coerce_source_date(True) is True


def entry(df, **metadata):
    return {'df': df, 'columns': list(df.columns), **metadata}


def test_multi_key_join_preserves_left_order_blanks_and_schemas():
    left = pd.DataFrame({'工号': [1, '001', None, 1], '月份': [8, 8, 8, 9]}, index=[5, 1, 7, 3])
    right = pd.DataFrame({'工号': ['1', '001', None, '1'], '月份': [8, 8, 8, 9], '金额': [10, 20, 999, 30]})
    schema = {'金额': {'field_type': 'decimal', 'number_format': '0.00'}}
    source = {'主表': entry(left), '明细': entry(right, column_schemas=schema)}
    result = merge_source_tables(source, {'horizontal_joins': [{'left': '主表', 'right': '明细', 'on': ['工号', '月份']}]})['主表']
    assert result['df'].index.tolist() == [5, 1, 7, 3]
    assert result['df']['金额'].iloc[[0, 1, 3]].tolist() == [10, 20, 30]
    assert pd.isna(result['df']['金额'].iloc[2])
    assert result['column_schemas'] == schema
    assert list(left.columns) == ['工号', '月份']
    assert result['df']['工号'].iloc[1] == '001'


def test_duplicate_lookup_key_is_not_silently_expanded():
    sources = {'A': entry(pd.DataFrame({'工号': [1]})),
               'B': entry(pd.DataFrame({'工号': [1, '1'], '工资': [100, 200]}))}
    with pytest.raises(ValueError, match='存在重复'):
        merge_source_tables(sources, {'horizontal_joins': [{'left': 'A', 'right': 'B', 'on': '工号'}]})


def test_vertical_merge_keeps_union_and_metadata():
    sources = {'A': entry(pd.DataFrame({'工号': ['001']}), column_schemas={'工号': {'field_type': 'text'}}),
               'B': entry(pd.DataFrame({'工号': ['002'], '金额': [2.5]}), column_formats={'金额': '0.00'})}
    combined = merge_source_tables(sources, {'vertical_groups': [['A', 'B']]})['derived_main_table']
    assert combined['columns'] == ['工号', '金额']
    assert combined['column_schemas']['工号']['field_type'] == 'text'
    assert combined['column_formats']['金额'] == '0.00'


def test_explicit_main_table_overrides_column_coverage():
    tables = {'人员名单': entry(pd.DataFrame({'工号': ['001']})),
              '金额明细': entry(pd.DataFrame({'工号': ['001'], '工资': [1], '奖金': [2], '扣款': [3]}))}
    analyzer = TableAnalyzer()
    result = analyzer._determine_main_table(tables, ['工号', '工资', '奖金', '扣款'], 1, '工号', '### 主表: 人员名单')
    assert result[:2] == ('single', ['人员名单'])
    assert analyzer._fuzzy_match_sheet('表', {'表备份': {}, '表': {}}) == '表'
    assert analyzer._fuzzy_match_sheet('表', {'表备份': {}, '表历史': {}}) is None
    assert not analyzer._sheet_has_column({'columns': ['原工号']}, '工号')


def test_key_heuristic_prefers_shared_identifier_over_name():
    tables = {'名单': entry(pd.DataFrame({'姓名': ['甲', '乙'], '工号': ['001', '002']})),
              '明细': entry(pd.DataFrame({'工号': ['001', '001'], '金额': [100, 200]}))}
    assert TableAnalyzer()._detect_primary_key(tables, ['姓名', '工号', '金额'], '')[0] == '工号'
    assert TableAnalyzer()._detect_primary_key({'A': entry(pd.DataFrame({'金额': [1, 2]}))}, ['金额'], '')[0] == ''


def test_sheet_keys_do_not_collide_with_result_names_case_or_invalid_chars():
    keys = assign_sheet_keys([('原文件', '结果'), ('X', '数据')], reserved_names={'结果', '原文件_结果'})
    assert keys[('原文件', '结果')] == '原文件_结果_2'
    used = {'DATA'}
    assert make_unique_sheet_key('data', used) == 'data_2'
    assert make_unique_sheet_key('A/B:C', used) == 'A_B_C'


def test_source_roundtrip_preserves_dates_keys_and_literal_equals(tmp_path):
    df = pd.DataFrame({'工号': ['001', 1002.0], '入职日期': [46271, None], '备注': ['=1+1', '正常'], '金额': [1.25, 3]})
    schemas = {'工号': {'field_type': 'text', 'number_format': '@'},
               '入职日期': {'field_type': 'date', 'number_format': 'yyyy-mm-dd'},
               '金额': {'field_type': 'decimal', 'number_format': '0.00'}}
    wb = openpyxl.Workbook()
    write_source_dataframe(wb.active, df, schemas)
    file = tmp_path / 'source.xlsx'
    wb.save(file)
    out = openpyxl.load_workbook(file)
    ws = out.active
    assert ws['A2'].value == '001'
    assert ws['A3'].value == '1002'
    assert ws['B2'].value == datetime(2026, 9, 6)
    assert ws['B3'].value is None
    assert ws['C2'].value == '=1+1' and ws['C2'].data_type == 's'
    assert ws['D2'].number_format == '0.00'
    out.close()


def formula_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '结果'
    ws.append(['工号', '金额', '税后'])
    for row in range(2, 18):
        ws.append([f'{row:03}', row, f'=B{row}*0.9'])
    ws['C18'] = '=IF(B18>100,B18*0.8,B18*0.9)'
    ws['D18'] = '=IF(A18="A18",2,3)'
    ws['D19'] = '=IF(A19="A19",2,3)'
    wb.save(path)


def test_full_formula_scan_reads_uncached_late_formulas(tmp_path):
    path = tmp_path / 'formulas.xlsx'
    formula_book(path)
    evidence = collect_formula_evidence(path)['结果']
    assert evidence['formula_count'] == 19
    assert len(evidence['formulas']) == 4
    assert evidence['formulas']['C18'].startswith('=IF(')
    assert 'D18' in evidence['formulas'] and 'D19' in evidence['formulas']


def test_document_excel_keeps_coordinates_merges_and_formulas(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells('A1:B1')
    ws['A1'] = '第一章'
    ws['B3'] = '=A3*2'
    wb.create_sheet('末尾规则')['A1'] = '最后一步'
    path = tmp_path / 'sop.xlsx'
    wb.save(path)
    text = DocumentParser().parse_document(str(path))
    assert 'A1:B1' in text and 'B3: =A3*2' in text and '最后一步' in text


def test_word_keeps_table_between_surrounding_paragraphs(tmp_path):
    import docx
    doc = docx.Document()
    doc.add_heading('薪资步骤', level=1)
    doc.add_paragraph('计算前')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '公式所在表格'
    doc.add_paragraph('计算后')
    path = tmp_path / 'sop.docx'
    doc.save(path)
    text = DocumentParser().parse_document(str(path))
    assert text.index('计算前') < text.index('公式所在表格') < text.index('计算后')
    assert '# 薪资步骤' in text


def test_long_sop_reads_every_segment_with_outline(monkeypatch):
    import backend.ai_engine.rule_organizer as module
    monkeypatch.setattr(module, 'MAX_DOC_CHARS', 1000)
    calls = []
    provider = SimpleNamespace(chat=lambda messages: calls.append(messages) or '保留规则条款')
    organizer = RuleOrganizer.__new__(RuleOrganizer)
    organizer.ai_provider = provider
    content = '# 总则\n' + '边界说明\n' * 350 + '\n# 尾部例外\n唯一末尾规则'
    result = organizer._prepare_design_context(content)
    assert len(calls) > 1
    assert all('# 尾部例外' in c[1]['content'] for c in calls)
    assert '唯一末尾规则' in calls[-1][1]['content']
    assert f', {len(content)})' in result


def test_training_prompt_retains_evidence_even_over_recommended_size():
    structure = {'sheets': {'结果': {'headers': {'工资': 'A'}, 'formulas': {'A2': '=B2*2'},
                                  'data_sample': [{'A': 123}], 'column_schemas': {'工资': {'field_type': 'decimal'}}}}}
    result = PromptGenerator()._compress_structure(structure, max_length=20)
    assert '=B2*2' in result and '123' in result and 'decimal' in result


def function_namespace(code, extra=None):
    """只执行生成代码的函数定义，避免 main 触及业务文件。"""
    tree = ast.parse(code)
    namespace = {'pd': pd, 'datetime': datetime, 'PatternFill': openpyxl.styles.PatternFill,
                 'Font': openpyxl.styles.Font, '_SK_TO_SHEET': {}, 'SOURCE_PREFIX': '源_'}
    from backend.utils.data_helpers import build_prefixed_sheet_names
    namespace['build_prefixed_sheet_names'] = build_prefixed_sheet_names
    namespace.update(extra or {})
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            exec(compile(ast.Module(body=[node], type_ignores=[]), '<generated>', 'exec'), namespace)
    return namespace


def test_generated_formula_and_template_scripts_use_runtime_helpers(tmp_path):
    from backend.ai_engine.formula_code_generator import FormulaCodeGenerator
    from backend.ai_engine.template_code_generator import TemplateCodeGenerator
    formula = FormulaCodeGenerator(ai_provider=object())._build_complete_code('def fill_result_sheets(*args):\n    pass')
    compiled = function_namespace(formula)
    assert compiled['_to_native_datetime'](46271) == datetime(2026, 9, 6)
    data = {'A': entry(pd.DataFrame({'工号': ['001'], '入职日期': [46271], '备注': ['=1+1']}))}
    wb = openpyxl.Workbook()
    result = compiled['write_source_sheets'](wb, data)
    assert result['A']['ws']['B2'].value == datetime(2026, 9, 6)
    assert result['A']['ws']['C2'].data_type == 's'
    template = TemplateCodeGenerator(ai_provider=object())._build_complete_code(
        'def fill_template(*args):\n    pass', str(tmp_path / 'template.xlsx'), str(tmp_path),
        {'sheets': {}}, {}, False)
    compiled = function_namespace(template)
    wb = openpyxl.Workbook()
    compiled['_append_source_sheets'](wb, data)
    assert wb['源_A']['C2'].data_type == 's'


def test_bulk_parser_sample_exports_only_needed_rows():
    from excel_parser import IntelligentExcelParser, ExcelRegion
    parser = IntelligentExcelParser()
    calls = []
    class Cells:
        def ExportArray(self, start, col, rows, columns):
            calls.append((start, rows))
            return np.array([[start + i + 1, 123] for i in range(rows)], dtype=object)
        def __getitem__(self, item):
            raise RuntimeError('schema sampling omitted in this controlled test')
    region = ExcelRegion()
    region.data_row_start, region.data_row_end = 2, 100000
    region.head_data = {'工号': 'A', '金额': 'B'}
    parser._is_title_row_from_values = lambda *_: False
    parser._is_summary_row_from_values = lambda *_: False
    parser._collect_data_bulk(SimpleNamespace(_ws=SimpleNamespace(Cells=Cells())), region, 2, max_data_rows=3)
    assert len(region.data) == 3
    assert calls == [(1, 3)]


def test_target_schema_postprocessing_is_idempotent(tmp_path):
    from backend.utils.output_postprocess import apply_expected_column_schemas, _canon_key_text
    assert _canon_key_text(123456789012345678) == '123456789012345678'
    wb = openpyxl.Workbook()
    wb.active.title = '结果'
    wb.active.append(['工号', '入职日期'])
    wb.active.append([123, 46271])
    path = tmp_path / 'result.xlsx'
    wb.save(path)
    structure = {'sheets': {'结果': {'headers': {'工号': 'A', '入职日期': 'B'}, 'column_schemas': {
        '工号': {'field_type': 'text', 'number_format': '@'},
        '入职日期': {'field_type': 'date', 'number_format': 'yyyy-mm-dd'}}}}}
    assert apply_expected_column_schemas(path, structure) > 0
    before = path.read_bytes()
    assert apply_expected_column_schemas(path, structure) == 0
    assert path.read_bytes() == before


def test_key_normalization_preserves_cached_formula(tmp_path):
    from backend.utils.output_postprocess import normalize_key_columns_to_text, _open_workbook
    wb = openpyxl.Workbook()
    wb.active.append(['工号'])
    wb.active.append(['=ROW()'])
    wb.active.append([123])
    path = tmp_path / 'cached.xlsx'
    wb.save(path)
    native = _open_workbook(str(path))
    try:
        native.CalculateFormula()
        native.Save(str(path))
    finally:
        native.Dispose()
    assert normalize_key_columns_to_text(path) == 1
    assert normalize_key_columns_to_text(path) == 0
    out = openpyxl.load_workbook(path)
    assert out.active['A2'].value == '=ROW()'
    assert out.active['A3'].value == '123'
    out.close()


def test_generated_script_cleans_then_joins_and_recalculates(tmp_path):
    from backend.ai_engine.formula_code_generator import FormulaCodeGenerator
    from backend.utils.output_postprocess import _open_workbook
    generator = FormulaCodeGenerator(ai_provider=object())
    code = generator._build_complete_code('''
def clean_source_data(source_data):
    data = source_data['明细']
    data['df'] = data['df'].groupby('工号', as_index=False)['金额'].sum()
    return source_data

def fill_result_sheets(wb, source_sheets, *args):
    ws = wb.create_sheet('结果')
    ws.append(['工号', '入职日期', '金额', '实发'])
    for r, record in enumerate(source_sheets['主表']['df'].itertuples(index=False, name=None), 2):
        ws.append(list(record) + [f'=C{r}*0.9'])
''', merge_config={'horizontal_joins': [{'left': '主表', 'right': '明细', 'on': '工号'}]})
    source = {'主表': entry(pd.DataFrame({'工号': ['001', '002'], '入职日期': [46271, None]})),
              '明细': entry(pd.DataFrame({'工号': ['001', '001', '002'], '金额': [100, 50, 200]}))}
    namespace = {'__name__': 'offline_pipeline', '_pre_loaded_source_data': source,
                 'output_folder': str(tmp_path), 'salary_year': 2026, 'salary_month': 9}
    exec(compile(code, '<generated>', 'exec'), namespace)
    assert namespace['main']()
    file = tmp_path / '薪资汇总表.xlsx'
    wb = _open_workbook(str(file))
    try:
        wb.CalculateFormula()
        wb.Save(str(file))
    finally:
        wb.Dispose()
    out = openpyxl.load_workbook(file, data_only=True)
    ws = out['结果']
    assert ws['A2'].value == '001'
    assert ws['B2'].value == datetime(2026, 9, 6)
    assert ws['C2'].value == 150 and ws['D2'].value == 135
    assert ws['D3'].value == 180
    assert ws.max_row == 3
    out.close()


def test_rule_generator_does_not_reduce_document_to_title():
    from backend.ai_engine.rule_generator import AIRuleGenerator
    generator = AIRuleGenerator(ai_provider=object())
    content = '# 总则\n' + '条款\n' * 500 + '\n不常见尾部业务规则'
    prompt = generator._create_rule_generation_prompt(content, {'files': {}}, {'sheets': {}})
    assert content in prompt


def test_rule_organizer_retains_full_doc_before_analysis(tmp_path):
    path = tmp_path / 'long.txt'
    content = '规则\n' * 11000 + '独有尾部步骤'
    path.write_text(content, encoding='utf-8')
    organizer = RuleOrganizer.__new__(RuleOrganizer)
    organizer.doc_parser = DocumentParser()
    assert content in organizer._extract_design_docs([str(path)])
