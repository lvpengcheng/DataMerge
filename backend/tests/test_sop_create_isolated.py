"""SOP 新建字段与数据归属校验，使用独立内存数据库。"""
import ast
import asyncio
import io
import pandas as pd
from datetime import datetime
from pathlib import Path
import re
import tempfile
import shutil
import uuid
from types import SimpleNamespace
from typing import Optional
import unittest

from sqlalchemy import Boolean, Column, DateTime, Integer, String, Index, JSON, ForeignKey, create_engine
from sqlalchemy.orm import declarative_base, sessionmaker, relationship
from starlette.datastructures import UploadFile
from fastapi.responses import FileResponse
from backend.utils.upload_stream import save_upload_file


Base = declarative_base()


class Entry(Base):
    __tablename__ = "entries"
    id = Column(Integer, primary_key=True)
    customer_name = Column(String)
    __table_args__ = (Index("uq_test_customer", "customer_name", unique=True),)
    batch_month = Column(String)
    description = Column(String)
    status = Column(String)
    created_by = Column(Integer)
    assignee_user_id = Column(Integer, ForeignKey("users.id"))
    assignee = relationship("User")
    supplementary_files = Column(JSON(none_as_null=True))
    rounds = relationship("Round")
    created_at = Column(DateTime, default=datetime.now)
    salesperson_name = Column(String)
    salesperson_email = Column(String)
    updated_at = Column(DateTime, default=datetime.now)


class User(Base):
    __tablename__ = "users"
    id = Column(Integer, primary_key=True)
    username = Column(String)
    display_name = Column(String)
    email = Column(String)
    is_active = Column(Boolean)
    can_upload = Column(Boolean)


class Round(Base):
    __tablename__ = "rounds"
    id = Column(Integer, primary_key=True)
    sop_id = Column(Integer, ForeignKey("entries.id"))
    round_no = Column(Integer)
    status = Column(String)
    review_comment = Column(String)
    ai_comment = Column(String)
    reviewed_at = Column(DateTime)
    created_at = Column(DateTime, default=datetime.now)


class HTTPException(Exception):
    def __init__(self, status_code, detail):
        self.status_code, self.detail = status_code, detail


class SopCreateTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        self.addCleanup(self.engine.dispose)
        Base.metadata.create_all(self.engine)
        self.session = sessionmaker(bind=self.engine)
        with self.session() as db:
            db.add(User(id=7, username="seller", display_name="业务员", email="seller@example.com", is_active=True, can_upload=True))
            db.commit()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        (Path(self.temp.name) / "客户甲").mkdir()
        tree = ast.parse(Path(__file__).parent.parent.joinpath("api/tools.py").read_text(encoding="utf-8"))
        nodes = [node for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
                 and node.name in {"sop_entries_create", "_sop_customer_choices", "sop_entries_list",
                                   "_sop_status_groups", "_sop_business_status", "_sop_filter_status", "_sop_status_label",
                                   "_sop_check_customer_name", "sop_import_assignments",
                                   "sop_report", "sop_supplementary_upload", "sop_supplementary_download",
                                   "_sop_check_upload", "_sop_check_read", "_sop_is_owner", "_sop_can_read"}]
        for node in nodes:
            node.decorator_list = []
        self.ns = dict(SopCreateRequest=object, Depends=lambda value: None, require_permission=lambda key: None,
                       SessionLocal=self.session, SopEntry=Entry, User=User, HTTPException=HTTPException,
                       SOP_ROOT=Path(self.temp.name) / "__tools_sop__", re=re, datetime=datetime,
                       has_permission=lambda user, key: user.can_upload,
                       _sop_entry_to_out=lambda entry, db: {"id":entry.id},
                       Optional=Optional, Query=lambda value: value,
                       UploadFile=object, File=lambda value: value, Form=lambda value: value, pd=pd, io=io,
                       _sop_is_superadmin=lambda user: user.role == "admin",
                       List=list, Path=Path, _require_sop_action=lambda key: None,
                       uuid=uuid, shutil=shutil, save_upload_file=save_upload_file, FileResponse=FileResponse,
                       PROJECT_ROOT=Path(self.temp.name), StreamingResponse=lambda stream, **kwargs: stream)
        exec(compile(ast.Module(body=nodes, type_ignores=[]), "sop_create", "exec"), self.ns)

    def create(self, **overrides):
        values = dict(customer_name="客户甲", batch_month="202610", assignee_user_id=7, description="说明")
        values.update(overrides)
        return asyncio.run(self.ns["sop_entries_create"](SimpleNamespace(**values), SimpleNamespace(id=1)))

    def test_four_status_filters_counts_and_reviewer_scope(self):
        states = ["draft", "awaiting_files", "ai_queued", "ai_analyzing", "ai_passed", "ai_failed", "rejected", "failed", "completed"]
        with self.session() as db:
            db.add_all(Entry(customer_name=f"客户甲{i}", batch_month="202610", status=state, assignee_user_id=7) for i, state in enumerate(states))
            db.add(Entry(customer_name="其他客户", batch_month="202611", status="draft", assignee_user_id=8))
            db.commit()
        self.ns["has_permission"] = lambda user, key: key in user.permissions
        reviewer = SimpleNamespace(id=3, role="reviewer", permissions={"sop.review"})
        admin = SimpleNamespace(id=1, role="admin", permissions=set())
        seller = SimpleNamespace(id=7, role="seller", permissions={"sop.upload"})
        def listing(user, status=None, month="202610"):
            return asyncio.run(self.ns["sop_entries_list"](keyword=None, status=status, batch_month=month, current_user=user))
        self.assertEqual(listing(reviewer), listing(admin))
        expected = {"pending_upload":2, "pending_review":3, "rejected":3, "completed":1}
        for status, count in expected.items():
            data = listing(reviewer, status)
            self.assertEqual(len(data["items"]), count)
            self.assertEqual(data["counts"], {"total":9, **expected})
        self.assertEqual(len(listing(seller, month=None)["items"]), 9)
        self.assertEqual(len(listing(reviewer, month=None)["items"]), 10)
        with self.assertRaises(HTTPException):
            listing(reviewer, "ai_passed")
        for status in states:
            self.assertIn(self.ns["_sop_status_label"](status), {"待上传", "待审核", "驳回", "通过"})

    def test_creation_persists_selected_owner_month_and_email(self):
        self.create()
        with self.session() as db:
            entry = db.query(Entry).one()
            self.assertEqual((entry.assignee_user_id, entry.batch_month, entry.status), (7, "202610", "draft"))
            self.assertEqual((entry.salesperson_name, entry.salesperson_email), ("业务员", "seller@example.com"))
        with self.assertRaises(HTTPException) as duplicate:
            self.create()
        self.assertEqual(duplicate.exception.status_code, 409)

    def test_salesperson_filter_applies_to_list_and_report(self):
        self.create()
        self.ns["has_permission"] = lambda user, key: key in user.permissions
        reviewer = SimpleNamespace(id=3, role="reviewer", permissions={"sop.review"})
        for word, expected in [("业务", 1), ("seller", 1), ("不存在", 0), ("%", 0)]:
            data = asyncio.run(self.ns["sop_entries_list"](salesperson=word, current_user=reviewer))
            self.assertEqual(len(data["items"]), expected)
            self.assertEqual(data["counts"]["total"], expected)
            from openpyxl import load_workbook
            book = load_workbook(asyncio.run(self.ns["sop_report"](salesperson=word)))
            self.assertEqual(book.active.max_row, expected + 1)
            book.close()

    def test_report_headers_and_last_rejection_survive_later_approval_and_template(self):
        self.create()
        with self.session() as db:
            entry = db.query(Entry).one()
            entry.status = "completed"
            db.add_all([Round(sop_id=entry.id, round_no=1, status="ai_failed", ai_comment="旧原因"),
                        Round(sop_id=entry.id, round_no=2, status="rejected", review_comment="最近原因"),
                        Round(sop_id=entry.id, round_no=3, status="completed")])
            db.commit()
        from openpyxl import Workbook, load_workbook
        for with_template in (False, True):
            if with_template:
                path = Path(self.temp.name) / "参考" / "图片和附件" / "合杰开户流程SOP_v3.xlsx"
                path.parent.mkdir(parents=True)
                template = Workbook()
                template.active.append(["旧标题"] * 12 + ["项目编号"])
                template.active.append(["旧数据"] * 13)
                template.save(path)
                template.close()
            book = load_workbook(asyncio.run(self.ns["sop_report"]()))
            rows = list(book.active.values)
            data = dict(zip(rows[0], rows[1]))
            self.assertEqual(len(rows), 2)
            self.assertNotIn("项目编号", data)
            self.assertEqual(len(rows[0]), 12)
            self.assertEqual(data["集约月份批次"], "202610")
            self.assertEqual(data["最近驳回原因"], "最近原因")
            self.assertEqual(data["当前状态"], "通过")
            self.assertEqual(data["驳回次数"], 2)
            book.close()

    def test_supplementary_upload_replaces_appends_and_remains_authorized_and_independent(self):
        self.create()
        self.ns["has_permission"] = lambda user, key: key in user.permissions
        seller = SimpleNamespace(id=7, role="seller", permissions={"sop.upload"})
        stranger = SimpleNamespace(id=8, role="seller", permissions={"sop.upload"})
        reviewer = SimpleNamespace(id=3, role="reviewer", permissions={"sop.review"})
        with self.session() as db:
            entry = db.query(Entry).one()
            entry.status = "completed"
            identifier = entry.id
            db.commit()
        def upload(who, name="注意事项.txt", content=b"supplement"):
            file = UploadFile(io.BytesIO(content), filename=name)
            try:
                return asyncio.run(self.ns["sop_supplementary_upload"](identifier, [file], who))
            finally:
                file.file.close()
        with self.assertRaises(HTTPException) as denied:
            upload(stranger)
        self.assertEqual(denied.exception.status_code, 403)
        with self.assertRaises(HTTPException):
            upload(seller, "bad.exe")
        self.assertTrue(upload(seller)["ok"])
        original = asyncio.run(self.ns["sop_supplementary_download"](identifier, 0, reviewer)).path
        self.assertTrue(upload(seller, content=b"replacement")["ok"])
        self.assertFalse(Path(original).exists())
        result = upload(seller, "历史数据.xlsx", b"history")
        self.assertEqual([file["name"] for file in result["files"]], ["注意事项.txt", "历史数据.xlsx"])
        with self.assertRaises(HTTPException):
            upload(seller, content=b"")
        with self.session() as db:
            entry = db.get(Entry, identifier)
            self.assertEqual(entry.status, "completed")
            self.assertEqual(entry.rounds, [])
            self.assertEqual(len(entry.supplementary_files), 2)
        response = asyncio.run(self.ns["sop_supplementary_download"](identifier, 0, reviewer))
        self.assertEqual(Path(response.path).read_bytes(), b"replacement")
        response = asyncio.run(self.ns["sop_supplementary_download"](identifier, 1, reviewer))
        self.assertEqual(Path(response.path).read_bytes(), b"history")
        with self.assertRaises(HTTPException) as private:
            asyncio.run(self.ns["sop_supplementary_download"](identifier, 0, stranger))
        self.assertEqual(private.exception.status_code, 404)

    def test_invalid_choices_and_month_are_rejected(self):
        for change in [dict(customer_name=""), dict(customer_name="手填未知客户"), dict(batch_month="202613"),
                       dict(batch_month="2026-10"), dict(assignee_user_id=999)]:
            with self.subTest(change=change), self.assertRaises(HTTPException):
                self.create(**change)
        with self.session() as db:
            db.query(User).filter_by(id=7).update({"can_upload":False})
            db.commit()
        with self.assertRaises(HTTPException):
            self.create()
        with self.session() as db:
            self.assertEqual(db.query(Entry).count(), 0)

    def test_customer_name_alone_is_unique_across_months(self):
        self.create()
        (Path(self.temp.name) / "客户乙").mkdir()
        self.create(customer_name="客户乙", batch_month="202611")
        with self.assertRaises(HTTPException):
            self.create(batch_month="202611")
        from sqlalchemy.exc import IntegrityError
        with self.session() as db:
            db.add(Entry(customer_name="客户甲"))
            with self.assertRaises(IntegrityError):
                db.commit()

    def import_rows(self, rows, legacy=False):
        columns = ["业务员", "邮箱", "租户", "所属月份批次"]
        if legacy:
            columns.append("项目编号")
        frame = pd.DataFrame(rows, columns=columns)
        async def read():
            return frame.to_csv(index=False).encode("utf-8-sig")
        return asyncio.run(self.ns["sop_import_assignments"](
            SimpleNamespace(filename="list.csv", read=read), "draft", SimpleNamespace(id=1)))

    def test_import_upserts_by_customer_without_number_and_preserves_progress(self):
        self.create()
        with self.session() as db:
            entry = db.query(Entry).one()
            entry.status = "completed"
            entry.description = "已有审核及文件记录"
            original_id = entry.id
            db.add(Entry(customer_name="旧客户", status="draft"))
            db.commit()
        result = self.import_rows([
            ["新业务员", "seller@example.com", "客户甲", "202611"],
            ["业务员", "seller@example.com", "客户乙", "202611"],
            ["业务员", "seller@example.com", " 客户甲 ", "202612"],
            ["业务员", "seller@example.com", "客户戊", "202611"],
            ["业务员", "seller@example.com", "旧客户", "202610"],
        ])
        self.assertEqual((result["created"], result["updated"], len(result["errors"])), (2, 2, 1))
        with self.session() as db:
            entry = db.get(Entry, original_id)
            self.assertEqual((entry.batch_month, entry.salesperson_name, entry.status), ("202611", "新业务员", "completed"))
            self.assertEqual(entry.description, "已有审核及文件记录")
            self.assertEqual(db.query(Entry).filter_by(customer_name="旧客户").count(), 1)
        repeat = self.import_rows([["业务员", "seller@example.com", "客户乙", "202612", "任意旧编号"]], legacy=True)
        self.assertEqual((repeat["created"], repeat["updated"], repeat["errors"]), (0, 1, []))


if __name__ == "__main__":
    unittest.main()
