"""用户模板与名称导入回归，使用独立数据库。"""
import ast
import io
from pathlib import Path
import unittest
from openpyxl import load_workbook, Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from backend.database.models import User, Role, Organization


class UserImportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        self.addCleanup(self.engine.dispose)
        for model in (Role, Organization, User):
            model.__table__.create(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.addCleanup(self.db.close)
        self.db.add_all([
            Role(id=1, name="业务员", permissions={}),
            Role(id=2, name="admin", permissions={"admin": True}),
            Organization(id=42, name="上海组织", is_active=True),
            Organization(id=43, name="停用组织", is_active=False),
        ])
        self.db.commit()
        path = Path(__file__).parent.parent / "admin/user_import.py"
        tree = ast.parse(path.read_text(encoding="utf-8"))
        tree.body = [node for node in tree.body if not (isinstance(node, ast.ImportFrom) and node.level)]
        self.ns = dict(__file__=str(path), User=User, Role=Role, Organization=Organization,
                       get_password_hash=lambda value: "test-hash")
        exec(compile(tree, str(path), "exec"), self.ns)

    def upload(self, org_name):
        book = Workbook()
        sheet = book.active
        sheet.title = "用户导入"
        sheet.append(self.ns["HEADERS"])
        sheet.append(["newuser", "password1", "", "", "", "业务员", org_name, "启用"])
        content = io.BytesIO()
        book.save(content)
        return self.ns["import_users"](content.getvalue(), "users.xlsx", self.db)

    def test_template_has_names_centered_headers_and_no_body_fill(self):
        book = load_workbook(self.ns["build_user_import_template"](self.db))
        sheet = book["用户导入"]
        refs = book["填写选项"]
        self.assertEqual(sheet["G1"].value, "组织名称")
        self.assertEqual(refs["D1"].value, "组织名称")
        self.assertEqual(refs["D2"].value, "上海组织")
        self.assertEqual(refs.max_column, 4)
        self.assertEqual(refs.max_row, 2)
        validation = next(v for v in sheet.data_validations.dataValidation if str(v.sqref) == "G2:G501")
        self.assertEqual(validation.formula1, "'填写选项'!$D$2:$D$2")
        self.assertEqual(sheet["A1"].fill.patternType, "solid")
        for worksheet in book:
            self.assertTrue(worksheet.sheet_view.showGridLines)
            for cell in worksheet[1]:
                if cell.value is not None:
                    self.assertEqual(cell.alignment.horizontal, "center")
                    self.assertEqual(cell.alignment.vertical, "center")
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    self.assertIsNone(cell.fill.patternType)
                    self.assertNotIn("组织ID", str(cell.value))
        self.db.get(Organization, 42).name = "新名称"
        self.db.commit()
        updated = load_workbook(self.ns["build_user_import_template"](self.db))
        self.assertEqual(updated["填写选项"]["D2"].value, "新名称")

    def test_name_resolves_to_id(self):
        self.assertEqual(self.upload(" 上海组织 ")["created"], 1)
        self.assertEqual(self.db.query(User).one().org_id, 42)

    def test_empty_organization_is_optional(self):
        self.assertEqual(self.upload("")["created"], 1)
        self.assertIsNone(self.db.query(User).one().org_id)

    def test_unknown_disabled_or_id_instead_of_name_rejected(self):
        for name in ("未知组织", "停用组织", "42"):
            result = self.upload(name)
            self.assertEqual(result["created"], 0)
            self.assertIn("组织名称不存在或组织已停用", result["errors"][0]["message"])
        self.assertEqual(self.db.query(User).count(), 0)

    def test_duplicate_active_names_rejected(self):
        self.db.add(Organization(name="上海组织", is_active=True))
        self.db.commit()
        result = self.upload("上海组织")
        self.assertEqual(result["created"], 0)
        self.assertIn("多个有效组织", result["errors"][0]["message"])
        self.assertEqual(self.db.query(User).count(), 0)

    def test_disabled_duplicate_does_not_block_active_name(self):
        self.db.add(Organization(name="上海组织", is_active=False))
        self.db.commit()
        self.assertEqual(self.upload("上海组织")["created"], 1)
        self.assertEqual(self.db.query(User).one().org_id, 42)
