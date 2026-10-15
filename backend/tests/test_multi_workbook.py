"""Multi-target training/execution/delivery, without AI or a native Excel runtime."""
import ast
import io
import json
import logging
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook
from backend.utils import multi_workbook as multi
from backend.utils.training_assets import stage_revision, asset_manifest


ROOT = Path(__file__).resolve().parents[2]


def load_function(relative_path, name, namespace=None, class_name=None):
    tree = ast.parse((ROOT / relative_path).read_text(encoding='utf-8-sig'))
    body = tree.body
    if class_name:
        body = next(node.body for node in body if isinstance(node, ast.ClassDef) and node.name == class_name)
    node = next(node for node in body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name == name)
    node.decorator_list = []
    for argument in node.args.args:
        argument.annotation = None
    node.returns = None
    node.args.defaults = [value if isinstance(value, ast.Constant) else ast.Constant(None) for value in node.args.defaults]
    scope = {'__package__': 'backend.api', 'Path': Path, 'os': os,
             'logger': logging.getLogger('multi-test'), **(namespace or {})}
    exec(compile(ast.fix_missing_locations(ast.Module(body=[node], type_ignores=[])), relative_path, 'exec'), scope)
    return scope[name]


def child_script(value, template=None, fail=False):
    return (f'TEMPLATE_PATH = {str(template)!r}\n' if template else '') + f'''
from pathlib import Path
from openpyxl import Workbook, load_workbook
def main():
    {'raise ValueError("business failure")' if fail else 'pass'}
    wb = load_workbook(TEMPLATE_PATH) if {'True' if template else 'False'} else Workbook()
    ws = wb.active
    ws.title = '结果'
    ws['A1'] = 'Contract ends'
    ws['A2'] = '=DATE(2028,12,31)'
    ws['A2'].number_format = 'mm-dd-yy'
    ws['B1'] = 'amount'
    ws['B2'] = {value}
    wb.save(str(Path(output_folder) / 'same_name.xlsx'))
'''


class MultiWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source'
        self.source.mkdir()
        self.output = self.root / 'out'
        self.output.mkdir()

    def targets(self, template=False, fail=False):
        targets = {}
        for index, name in enumerate(['明细.xlsx', '汇总.xlsx']):
            path = self.root / name
            book = Workbook(); book.active['D1'] = f'template-{index}'; book.save(path); book.close()
            targets[name] = {'code': child_script(index + 10, path if template else None, fail and index == 1),
                             'structure': {'expected_file': str(path), 'sheets': {'结果': {'headers': ['Contract ends']}}}}
        return targets

    def execute(self, targets):
        code = multi.build_script(targets)
        scope = {'input_folder': str(self.source), 'output_folder': str(self.output), '__name__': 'test'}
        exec(compile(code, 'multi.py', 'exec'), scope)
        scope['main']()
        return code

    def test_formula_and_template_execution_isolate_names_and_preserve_styles(self):
        for template in [False, True]:
            with self.subTest(template=template):
                targets = self.targets(template)
                self.execute(targets)
                self.assertEqual({p.name for p in self.output.glob('*.xlsx')}, set(targets))
                for index, name in enumerate(targets):
                    wb = load_workbook(self.output / name)
                    self.assertEqual(wb.active['B2'].value, index + 10)
                    self.assertEqual(wb.active['A2'].value, '=DATE(2028,12,31)')
                    self.assertEqual(wb.active['A2']._style.numFmtId, 14)
                    if template:
                        self.assertEqual(wb.active['D1'].value, f'template-{index}')
                    wb.close()

    def test_child_failure_does_not_publish_partial_outputs(self):
        with self.assertRaisesRegex(ValueError, 'business failure'):
            self.execute(self.targets(fail=True))
        self.assertFalse(list(self.output.glob('*.xlsx')))

    def test_missing_target_is_not_silently_ignored(self):
        targets = self.targets()
        structure = multi.combine_structures({n: t['structure'] for n, t in targets.items()})
        with self.assertRaisesRegex(ValueError, '缺少交付文件'):
            multi.select_outputs([], structure)

    def test_generation_scopes_same_named_sheets_and_template_paths(self):
        for template in [False, True]:
            targets = self.targets(template)
            structure = multi.combine_structures({n: t['structure'] for n, t in targets.items()})
            calls = []
            class Generator:
                def generate_code(owner, **kwargs):
                    calls.append(kwargs)
                    return child_script(len(calls)), 'response'
            kwargs = {'input_folder': str(self.source), 'rules_content': '明细规则与汇总规则',
                      'expected_structure': structure}
            if template:
                kwargs.update(template_path='unused', target_sheets={'明细.xlsx': ['结果'], '汇总.xlsx': []})
            code, _ = multi.generate_workbooks(Generator(), kwargs)
            self.assertEqual(set(multi.script_targets(code)), set(targets))
            for call, name in zip(calls, targets):
                self.assertIn(f'【{name}】', call['rules_content'])
                self.assertNotIn('workbooks', call['expected_structure'])
                if template:
                    self.assertEqual(call['template_path'], targets[name]['structure']['expected_file'])
                    self.assertIsInstance(call['target_sheets'], list)

    def test_validation_counts_every_file_and_reports_weighted_accuracy(self):
        targets = self.targets()
        code = self.execute(targets)
        structure = multi.combine_structures({n: t['structure'] for n, t in targets.items()})
        counts = [(3, 2), (7, 7)]
        compared = []
        def compare(actual, expected, diff, **kwargs):
            compared.append(Path(actual).name)
            total, matched = counts[len(compared) - 1]
            return {'total_cells': total, 'matched_cells': matched,
                    'total_differences': total - matched, 'detailed_text': '差异'}
        # Inject lightweight comparator functions: native formula calculation is tested separately.
        import sys
        comparator = SimpleNamespace(compare_excel_files_multi_sheet=compare,
            extract_primary_keys_from_rules=lambda rules: None, require_complete_comparison=lambda result: None)
        with patch.dict(sys.modules, {'backend.utils.excel_comparator': comparator}), patch(
                'backend.utils.training_validation.finalize_training_output', return_value={'formula_report': {}}):
            result = multi.validate_workbooks(self.output, code, structure, 'rules')
        self.assertEqual(result['accuracy'], .9)
        self.assertEqual(set(compared), set(targets))
        self.assertEqual(len(result['output_files']), 2)
        self.assertIn('明细.xlsx', result['detailed_diff'])

    def test_zip_delivery_finalizes_every_file_and_excludes_inputs(self):
        import sys
        from backend.utils import output_postprocess as post
        targets = self.targets()
        code = self.execute(targets)
        structure = multi.combine_structures({n: t['structure'] for n, t in targets.items()})
        (self.output / 'private.txt').write_text('not a deliverable')
        calls = []
        def finalize(path, template, schema, child_code):
            calls.append(Path(path).name)
            return {'rows_processed': 2, 'sheet_summary': []}
        with patch.dict(sys.modules, {'backend.utils.excel_comparator': SimpleNamespace(
                inspect_formula_cache=lambda path: {})}), patch.object(post, 'finalize_output_workbook', finalize), patch.object(
                post, 'dual_output_enabled', return_value=False):
            result = multi.deliver_workbooks(self.output, self.root / 'delivery', code, structure, 'GWS', '计算')
        self.assertEqual(calls, list(targets))
        self.assertEqual(result['output_files'], [result['output_file']])
        self.assertEqual(result['rows_processed'], 4)
        with zipfile.ZipFile(self.root / 'delivery' / result['output_file']) as archive:
            self.assertEqual(set(archive.namelist()), set(targets))
            for name in targets:
                wb = load_workbook(io.BytesIO(archive.read(name)))
                self.assertEqual(wb.active['A2']._style.numFmtId, 14)
                wb.close()

    def test_single_file_selection_and_duplicate_zip_names(self):
        path = self.root / 'only.xlsx'; path.write_bytes(b'test')
        self.assertEqual(multi.select_outputs([path], {'sheets': {}}), [path])
        with self.assertRaises(ValueError):
            multi.create_zip(self.root / 'bad.zip', [path, path])

    def test_asset_revision_keeps_all_targets_and_original_config_unchanged(self):
        targets = self.targets()
        originals = [t['structure']['expected_file'] for t in targets.values()]
        config = {'mode': 'formula'}
        cfg, revision, changed = stage_revision(config, self.root / 'session',
            {'expected': originals}, {}, 'merge')
        self.assertTrue(changed)
        self.assertNotIn('expected_files', config)
        self.assertEqual(len(cfg['expected_files']), 2)
        manifest = asset_manifest(cfg, self.root / 'session')
        self.assertEqual(set(manifest['expected_file_names']), set(targets))

    def test_composite_script_does_not_resolve_one_global_template(self):
        from backend.utils.template_resolver import extract_template_ref
        code = multi.build_script(self.targets(template=True))
        self.assertEqual(extract_template_ref(code), (None, None, None))

    def test_reject_path_names(self):
        for name in ['../bad.xlsx', 'a/b.xlsx', 'a\\b.xlsx', '_diff.xlsx']:
            with self.assertRaises(ValueError):
                multi.validate_name(name)

    def test_both_actual_generators_dispatch_multi_target_contract(self):
        targets = self.targets()
        structure = multi.combine_structures({n: t['structure'] for n, t in targets.items()})
        for class_name, path in [
                ('FormulaCodeGenerator', 'backend/ai_engine/formula_code_generator.py'),
                ('TemplateCodeGenerator', 'backend/ai_engine/template_code_generator.py')]:
            method = load_function(path, 'generate_code', class_name=class_name)
            owner = SimpleNamespace(generate_code=MagicMock(return_value=(child_script(1), 'AI')))
            kwargs = {'input_folder': str(self.source), 'rules_content': '规则', 'expected_structure': structure}
            if class_name == 'TemplateCodeGenerator':
                kwargs['template_path'] = 'ignored global template'
            code, response = method(owner, **kwargs)
            self.assertEqual(set(multi.script_targets(code)), set(targets))
            self.assertEqual(owner.generate_code.call_count, 2)

    def test_multi_target_correction_calls_each_child_and_retains_schema(self):
        targets = self.targets()
        generator = SimpleNamespace(generate_correction_code=MagicMock(side_effect=lambda **kw: kw['original_code']))
        code = multi.correct_workbooks(generator, {'original_code': multi.build_script(targets), 'rules_content': 'rules'})
        corrected = multi.script_targets(code)
        self.assertEqual(set(corrected), set(targets))
        for name in targets:
            self.assertEqual(corrected[name]['structure'], targets[name]['structure'])
        self.assertEqual(generator.generate_correction_code.call_count, 2)

    def test_named_file_precise_edit_keeps_other_file_unchanged(self):
        import sys
        targets = self.targets()
        edits = MagicMock(side_effect=lambda provider, code, feedback, **kw: code.replace('= 10', '= 99'))
        with patch.dict(sys.modules, {'backend.ai_engine.precise_edit': SimpleNamespace(run_precise_edit=edits)}):
            code = multi.precise_edit_workbooks(object(), multi.build_script(targets), '请修改明细.xlsx的金额')
        changed = multi.script_targets(code)
        self.assertEqual(edits.call_count, 1)
        self.assertEqual(changed['汇总.xlsx']['code'], targets['汇总.xlsx']['code'])

    def test_training_iteration_persists_and_downloads_zip_with_zip_headers(self):
        import shutil
        import sys
        from datetime import datetime
        targets = self.targets()
        code = self.execute(targets)
        manager = SimpleNamespace(get_tenant_dir=lambda tenant: self.root)
        persistence = load_function('backend/api/training_chat.py', '_persist_iteration_files', {'shutil': shutil})
        with patch.dict(sys.modules, {'backend.storage.storage_manager': SimpleNamespace(StorageManager=lambda: manager)}):
            paths = persistence('GWS', 1, 1, code, {'output_dir': str(self.output),
                'output_files': [str(self.output / name) for name in targets]})
        with zipfile.ZipFile(paths['output_file']) as archive:
            self.assertEqual(set(archive.namelist()), set(targets))
        session = SimpleNamespace(config={'latest_files': paths}, tenant_id='GWS')
        db = MagicMock(); db.query.return_value.filter_by.return_value.first.return_value = session
        download = load_function('backend/api/training_chat.py', 'download_iteration_file', {
            'TrainingSession': object(), 'datetime': datetime, 're': __import__('re'), '_NO_STORE_HEADERS': {}})
        response_module = SimpleNamespace(FileResponse=lambda path, **kwargs: SimpleNamespace(path=path, **kwargs))
        with patch.dict(sys.modules, {'fastapi.responses': response_module}):
            response = download(1, 'output', db=db)
        self.assertEqual(response.media_type, 'application/zip')
        self.assertTrue(response.filename.endswith('.zip'))
        self.assertEqual(response.path, paths['output_file'])

    def test_multi_asset_paths_rebase_after_windows_to_docker_migration(self):
        from backend.utils.training_assets import resolve_asset_config
        targets = self.targets(template=True)
        old_paths = [rf'C:\old\training_chat\1\{name}' for name in targets]
        structure = multi.combine_structures({n: {'expected_file': old, 'sheets': {}} for n, old in zip(targets, old_paths)})
        cfg = resolve_asset_config({'mode': 'template', 'expected_files': old_paths,
            'expected_file': old_paths[0], 'expected_structure': structure}, self.root)
        self.assertEqual(cfg['expected_files'], [str(self.root / name) for name in targets])
        for name, info in cfg['expected_structure']['workbooks'].items():
            self.assertEqual(info['template_path'], str(self.root / name))

    def test_upload_preprocessing_visits_every_target_and_preserves_single_file(self):
        import shutil
        import sys
        targets = self.targets()
        paths = [t['structure']['expected_file'] for t in targets.values()]
        prepare = load_function('backend/api/training_chat.py', '_prepare_training_uploads_subprocess', {'shutil': shutil})
        shrink = MagicMock()
        with patch.dict(sys.modules, {
            'backend.utils.aspose_helper': SimpleNamespace(is_encrypted=lambda path: False, decrypt_excel=MagicMock()),
            'backend.utils.source_normalizer': SimpleNamespace(convert_xls_to_xlsx=lambda path: path, shrink_inflated_columns=shrink)}):
            result = prepare({'source_dir': str(self.source), 'expected_file': paths[0], 'expected_files': paths})
            self.assertEqual(result['expected_files'], paths)
            self.assertEqual(shrink.call_count, 2)
            result = prepare({'source_dir': str(self.source), 'expected_file': paths[0]})
            self.assertEqual(result['expected_files'], [paths[0]])

    def test_failed_finalization_publishes_no_zip(self):
        import sys
        from backend.utils import output_postprocess as post
        targets = self.targets()
        code = self.execute(targets)
        structure = multi.combine_structures({n: t['structure'] for n, t in targets.items()})
        with patch.dict(sys.modules, {'backend.utils.excel_comparator': SimpleNamespace(
                inspect_formula_cache=lambda path: {})}), patch.object(post, 'finalize_output_workbook',
                side_effect=[{'rows_processed': 1}, RuntimeError('calculation failed')]), patch.object(
                post, 'dual_output_enabled', return_value=False):
            with self.assertRaisesRegex(RuntimeError, 'calculation failed'):
                multi.deliver_workbooks(self.output, self.root / 'delivery', code, structure, 'GWS', '计算')
        self.assertFalse(list((self.root / 'delivery').glob('*.zip')))


if __name__ == '__main__':
    unittest.main()
