"""SOP AI payload regression tests without a server or external AI calls."""
import ast
from datetime import datetime
import json
import logging
import os
from pathlib import Path
import re
import tempfile
import sys
from types import SimpleNamespace
from typing import Optional
import unittest
from unittest.mock import Mock, patch

import openpyxl

from backend.ai_engine.document_parser import DocumentParser
from backend.utils.sop_privacy import SopPrivacy
from backend.utils.sop_reader import read_sop_excel


source = Path(__file__).parent.parent.joinpath('api/tools.py').read_text(encoding='utf-8')
tree = ast.parse(source)
nodes = [node for node in tree.body if isinstance(node, ast.FunctionDef)
         and node.name in {'_sop_read_file', '_build_sop_prompt'}]
namespace = dict(Path=Path, os=os, re=re, DocumentParser=DocumentParser,
                 SopEntry=object, Optional=Optional, SOP_SYSTEM_PROMPT='Review samples only.')
exec(compile(ast.Module(body=nodes, type_ignores=[]), 'sop_payload', 'exec'), namespace)


class SopSheetPrivacyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.parser = Mock()
        self.parser.parse_excel_file.side_effect = self.parsed_workbook
        parser_module = SimpleNamespace(IntelligentExcelParser=lambda: self.parser)
        parser_patch = patch.dict(sys.modules, {'excel_parser': parser_module})
        parser_patch.start()
        self.addCleanup(parser_patch.stop)

    def parsed_workbook(self, path, **options):
        # Supply parser contract fixtures without loading the .NET/Aspose runtime.
        wb = openpyxl.load_workbook(path, data_only=False)
        sheets = []
        try:
            for ws in wb.worksheets:
                if ws.sheet_state != 'visible':
                    continue
                rows = list(ws.iter_rows())
                if not rows or all(cell.value is None for cell in rows[0]):
                    continue
                headers = {str(cell.value): cell.column_letter for cell in rows[0] if cell.value is not None}
                data, formulas = [], {}
                for row in rows[1:4]:
                    sample = {}
                    for cell in row:
                        if cell.data_type == 'f':
                            formulas[cell.coordinate] = cell.value
                            sample[cell.column_letter] = None
                        else:
                            sample[cell.column_letter] = cell.value
                    data.append(sample)
                region = SimpleNamespace(head_data=headers, data=data, formula=formulas,
                                         head_row_start=1, head_row_end=1, data_row_start=2,
                                         data_row_end=ws.max_row, column_schemas={'E': {'field_type': 'decimal', 'number_format': '0.00'}})
                sheets.append(SimpleNamespace(sheet_name=ws.title, regions=[region]))
        finally:
            wb.close()
        return sheets

    def workbook(self, name='source.xlsx'):
        wb = openpyxl.Workbook()
        wb.active.title = '结果表'
        for title in ['结果表', '人员源数据', '考勤源数据']:
            ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
            ws.append(['姓名', '身份证号', '通讯地址', '工号', '基本工资', '计算结果'])
            ws.append(['张三', '31012319900101123X', '上海市测试路123号', 'EMP001', 5200, '=E2*2'])
            for i in range(3, 40):
                ws.append(['李四', '310123199001011234', '北京市测试路456号', 'EMP002', i, '=E3*2'])
        for title, state in [('隐藏表', 'hidden'), ('深度隐藏表', 'veryHidden')]:
            ws = wb.create_sheet(title)
            ws.sheet_state = state
            ws['A1'] = 'HIDDEN_SECRET'
        wb.create_sheet('空表')
        path = self.root / name
        wb.save(path)
        wb.close()
        return path

    def test_all_visible_sheets_have_samples_and_hidden_sheets_are_excluded(self):
        path = self.workbook()
        before = path.read_bytes()
        text = read_sop_excel(str(path), 180000)
        self.assertEqual(path.read_bytes(), before)
        for sheet in ['结果表', '人员源数据', '考勤源数据', '空表']:
            self.assertIn(f'=== Sheet: {sheet} ===', text)
        self.assertNotIn('隐藏表', text)
        self.assertNotIn('HIDDEN_SECRET', text)
        self.assertIn('示例数据 3:', text)
        self.assertNotIn('示例数据 4:', text)
        self.assertIn('E=5200', text)
        self.assertIn('E=4', text)
        self.assertNotRegex(text, r'\bE=5(?:\s*\||$)')
        self.assertIn('=E2*2', text)
        self.assertIn('EMP001', text)
        for secret in ['张三', '李四', '31012319900101123X', '上海市测试路123号', '北京市测试路456号']:
            self.assertNotIn(secret, text)
        self.assertIn('姓名', text)
        self.assertIn('身份证号', text)
        options = self.parser.parse_excel_file.call_args.kwargs
        self.assertEqual(options['max_data_rows'], 3)
        self.assertTrue(options['read_formulas'])
        self.assertTrue(options['skip_hidden_sheets'])
        self.assertFalse(options['active_sheet_only'])
        self.assertFalse(options['best_region_only'])

    def test_shared_aliases_keep_matches_across_files_and_mask_final_metadata(self):
        privacy = SopPrivacy()
        first = namespace['_sop_read_file'](str(self.workbook('张三_source.xlsx')), privacy=privacy)
        second = namespace['_sop_read_file'](str(self.workbook('result.xlsx')), privacy=privacy)
        token = privacy.alias('张三', '姓名')
        self.assertIn(token, first)
        self.assertIn(token, second)
        entry = SimpleNamespace(customer_name='测试客户', description='联系人：张三，手机：13812345678，邮箱：person@example.com')
        privacy.alias(entry.customer_name, '单位')
        prompt = namespace['_build_sop_prompt'](entry, '按规则审核', '张三_source.xlsx\n' + first, second, '姓名：张三')
        prompt = privacy.redact(prompt)
        for secret in ['张三', '测试客户', '13812345678', 'person@example.com']:
            self.assertNotIn(secret, prompt)
        self.assertEqual(privacy.redact(prompt), prompt)

    def test_markdown_and_word_style_tables_mask_names_and_keep_amounts(self):
        for table in ['| 姓名 | 身份证号 | 地址 | 工资 |\n|---|---|---|---|\n| 王五 | 31012319900101123X | 测试街道 | 5200 |',
                      '姓名 | 身份证号 | 地址 | 工资\n王五 | 31012319900101123X | 测试街道 | 5200']:
            text = SopPrivacy().redact(table)
            for secret in ['王五', '31012319900101123X', '测试街道']:
                self.assertNotIn(secret, text)
            self.assertIn('5200', text)

    def test_blank_identifiers_stay_blank_and_distinct_people_do_not_merge(self):
        privacy = SopPrivacy()
        self.assertIsNone(privacy.alias(None, '证件'))
        self.assertEqual(privacy.alias('', '证件'), '')
        self.assertNotEqual(privacy.alias('张三', '姓名'), privacy.alias('李四', '姓名'))
        self.assertEqual(privacy.alias('张三', '姓名'), privacy.alias('张三', '姓名'))
        privacy.alias('1', '地址')
        self.assertIn('=E2*1', privacy.redact('=E2*1；地址：1'))

    def test_large_first_sheet_does_not_remove_later_sheet(self):
        wb = openpyxl.Workbook()
        wb.active.append(['说明'])
        for _ in range(3):
            wb.active.append(['业务说明' * 1000])
        wb.create_sheet('最后的源表').append(['需要保留的字段'])
        path = self.root / 'long.xlsx'
        wb.save(path)
        wb.close()
        text = namespace['_sop_read_file'](str(path))
        self.assertGreater(len(text), 6000)
        self.assertIn('=== Sheet: 最后的源表 ===', text)
        with self.assertRaisesRegex(ValueError, '没有截掉后续 Sheet'):
            read_sop_excel(str(path), 6000)

    def test_document_failure_and_prompt_overflow_stop_review(self):
        path = self.root / 'empty.txt'
        path.touch()
        with self.assertRaises(ValueError):
            namespace['_sop_read_file'](str(path))
        entry = SimpleNamespace(customer_name='客户', description='')
        with self.assertRaisesRegex(ValueError, '没有截掉后续 Sheet'):
            namespace['_build_sop_prompt'](entry, '规则' * 90000, '源数据', '结果', None)

    def test_actual_ai_call_and_saved_prompt_are_redacted(self):
        result_path = str(self.workbook('result.xlsx'))
        rule_path = self.root / 'rules.txt'
        rule_path.write_text('联系人：王五，身份证号：31012319900101123X', encoding='utf-8')
        for unreadable in [False, True]:
            with self.subTest(unreadable=unreadable):
                current = SimpleNamespace(id=7, status='ai_analyzing', source_file_paths=[], source_file_path='',
                                          result_file_path=result_path, rule_file_path=str(rule_path), round_no=1)
                entry = SimpleNamespace(id=3, latest_round_id=7, customer_name='隐私客户', description='联系人：王五')
                round_model, entry_model, rule_model, job_model = object(), object(), object(), object()
                queries = {model: Mock() for model in [round_model, entry_model, rule_model, job_model]}
                queries[round_model].filter_by.return_value.first.return_value = current
                queries[entry_model].filter_by.return_value.first.return_value = entry
                queries[rule_model].all.return_value = []
                db = Mock()
                db.query.side_effect = lambda model: queries[model]
                provider = SimpleNamespace(model='fake', chat=Mock(return_value='{"passed":true}'))
                verdict = SimpleNamespace(passed=True, summary='通过', issues=[], model_dump=lambda: {'passed': True})
                context = dict(namespace, SopPrivacy=SopPrivacy, SessionLocal=lambda: db, SopRound=round_model,
                               SopEntry=entry_model, SopRuleFile=rule_model, SopAiJob=job_model,
                               _sop_effective_rule=lambda *args: None, lock_job=lambda *args: True,
                               AIProviderFactory=SimpleNamespace(create_with_fallback=lambda: provider),
                               SopAiVerdict=SimpleNamespace(model_validate=lambda data: verdict),
                               _sop_notify_ai_pass=Mock(), _sop_notify_rejection=Mock(), datetime=datetime,
                               json=json, logger=logging.getLogger('sop_privacy_test'))
                if unreadable:
                    context['_sop_read_file'] = Mock(side_effect=ValueError('文件无法读取'))
                node = next(node for node in ast.parse(source).body if isinstance(node, ast.FunctionDef)
                            and node.name == '_run_sop_ai_analysis')
                node.body = [statement for statement in node.body if not isinstance(statement, ast.ImportFrom)]
                exec(compile(ast.Module(body=[node], type_ignores=[]), 'sop_analyzer', 'exec'), context)
                if unreadable:
                    with self.assertLogs('sop_privacy_test', level='ERROR'):
                        context['_run_sop_ai_analysis'](7, 3, 1, 'claim')
                    provider.chat.assert_not_called()
                    self.assertEqual(current.status, 'failed')
                else:
                    context['_run_sop_ai_analysis'](7, 3, 1, 'claim')
                    payload = provider.chat.call_args.args[0][1]['content']
                    self.assertEqual(current.prompt_text, payload)
                    for secret in ['张三', '王五', '隐私客户', '31012319900101123X', '上海市测试路123号']:
                        self.assertNotIn(secret, payload)
                    self.assertIn('=== Sheet: 考勤源数据 ===', payload)
                    self.assertIn('E=5200', payload)
                    self.assertEqual(current.status, 'ai_passed')

    def test_multiline_header_is_structure_and_three_data_rows_are_separate(self):
        path = self.workbook()
        region = SimpleNamespace(head_row_start=3, head_row_end=4, data_row_start=5, data_row_end=20,
                                 head_data={'姓名': 'A', '基本工资': 'B'},
                                 data=[{'A': '张三', 'B': 5101}, {'A': '李四', 'B': 5102}, {'A': '王五', 'B': 5103},
                                       {'A': '不应提交', 'B': 5104}],
                                 formula={'B5': '=C5+D5'}, column_schemas={'B': {'field_type': 'decimal'}})
        self.parser.parse_excel_file.side_effect = None
        self.parser.parse_excel_file.return_value = [SimpleNamespace(sheet_name='结果表', regions=[region])]
        text = read_sop_excel(str(path), 180000)
        self.assertIn('表头行 3-4', text)
        self.assertIn('表头：姓名(A) | 基本工资(B)', text)
        for amount in (5101, 5102, 5103):
            self.assertIn(str(amount), text)
        self.assertNotIn('5104', text)
        self.assertNotIn('不应提交', text)
        self.assertNotIn('A=姓名', text)
        self.assertIn('=== Sheet: 人员源数据 ===', text)
        self.assertIn('未识别到数据区域', text)
        self.assertEqual(region.data[0]['A'], '张三')


if __name__ == '__main__':
    unittest.main()
