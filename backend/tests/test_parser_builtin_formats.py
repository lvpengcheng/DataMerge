"""内置格式在 Aspose Custom 为空时仍应保留可复用的 Excel 格式码。"""
from datetime import datetime

import openpyxl

from excel_parser import IntelligentExcelParser, _licensed_workbook


def test_builtin_and_custom_formats_survive_schema_sampling(tmp_path):
    path = tmp_path / "formats.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    cases = [
        ("Contract ends", datetime(2026, 10, 14), "mm-dd-yy", "date", "date"),
        ("数量", 12, "0", "integer", "integer"),
        ("占比", 0.25, "0.00%", "decimal", "percentage"),
        ("备注", "example", "@", "text", "text"),
        ("金额", 12.5, "General", "decimal", "general"),
        ("日期", datetime(2026, 10, 14), "yyyy/m/d;@", "date", "date"),
    ]
    for column, (header, value, code, _, _) in enumerate(cases, 1):
        sheet.cell(1, column, header)
        sheet.cell(2, column, value).number_format = code
    workbook.save(path)
    workbook.close()

    source = _licensed_workbook(str(path))
    cells = source.Worksheets[0].Cells
    parser = IntelligentExcelParser()
    try:
        assert cells[1, 0].GetStyle().Number == 14
        assert not cells[1, 0].GetStyle().Custom
        for column, (header, _, code, field_type, format_type) in enumerate(cases):
            schema = parser._sample_column_schema(
                lambda row, col: cells[row, col], 1, 1, column, header,
            )
            assert schema["field_type"] == field_type
            assert schema["format_type"] == format_type
            assert schema["number_format"] == code
    finally:
        source.Dispose()
