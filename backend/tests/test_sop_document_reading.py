"""Exercise SOP rule reading through the shared document parser."""
import ast
import os
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import patch

from backend.ai_engine.document_parser import DocumentParser


source = Path(__file__).parent.parent.joinpath("api/tools.py").read_text(encoding="utf-8")
nodes = [node for node in ast.parse(source).body if isinstance(node, ast.FunctionDef)
         and node.name == "_sop_read_file"]
namespace = dict(Path=Path, os=os, DocumentParser=DocumentParser)
exec(compile(ast.Module(body=nodes, type_ignores=[]), "sop_reading", "exec"), namespace)


class SopDocumentReadingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def read(self, path):
        return namespace["_sop_read_file"](str(path), max_chars=30000)

    def test_text_and_markdown(self):
        for extension in ("txt", "md", "markdown"):
            with self.subTest(extension=extension):
                path = self.folder / ("rules." + extension)
                path.write_text("# 审核规则\n必须保留客户名称。", encoding="utf-8")
                self.assertIn("必须保留客户名称", self.read(path))
                self.assertTrue(DocumentParser().is_supported_format(str(path)))

    def test_word_paragraphs_and_tables(self):
        from docx import Document
        path = self.folder / "rules.docx"
        doc = Document()
        doc.add_paragraph("Customer name is required.")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Keep all customer records."
        doc.save(path)
        text = self.read(path)
        self.assertIn("Customer name is required.", text)
        self.assertIn("Keep all customer records.", text)

    def test_pdf_text(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.folder / "rules.pdf"
        canvas = Canvas(str(path))
        canvas.drawString(50, 750, "Customer name is required.")
        canvas.save()
        self.assertIn("Customer name is required.", self.read(path))

    def test_legacy_word_uses_existing_converter(self):
        path = self.folder / "rules.doc"
        path.write_bytes(b"legacy word fixture")
        converted = types.SimpleNamespace(returncode=0, stdout="Customer name is required.")
        with patch("subprocess.run", return_value=converted) as converter:
            self.assertEqual(self.read(path), converted.stdout)
        self.assertEqual(converter.call_args.args[0], ["antiword", str(path)])


if __name__ == "__main__":
    unittest.main()
