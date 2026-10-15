"""验证生成模板、公式计算和输出后处理保留日期格式。"""
import ast
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from backend.utils.column_formats import resolve_column_format
from backend.utils.output_postprocess import finalize_output_workbook


def test_date_fallback_preserves_explicit_formats():
    assert resolve_column_format({'field_type': 'date', 'number_format': 'mm/dd/yy'}) == 'mm/dd/yy'
    assert resolve_column_format({'field_type': 'date', 'number_format': 'General'}, 'yyyy/m/d;@') == 'yyyy/m/d;@'
    assert resolve_column_format({'field_type': 'datetime'}) == 'yyyy-mm-dd hh:mm:ss'


@pytest.mark.parametrize('legacy_schema', [False, True])
def test_generated_format_policy_survives_final_calculation(tmp_path, legacy_schema):
    # 提取实际生成模板中的函数，避免加载不相关的 AI 服务依赖。
    source = Path(__file__).parents[1] / 'ai_engine' / 'formula_code_generator.py'
    tree = ast.parse(source.read_text(encoding='utf-8'))
    template = next(n.value for n in ast.walk(tree) if isinstance(n, ast.Constant)
                    and isinstance(n.value, str) and 'def apply_target_column_schemas(wb):' in n.value)
    function = next(n for n in ast.parse(template).body if isinstance(n, ast.FunctionDef)
                    and n.name == 'apply_target_column_schemas')
    codes = ['mm-dd-yy', 'mm/dd/yy', 'yyyy-mm-dd', 'yyyy/m/d;@', 'yyyy-mm-dd hh:mm:ss']
    schemas = {f'Date {i}': {'field_type': 'datetime' if 'hh' in code else 'date',
                            'number_format': code, 'is_builtin_format': i == 0,
                            'number_format_id': 14 if i == 0 else 164 + i}
               for i, code in enumerate(codes)}
    if legacy_schema:
        for schema in schemas.values():
            schema.pop('is_builtin_format')
            schema.pop('number_format_id')
    ns = {'_TARGET_COLUMN_SCHEMAS': {'Result': schemas}}
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(source), 'exec'), ns)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Result'
    for i, name in enumerate(schemas, 1):
        ws.cell(1, i, name)
        ws.cell(2, i, '=DATE(2026,10,14)+TIME(12,30,0)')
        ws.cell(2, i).number_format = 'yyyy/mm/dd'
    ns['apply_target_column_schemas'](wb)
    path = tmp_path / 'result.xlsx'
    wb.save(path)
    wb.close()
    for after in (False, True):
        if after:
            finalize_output_workbook(path, expected_structure={'sheets': {'Result': {'column_schemas': schemas}}})
        # 检查原始 XML，避免 openpyxl 把同名自定义格式归一化为内置编号后掩盖问题。
        with zipfile.ZipFile(path) as archive:
            xml_ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            styles = ET.fromstring(archive.read('xl/styles.xml'))
            formats = {int(n.attrib['numFmtId']): n.attrib['formatCode']
                       for n in styles.findall('x:numFmts/x:numFmt', xml_ns)}
            xfs = styles.find('x:cellXfs', xml_ns)
            sheet_xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            cells = {c.attrib['r']: c for c in sheet_xml.findall('.//x:sheetData/x:row/x:c', xml_ns)}
            for i, code in enumerate(codes):
                cell = cells[f'{chr(65 + i)}2']
                number_id = int(xfs[int(cell.attrib['s'])].attrib['numFmtId'])
                if i == 0:
                    assert number_id == 14
                else:
                    assert number_id >= 164
                    assert formats[number_id] == code
        result = openpyxl.load_workbook(path, data_only=False)
        for i, code in enumerate(codes, 1):
            cell = result.active.cell(2, i)
            assert cell.number_format == code
            assert cell.data_type == 'f'
        assert result.active['A2']._style.numFmtId == 14
        result.close()
    values = openpyxl.load_workbook(path, data_only=True)
    assert values.active['E2'].value.hour == 12
    assert values.active['E2'].value.minute == 30
    values.close()
