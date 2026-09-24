"""人工确认到计算输入的回归：真实行值、跨 Sheet 同名列、落盘和失败阻断。"""
import copy
import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.utils.datetime import from_excel

from backend.utils.compute_ingest import IngestMeta, resolve_with_confirmations, build_preload, write_ingest, read_meta, read_sources
from backend.utils.confirmed_source_mapping import apply_confirmed_mapping, write_execution_sources
from backend.utils.fast_header_matcher import FastHeaderMatcher
from backend.utils.script_entry import invoke_script_main
from excel_parser import ExcelRegion


def test_file_confirmation_limits_automatic_sheet_matching(tmp_path, monkeypatch):
    """Even stronger header similarity cannot steal a manually assigned payroll file."""
    import backend.utils.compute_precheck as pre
    monkeypatch.setattr(pre, '_check_target_sheets', lambda *args: None)
    monkeypatch.setattr(FastHeaderMatcher, '_is_template_mode', lambda *args: False)
    monkeypatch.setattr(FastHeaderMatcher, '_match_headers',
                        lambda self, incoming, expected: ({c: c for c in expected},
                                                          1.0 if '薪酬标记' in incoming else 0.8))
    train = [dict(file_name=f, sheet_name=s, headers={'工号': 'A', '金额': 'B', '月份': 'C'})
             for f, s in [('导出上月.xlsx', '第一批'), ('薪酬上月.xlsx', '上月1'),
                          ('薪酬上月.xlsx', '上月2')]]
    inputs = [dict(file_name=f, file_path=str(tmp_path / f), sheet_name=s,
                   headers={'工号': 'A', '金额': 'B', '月份': 'C', marker: 'D'})
              for f, s, marker in [('七月薪酬.xlsx', '202607（1）', '薪酬标记'),
                                   ('七月薪酬.xlsx', '202607（2)', '薪酬标记'),
                                   ('七月导出.xlsx', '第一批', '导出标记')]]
    structure = {'files': {f: {'sheets': {s['sheet_name']: {'headers': s['headers']}
                                        for s in train if s['file_name'] == f}}
                           for f in {s['file_name'] for s in train}}}
    meta = IngestMeta(source_dir=str(tmp_path), source_structure=structure,
                      train_sheets=train, input_sheets=inputs)
    result = resolve_with_confirmations(meta, confirmed_renames={
        '七月薪酬.xlsx': '薪酬上月.xlsx', '七月导出.xlsx': '导出上月.xlsx'}, skip_history_check=True)
    assert result.ok
    assert result.file_mapping['七月薪酬.xlsx']['expected_file'] == '薪酬上月.xlsx'
    assert result.file_mapping['七月薪酬.xlsx']['sheet_mapping'] == {
        '202607（1）': '上月1', '202607（2)': '上月2'}
    assert result.file_mapping['七月导出.xlsx']['sheet_mapping'] == {'第一批': '第一批'}


def fixture_data(tmp_path):
    source = tmp_path / 'uploaded.xlsx'
    source.write_bytes(b'original upload')
    headers = {'工号': 'A', '金额': 'B', '本月金额': 'C', '日期': 'D'}
    sheets = [dict(file_name=source.name, file_path=str(source), sheet_name=name, headers=headers)
              for name in ['当月工资', '当月补贴']]
    train = [dict(file_name='trained.xlsx', sheet_name=name, headers={col: letter for col, letter in
             [('工号', 'A'), (amount, 'B'), ('日期', 'C')]})
             for name, amount in [('工资表', '工资'), ('补贴表', '补贴')]]
    structure = {'files': {'trained.xlsx': {'sheets': {s['sheet_name']: {'headers': s['headers']} for s in train}}},
                 'multi_sheet_source': True}
    meta = IngestMeta(source_dir=str(tmp_path), source_structure=structure, train_sheets=train, input_sheets=sheets)
    parsed = {(str(source), name): SimpleNamespace(sheet_name=name, regions=[ExcelRegion(
        head_data=headers, data=[{'A': '001', 'B': old, 'C': new, 'D': 45992}],
        column_formats={'D': 'yyyy-mm-dd'})]) for name, old, new in [('当月工资', 10, 100), ('当月补贴', 20, 200)]}
    confirmed = {'uploaded.xlsx': {'expected_file': 'trained.xlsx',
        'sheet_mapping': {'当月工资': '工资表', '当月补贴': '补贴表'},
        'header_mapping_by_sheet': {'当月工资': {'工号': '工号', '本月金额': '工资', '日期': '日期'},
                                    '当月补贴': {'工号': '工号', '本月金额': '补贴', '日期': '日期'}}}}
    return meta, parsed, confirmed


def test_partial_match_preserved_when_another_sheet_is_absent(tmp_path, monkeypatch):
    import backend.utils.compute_precheck as pre
    monkeypatch.setattr(pre, '_check_target_sheets', lambda *args: None)
    meta, _, _ = fixture_data(tmp_path)
    meta.input_sheets = meta.input_sheets[:1]
    meta.input_sheets[0]['headers'] = dict(meta.train_sheets[0]['headers'])
    result = resolve_with_confirmations(meta, skip_history_check=True)
    assert not result.ok
    assert result.file_mapping['uploaded.xlsx']['sheet_mapping'] == {'当月工资': '工资表'}
    assert result.missing_columns == []
    assert [(x['expected_file'], x['expected_sheet']) for x in result.source_sheet_reviews] == [
        ('trained.xlsx', '补贴表')]


def test_month_number_mapping_preserves_fixed_aliases_and_skips_compensation(tmp_path):
    headers = {'工号': 'A', '金额': 'B', '月份': 'C'}
    train = [dict(file_name='薪酬.xlsx', sheet_name=f'上月{i}', headers=headers) for i in [1, 2, 3]]
    inputs = [dict(file_name='薪酬.xlsx', file_path=str(tmp_path / '七月.xlsx'),
                   sheet_name=s, headers=headers, _confirmed_file='薪酬.xlsx')
              for s in ['经济补偿金', '202607（2)', '202607（1）']]
    result = FastHeaderMatcher().match_headers_only(train, inputs)
    assert not result['success']
    assert result['mapping']['file_mapping']['薪酬.xlsx']['sheet_mapping'] == {
        '202607（1）': '上月1', '202607（2)': '上月2'}


def test_ai_cannot_reassign_confirmed_month_sheet(tmp_path):
    from backend.utils.ai_source_mapping import validate_mapping
    headers = {'金额': 'A'}
    train = [dict(file_name='薪酬.xlsx', sheet_name='上月1', headers=headers)]
    inputs = [dict(file_name='薪酬.xlsx', file_path=str(tmp_path / '七月.xlsx'),
                   sheet_name='202607(2)', headers=headers, _confirmed_file='薪酬.xlsx')]
    with pytest.raises(ValueError, match='人工确认'):
        validate_mapping(FastHeaderMatcher(), train, inputs,
                         {'mappings': [{'training_id': 0, 'actual_id': 0, 'columns': {'金额': '金额'}}]})


def test_confirmed_mapping_survives_disk_and_direct_file_script(tmp_path, monkeypatch):
    meta, parsed, confirmed = fixture_data(tmp_path)
    monkeypatch.setattr(FastHeaderMatcher, 'match_headers_only',
                        lambda *args: pytest.fail('人工已确认全部表，不应重新调用匹配器'))
    import backend.utils.compute_precheck as pre
    monkeypatch.setattr(pre, '_check_target_sheets', lambda *args: None)
    pc = resolve_with_confirmations(meta, confirmed_mapping=confirmed, skip_history_check=True)
    assert pc.ok
    params = tmp_path / '_compute_params.json'
    params.write_text(json.dumps({'pre_validated_mapping': pc.file_mapping}, ensure_ascii=False), encoding='utf-8')
    write_ingest(str(tmp_path), meta, parsed)
    mapping = json.loads(params.read_text(encoding='utf-8'))['pre_validated_mapping']
    preload = build_preload(read_meta(str(tmp_path)), read_sources(str(tmp_path)), mapping)
    assert preload['工资表']['df']['工资'].tolist() == [100]
    assert preload['补贴表']['df']['补贴'].tolist() == [200]
    assert '金额' not in preload['工资表']['columns']
    folder = write_execution_sources(tmp_path / 'execution', mapping, preload, meta.source_structure)
    def direct_script(input_folder, output_folder):
        wb = openpyxl.load_workbook(Path(input_folder) / 'trained.xlsx', data_only=True)
        try:
            assert wb['工资表']['C2'].value == from_excel(45992)
            return wb['工资表']['B2'].value + wb['补贴表']['B2'].value
        finally:
            wb.close()
    assert invoke_script_main(direct_script, {'input_folder': folder, 'output_folder': str(tmp_path)}) == 300
    assert (tmp_path / 'uploaded.xlsx').read_bytes() == b'original upload'
    assert parsed[(str(tmp_path / 'uploaded.xlsx'), '当月工资')].regions[0].data[0]['B'] == 10


def test_manual_column_replaces_existing_same_named_column(tmp_path):
    meta, parsed, confirmed = fixture_data(tmp_path)
    meta.train_sheets[0]['headers'] = {'工号': 'A', '金额': 'B', '日期': 'C'}
    confirmed['uploaded.xlsx']['header_mapping_by_sheet']['当月工资']['本月金额'] = '金额'
    mapping = apply_confirmed_mapping(meta, {}, confirmed)
    data = build_preload(meta, parsed, mapping)
    assert data['工资表']['df']['金额'].tolist() == [100]


@pytest.mark.parametrize('change', ['missing_column', 'missing_sheet', 'duplicate_target', 'incomplete'])
def test_invalid_confirmation_stops_instead_of_falling_back(tmp_path, monkeypatch, change):
    meta, _, confirmed = fixture_data(tmp_path)
    info = confirmed['uploaded.xlsx']
    if change == 'missing_column':
        info['header_mapping_by_sheet']['当月工资']['不存在'] = '工资'
    elif change == 'missing_sheet':
        info['sheet_mapping']['当月工资'] = '不存在'
    elif change == 'duplicate_target':
        info['header_mapping_by_sheet']['当月工资']['金额'] = '工资'
    else:
        del info['header_mapping_by_sheet']['当月工资']['本月金额']
    monkeypatch.setattr(FastHeaderMatcher, 'match_headers_only', lambda *a: pytest.fail('无效人工匹配不应回退自动匹配'))
    result = resolve_with_confirmations(meta, confirmed_mapping=confirmed, skip_history_check=True)
    assert not result.ok
    assert result.missing_columns[-1]['error']


def test_automatic_mapping_keeps_column_names_scoped_per_sheet(tmp_path):
    meta, parsed, _ = fixture_data(tmp_path)
    matches = [dict(input_file='uploaded.xlsx', input_file_path=str(tmp_path / 'uploaded.xlsx'),
                    train_file='trained.xlsx', input_sheet=source, train_sheet=target,
                    col_mapping={'工号': '工号', '本月金额': amount, '日期': '日期'}, needs_rewrite=True)
               for source, target, amount in [('当月工资', '工资表', '工资'), ('当月补贴', '补贴表', '补贴')]]
    mapping = FastHeaderMatcher()._build_file_mapping(matches)
    before = copy.deepcopy(mapping)
    data = build_preload(meta, parsed, mapping)
    assert '工资' in data['工资表']['columns'] and '补贴' not in data['工资表']['columns']
    assert '补贴' in data['补贴表']['columns'] and '工资' not in data['补贴表']['columns']
    assert mapping == before


def test_extra_source_columns_are_not_listed_as_expected_mappings(tmp_path):
    from backend.utils.compute_ingest import _build_virtual_sheets, _compose_file_mapping
    meta, _, _ = fixture_data(tmp_path)
    virtual, xlate = _build_virtual_sheets(meta.input_sheets[:1], {}, None)
    automatic = {'uploaded.xlsx': {'expected_file': 'trained.xlsx', 'file_path': virtual[0]['file_path'],
        'sheet_mapping': {'当月工资': '工资表'}, 'header_mapping': {'工号': '工号', '本月金额': '工资', '日期': '日期'}}}
    mapping = _compose_file_mapping(automatic, xlate)
    assert '金额' not in mapping['uploaded.xlsx']['header_mapping_by_sheet']['当月工资']


def test_pending_response_includes_final_mapping():
    import ast
    from typing import Optional
    from backend.utils.compute_precheck import PrecheckResult
    tree = ast.parse((Path(__file__).parent / 'app/main.py').read_text(encoding='utf-8'))
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == '_compute_pending_payload')
    ns = {'Optional': Optional, '_precheck_requires_user_action': lambda pc, skip=False: bool(
        pc.source_sheet_reviews or pc.rename_candidates or pc.missing_files
        or pc.target_candidates or (pc.history_warnings and not skip))}
    exec(compile(ast.Module(body=[node], type_ignores=[]), '<pending>', 'exec'), ns)
    result = PrecheckResult(ok=False, file_mapping={'upload.xlsx': {'expected_file': 'training.xlsx'}})
    payload = ns['_compute_pending_payload'](result, 'session')
    assert payload['file_mapping'] == result.file_mapping
    assert payload['session_id'] == 'session'


def test_pending_response_separates_final_mapping_from_review_subset():
    import ast
    from typing import Optional
    from backend.utils.compute_precheck import PrecheckResult
    tree = ast.parse((Path(__file__).parent / 'app/main.py').read_text(encoding='utf-8'))
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == '_compute_pending_payload')
    ns = {'Optional': Optional, '_precheck_requires_user_action': lambda pc, skip=False: bool(
        pc.source_sheet_reviews or pc.rename_candidates or pc.missing_files
        or pc.target_candidates or (pc.history_warnings and not skip))}
    exec(compile(ast.Module(body=[node], type_ignores=[]), '<pending>', 'exec'), ns)
    result = PrecheckResult(
        ok=False,
        missing_columns=[{'file': 'review.xlsx', 'sheet': '待确认', 'expected_columns': ['金额']}],
        actual_paths=['base.xlsx > 基础 > 工号', 'upload.xlsx > Data > Amount'],
        actual_sources=[
            {'file': 'base.xlsx', 'sheet': '基础'},
            {'file': 'upload.xlsx', 'sheet': 'Data'},
        ],
        file_mapping={
            'base.xlsx': {
                'expected_file': 'base.xlsx', 'sheet_mapping': {'基础': '基础'},
                'header_mapping_by_sheet': {'基础': {'工号': '工号'}},
            },
            'upload.xlsx': {
                'expected_file': 'review.xlsx', 'sheet_mapping': {'Data': '待确认'},
                'header_mapping_by_sheet': {'Data': {'Amount': '金额'}},
            },
        },
    )

    payload = ns['_compute_pending_payload'](result, 'session')

    assert set(payload['file_mapping']) == {'base.xlsx', 'upload.xlsx'}
    assert set(payload['review_file_mapping']) == {'upload.xlsx'}
    assert payload['actual_paths'] == ['upload.xlsx > Data > Amount']
    assert payload['actual_sources'] == [{'file': 'upload.xlsx', 'sheet': 'Data'}]


def test_sheet_review_payload_never_carries_column_mapping():
    import ast
    from typing import Optional
    from backend.utils.compute_precheck import PrecheckResult
    tree = ast.parse((Path(__file__).parent / 'app/main.py').read_text(encoding='utf-8'))
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == '_compute_pending_payload')
    ns = {'Optional': Optional, '_precheck_requires_user_action': lambda pc, skip=False: bool(
        pc.source_sheet_reviews or pc.rename_candidates or pc.missing_files
        or pc.target_candidates or (pc.history_warnings and not skip))}
    exec(compile(ast.Module(body=[node], type_ignores=[]), '<pending>', 'exec'), ns)
    result = PrecheckResult(
        ok=False,
        source_sheet_reviews=[{
            'expected_file': 'train.xlsx', 'expected_sheet': '工资',
            'suggested_file': 'upload.xlsx', 'suggested_sheet': 'Payroll',
        }],
        actual_paths=['upload.xlsx > Payroll > ID'],
        actual_sources=[{'file': 'upload.xlsx', 'sheet': 'Payroll'}],
        file_mapping={'upload.xlsx': {
            'expected_file': 'train.xlsx',
            'sheet_mapping': {'Payroll': '工资'},
            'header_mapping': {'ID': '工号'},
            'header_mapping_by_sheet': {'Payroll': {'ID': '工号'}},
        }},
    )

    payload = ns['_compute_pending_payload'](result, 'session')

    review = payload['review_file_mapping']['upload.xlsx']
    assert review['sheet_mapping'] == {'Payroll': '工资'}
    assert review['header_mapping'] == {}
    assert review['header_mapping_by_sheet'] == {'Payroll': {}}
    assert review['selected_columns_by_sheet'] == {'Payroll': None}


def test_session_confirmation_keeps_prior_rounds_and_explicit_skips(tmp_path):
    from backend.utils.confirmed_source_mapping import save_confirmation_state
    meta, parsed, mapping = fixture_data(tmp_path)
    first = copy.deepcopy(mapping)
    del first['uploaded.xlsx']['sheet_mapping']['当月补贴']
    del first['uploaded.xlsx']['header_mapping_by_sheet']['当月补贴']
    save_confirmation_state(tmp_path, {'confirmed_mapping': {'file_mapping': first},
        'confirmed_target_map': {'结果': ''}, 'skip_history_check': True}, {'script_content': 'not persisted'})
    second = copy.deepcopy(mapping)
    del second['uploaded.xlsx']['sheet_mapping']['当月工资']
    del second['uploaded.xlsx']['header_mapping_by_sheet']['当月工资']
    state = save_confirmation_state(tmp_path, {'confirmed_mapping': {'file_mapping': second},
                                              'confirmed_renames': {'暂缺.xlsx': ''}})
    assert state['confirmed_mapping']['file_mapping'] == mapping
    assert state['confirmed_target_map'] == {'结果': ''}
    assert state['confirmed_renames'] == {'暂缺.xlsx': ''}
    assert state['skip_history_check'] is True
    assert 'script_content' not in json.loads((tmp_path / '_confirmations.json').read_text(encoding='utf-8'))
    # 下一轮只提交目标表，源匹配保持可用，不再要求重选。
    state = save_confirmation_state(tmp_path, {'confirmed_target_map': {'结果二': '本月结果'}})
    assert apply_confirmed_mapping(meta, {}, state['confirmed_mapping'])


def test_sheet_only_confirmation_survives_session_merge(tmp_path):
    """只选择源 Sheet、尚未选择列时，Sheet 关系也必须持久化到最终映射。"""
    from backend.utils.confirmed_source_mapping import save_confirmation_state
    state = save_confirmation_state(tmp_path, {'confirmed_mapping': {'file_mapping': {
        '智算上传.xlsx': {
            'expected_file': '智训源.xlsx',
            'sheet_mapping': {'202609': '工资表'},
            'header_mapping_by_sheet': {'202609': {}},
        },
    }}})
    info = state['confirmed_mapping']['file_mapping']['智算上传.xlsx']
    assert info['sheet_mapping'] == {'202609': '工资表'}
    assert info['header_mapping_by_sheet'] == {'202609': {}}


def test_incomplete_manual_columns_still_build_confirmed_source_sheet(tmp_path, monkeypatch):
    """最终放行时列未补齐也不能丢掉已经人工确认的源 Sheet。"""
    import backend.utils.compute_precheck as pre
    meta, parsed, _ = fixture_data(tmp_path)
    confirmed = {'uploaded.xlsx': {
        'expected_file': 'trained.xlsx',
        'sheet_mapping': {'当月工资': '工资表'},
        'header_mapping_by_sheet': {'当月工资': {}},
    }}
    monkeypatch.setattr(pre, '_check_target_sheets', lambda *args: None)

    result = resolve_with_confirmations(
        meta, confirmed_mapping=confirmed, skip_history_check=True)

    assert not result.ok
    info = result.file_mapping['uploaded.xlsx']
    assert info['file_path'] == str(tmp_path / 'uploaded.xlsx')
    assert info['sheet_mapping'] == {'当月工资': '工资表'}
    source_data = build_preload(meta, parsed, result.file_mapping)
    assert '工资表' in source_data
    assert source_data['工资表']['df']['工号'].tolist() == ['001']


def test_all_pending_categories_are_collected_even_if_source_mapping_invalid(tmp_path, monkeypatch):
    import backend.utils.compute_precheck as pre
    meta, _, mapping = fixture_data(tmp_path)
    mapping['uploaded.xlsx']['header_mapping_by_sheet']['当月工资']['不存在'] = '工资'
    def history(*args):
        args[-2].history_warnings = ['缺少历史数据']
    def target(*args):
        # 最后一项是 AI provider，PrecheckResult 位于倒数第二项。
        args[-2].target_candidates = [{'key': '结果', 'candidates': []}]
    monkeypatch.setattr(pre, '_check_history', history)
    monkeypatch.setattr(pre, '_check_target_sheets', target)
    result = resolve_with_confirmations(meta, confirmed_mapping=mapping)
    assert not result.ok and result.history_warnings and result.target_candidates
    info = result.file_mapping['uploaded.xlsx']
    assert info['expected_file'] == mapping['uploaded.xlsx']['expected_file']
    assert info['sheet_mapping'] == mapping['uploaded.xlsx']['sheet_mapping']
    assert info['header_mapping_by_sheet'] == mapping['uploaded.xlsx']['header_mapping_by_sheet']
    assert info['file_path'] == str(tmp_path / 'uploaded.xlsx')
    assert info['needs_rewrite'] is True


def test_explicit_unmatched_column_continues_without_fabricating_values(tmp_path, monkeypatch):
    import backend.utils.compute_precheck as pre
    from backend.utils.confirmed_source_mapping import save_confirmation_state
    meta, parsed, mapping = fixture_data(tmp_path)
    save_confirmation_state(tmp_path, {'confirmed_mapping': {'file_mapping': mapping}})
    # 用户把此前匹配的工资列改为“无匹配”，不能复用旧映射。
    state = save_confirmation_state(tmp_path, {'confirmed_mapping': {
        'file_mapping': {}, 'unmatched_columns': [['trained.xlsx', '工资表', '工资']]}})
    monkeypatch.setattr(FastHeaderMatcher, 'match_headers_only', lambda *a: pytest.fail('已确认的无匹配不应重新推断'))
    monkeypatch.setattr(pre, '_check_target_sheets', lambda *a: None)
    result = resolve_with_confirmations(meta, confirmed_mapping=state['confirmed_mapping'], skip_history_check=True)
    assert result.ok and result.unmatched_columns == [['trained.xlsx', '工资表', '工资']]
    data = build_preload(meta, parsed, result.file_mapping)
    assert '工资' not in data['工资表']['df'].columns
    assert data['工资表']['df']['工号'].tolist() == ['001']
    folder = write_execution_sources(tmp_path / 'execution', result.file_mapping, data, meta.source_structure)
    # 不依赖工资列的脚本能正常读取其余数据并完成计算。
    wb = openpyxl.load_workbook(Path(folder) / 'trained.xlsx', data_only=True)
    assert wb['工资表']['A2'].value == '001'
    assert wb['补贴表']['B2'].value == 200
    wb.close()
    restored = save_confirmation_state(tmp_path, {'confirmed_mapping': {'file_mapping': mapping}})
    assert restored['confirmed_mapping']['unmatched_columns'] == []


@pytest.mark.parametrize('all_sheets', [False, True])
def test_entire_sheet_or_all_columns_can_be_confirmed_unmatched(tmp_path, monkeypatch, all_sheets):
    import backend.utils.compute_precheck as pre
    from backend.utils.confirmed_source_mapping import fully_unmatched_sheets, merge_confirmation_state
    meta, parsed, mapping = fixture_data(tmp_path)
    skipped = [[s['file_name'], s['sheet_name'], c] for s in meta.train_sheets
               if all_sheets or s['sheet_name'] == '工资表' for c in s['headers']]
    state = merge_confirmation_state({'confirmed_mapping': {'file_mapping': mapping}},
                                    {'confirmed_mapping': {'file_mapping': {}, 'unmatched_columns': skipped}})
    monkeypatch.setattr(FastHeaderMatcher, 'match_headers_only', lambda *a: pytest.fail('无匹配是明确决定'))
    monkeypatch.setattr(pre, '_check_target_sheets', lambda *a: None)
    result = resolve_with_confirmations(meta, confirmed_mapping=state['confirmed_mapping'], skip_history_check=True)
    assert result.ok and not result.missing_columns
    assert len(fully_unmatched_sheets(meta.source_structure, result.unmatched_columns)) == (2 if all_sheets else 1)
    # “所有列无匹配”只清空列关系，不能删除已人工确认的源 Sheet。
    assert result.file_mapping
    expected_sheets = {'当月工资': '工资表', '当月补贴': '补贴表'}
    assert result.file_mapping['uploaded.xlsx']['sheet_mapping'] == expected_sheets
    if all_sheets:
        data = build_preload(meta, parsed, result.file_mapping)
        assert set(data) == {'工资表', '补贴表'}
        assert '本月金额' in data['工资表']['df'].columns


def test_fallback_rewrite_uses_training_file_sheet_and_column_names(tmp_path, monkeypatch):
    """预加载降级重写后，旧脚本读取到的仍应是智训侧名称。"""
    from types import SimpleNamespace
    from excel_parser import IntelligentExcelParser

    parsed = [
        SimpleNamespace(
            sheet_name='未匹配附表',
            regions=[SimpleNamespace(head_data={'说明': 'A'}, data=[{'A': '不应写入'}])],
        ),
        SimpleNamespace(
            sheet_name='202609工资',
            regions=[SimpleNamespace(
                head_data={'员工编号': 'A', '本月工资': 'B'},
                data=[{'A': '001', 'B': 100}],
            )],
        ),
    ]
    monkeypatch.setattr(IntelligentExcelParser, 'parse_excel_file', lambda *args, **kwargs: parsed)
    mapping = {
        'expected_file': '智训源.xlsx',
        'file_path': str(tmp_path / '智算上传.xlsx'),
        'sheet_mapping': {'202609工资': '工资表'},
        'header_mapping_by_sheet': {
            '202609工资': {'员工编号': '工号', '本月工资': '工资'},
        },
    }

    (tmp_path / 'mapped').mkdir()
    result = Path(FastHeaderMatcher.rewrite_excel(mapping, str(tmp_path / 'mapped')))
    wb = openpyxl.load_workbook(result, data_only=True)
    try:
        assert result.name == '智训源.xlsx'
        assert wb.sheetnames == ['工资表']
        assert [wb['工资表']['A1'].value, wb['工资表']['B1'].value] == ['工号', '工资']
        assert [wb['工资表']['A2'].value, wb['工资表']['B2'].value] == ['001', 100]
    finally:
        wb.close()
