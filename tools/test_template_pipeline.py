"""v2 template integration tests: preserve content, layouts and format ownership."""
import ast
import importlib
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch, Mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from backend.utils.script_entry import invoke_script_main
from backend.utils.script_formats import read_explicit_formats
from backend.utils.template_pipeline import (
    region_layout, source_origin, plan_original_sources, source_mode_from_rules,
    is_template_pipeline_v2, uses_original_template_sources,
    write_template_cell, clear_template_cells, rebuilt_source_layouts,
    _sheet_references, sync_source_mode,
)


class TemplatePipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'inputs'
        self.outputs = self.root / 'outputs'
        self.inputs.mkdir()
        self.outputs.mkdir()
        self.template = self.root / '模板.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '目标'
        ws.merge_cells('A1:F1')
        ws['A1'] = '已经存在的标题'
        ws['B4'], ws['C4'], ws['D4'] = '工号', '备注', '出勤天数'
        ws['B5'], ws['C5'] = '001', '保留备注'
        ws['B6'], ws['D6'] = '002', 8
        ws['D8'] = '=SUM(D5:D6)'
        ws['D5'].number_format = '0.00'
        ws['D5'].fill = PatternFill('solid', fgColor='FFFF00')
        ws.row_dimensions[5].height = 31
        wb.create_sheet('辅助')['A1'] = '不能删除的内容'
        wb.save(self.template)
        source = Workbook()
        source.active.title = '原始明细'
        source.active.merge_cells('A1:H1')
        source.active['A1'] = '原始标题'
        source.active['C7'], source.active['F7'] = '工号', '出勤天数'
        source.active['C8'], source.active['F8'] = '001', 20
        source.active['C9'], source.active['F9'] = '002', 22
        source.save(self.inputs / '考勤.xlsx')
        self.region = SimpleNamespace(head_data={'工号': 'C', '出勤天数': 'F'},
            head_row_start=7, head_row_end=7, data_row_start=8, data_row_end=9,
            column_schemas={'C': {'field_type': 'text', 'number_format': '@'},
                            'F': {'field_type': 'integer', 'number_format': '0'}})

    def generator(self, mode='copy'):
        # Code assembly does not require an AI provider or network packages.
        provider = SimpleNamespace(BaseAIProvider=object, AIProviderFactory=object)
        with patch.dict(sys.modules, {'backend.ai_engine.ai_provider': provider}):
            cls = importlib.import_module('backend.ai_engine.template_code_generator').TemplateCodeGenerator
        instance = object.__new__(cls)
        instance._source_sheet_mode = mode
        return instance

    def structures(self):
        target = {'sheets': {'目标': {'regions': [{'header_row': 4, 'data_start_row': 5, 'data_end_row': 6,
            'columns': [{'name': '工号', 'column_letter': 'B'}, {'name': '备注', 'column_letter': 'C'},
                        {'name': '出勤天数', 'column_letter': 'D'}]}]}}}
        layout = region_layout(self.region)
        source = {'files': {'考勤.xlsx': {'原始明细': {'columns': ['工号', '出勤天数'],
                                                        'layout': layout, 'regions': [layout]}}}}
        return target, source

    def build(self, fill, mode='copy'):
        target, source = self.structures()
        return self.generator(mode)._build_complete_code(fill, str(self.template), str(self.inputs),
            target, source, False, reserved_names={'目标', '辅助'})

    def test_real_layout_types_formats_and_actual_column_letters(self):
        generator = self.generator()
        target, source = self.structures()
        mapping, _ = generator._build_source_map_with_letters(source, target)
        layout = next(iter(mapping.values()))
        self.assertEqual(layout['data_start_row'], 8)
        self.assertEqual([c['letter'] for c in layout['columns']], ['C', 'F'])
        self.assertEqual(layout['columns'][0]['number_format'], '@')
        self.assertEqual(generator._guess_column_letter(self.region, '工号'), 'C')

    def test_generated_partial_fill_retains_values_formulas_styles_and_source(self):
        original_bytes = (self.inputs / '考勤.xlsx').read_bytes()
        code = self.build('''def fill_template(wb, source_data, salary_year, salary_month, monthly_hours):
    ws = wb['目标']
    src = next(iter(_SOURCE_MAP))
    assert _SOURCE_MAP[src]['data_start_row'] == 8
    assert _SOURCE_MAP[src]['columns'][0]['letter'] == 'C'
    for row in (5, 6):
        formula = f"=INDEX('{src}'!F:F,MATCH(B{row},'{src}'!C:C,0))"
        write_template_cell(ws, row, 4, formula)
    ws['D5'].number_format = '0.0000'
''')
        env = {'__name__': 'new_template_test', 'input_folder': str(self.inputs), 'output_folder': str(self.outputs)}
        exec(code, env)
        env['load_source_data'] = lambda: {'原始明细': {'df': None, 'source_file': '考勤', 'source_sheet': '原始明细',
            'source_origins': [source_origin(self.inputs / '考勤.xlsx', '原始明细', [self.region])]}}
        env['copy_original_source_sheets'] = Mock()
        invoke_script_main(env['main'], env)
        path = self.outputs / self.template.name
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        ws = wb['目标']
        self.assertEqual(ws['C5'].value, '保留备注')
        self.assertEqual(ws['D6'].value, 8)  # no overwrite authorization
        self.assertEqual(ws['D8'].value, '=SUM(D5:D6)')
        self.assertEqual(ws['D5'].number_format, '0.0000')
        self.assertEqual(ws['D5'].fill.fgColor.rgb, '00FFFF00')
        self.assertEqual(ws.row_dimensions[5].height, 31)
        self.assertEqual(wb['辅助']['A1'].value, '不能删除的内容')
        self.assertEqual(str(ws.merged_cells), 'A1:F1')
        self.assertTrue(ws['D5'].value.startswith('=INDEX'))
        env['copy_original_source_sheets'].assert_called_once()
        plan = env['copy_original_source_sheets'].call_args.args[1]
        self.assertEqual(plan[0]['sheet_name'], '原始明细')
        self.assertTrue(plan[0]['placeholder'])
        self.assertEqual((self.inputs / '考勤.xlsx').read_bytes(), original_bytes)
        self.assertEqual(read_explicit_formats(path), {})  # v2 has no restore manifest
        calls = {node.func.id for node in ast.walk(ast.parse(code))
                 if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)}
        self.assertNotIn('_restore_number_formats', calls)
        self.assertNotIn('_snapshot_number_formats', calls)
        self.assertNotIn('_shrink_inflated_columns', calls)

    def test_v2_retains_general_during_date_assignment_and_explicit_format_is_final(self):
        env = {'__name__': 'format_owner', 'TEMPLATE_PIPELINE_VERSION': 2, 'output_folder': str(self.outputs),
               'input_folder': str(self.inputs), 'Workbook': Workbook, 'datetime': datetime}
        exec('''def main():
    wb = Workbook()
    wb.active['A1'] = datetime(2026, 10, 23)
    wb.active['B1'] = datetime(2026, 10, 23)
    wb.active['B1'].number_format = 'yyyy/mm/dd'
    wb.save(output_folder + '/formats.xlsx')
''', env)
        invoke_script_main(env['main'], env)
        wb = load_workbook(self.outputs / 'formats.xlsx')
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['A1'].number_format, 'General')
        self.assertEqual(wb.active['B1'].number_format, 'yyyy/mm/dd')
        self.assertEqual(read_explicit_formats(self.outputs / 'formats.xlsx'), {})

    def test_explicit_overwrite_and_clear_preserve_formulas_by_default(self):
        wb = load_workbook(self.template)
        self.addCleanup(wb.close)
        ws = wb['目标']
        self.assertFalse(write_template_cell(ws, 6, 4, 22))
        self.assertTrue(write_template_cell(ws, 6, 4, 22, overwrite=True))
        clear_template_cells(ws, 'D5:D8')
        self.assertIsNone(ws['D6'].value)
        self.assertEqual(ws['D8'].value, '=SUM(D5:D6)')
        clear_template_cells(ws, 'D8:D8', include_formulas=True)
        self.assertIsNone(ws['D8'].value)
        self.assertEqual(ws['D5'].number_format, '0.00')

    def test_plan_uses_confirmed_upload_positions_instead_of_training_positions(self):
        layout = region_layout(self.region)
        original = source_origin(self.inputs / '考勤.xlsx', '原始明细', [self.region])
        original['regions'][0]['data_start_row'] = 12
        original['regions'][0]['columns'][0]['letter'] = 'E'
        plan, runtime = plan_original_sources({'训练名': {'source_origins': [original]}},
            {'源_训练名': layout}, {'训练名': '源_训练名'}, self.inputs, {'目标'})
        self.assertEqual(runtime['源_训练名']['data_start_row'], 12)
        self.assertEqual(runtime['源_训练名']['columns'][0]['letter'], 'E')
        self.assertEqual(plan[0]['sheet_name'], '原始明细')

    def test_rebuild_is_explicit_and_uses_normalized_coordinates(self):
        self.assertEqual(source_mode_from_rules('补充出勤天数'), 'copy')
        self.assertEqual(source_mode_from_rules('源表处理方式：重建'), 'rebuild')
        target, source = self.structures()
        mapping, _ = self.generator('rebuild')._build_source_map_with_letters(source, target)
        layout = next(iter(mapping.values()))
        self.assertEqual(layout['data_start_row'], 2)
        self.assertEqual([c['letter'] for c in layout['columns']], ['A', 'B'])
        import pandas as pd
        layouts = rebuilt_source_layouts({'新表': {'df': pd.DataFrame({'工号': ['001']})}}, {'新表': '源_新表'})
        self.assertEqual(layouts['源_新表']['data_end_row'], 2)
        with self.assertRaisesRegex(ValueError, 'rebuild'):
            plan_original_sources({'合并表': {'source_origins': [{}, {}]}}, {}, {}, self.inputs)

    def test_rebuild_executes_explicit_cleaning_and_never_copies_original(self):
        import pandas as pd
        code = self.build('''def clean_source_data(source_data):
    entry = source_data['原始明细']
    entry['df'] = entry['df'].iloc[:1].copy()
    return source_data
def fill_template(wb, source_data, salary_year, salary_month, monthly_hours):
    info = next(iter(_SOURCE_MAP.values()))
    assert info['header_row'] == 1
    assert info['data_start_row'] == 2
''', 'rebuild')
        env = {'__name__': 'rebuilt_template_test', 'input_folder': str(self.inputs), 'output_folder': str(self.outputs)}
        exec(code, env)
        env['load_source_data'] = lambda: {'原始明细': {'df': pd.DataFrame({'工号': ['001', '002'], '出勤天数': [20, 22]})}}
        env['copy_original_source_sheets'] = Mock()
        invoke_script_main(env['main'], env)
        env['copy_original_source_sheets'].assert_not_called()
        wb = load_workbook(self.outputs / self.template.name)
        self.addCleanup(wb.close)
        source = wb['源_原始明细']
        self.assertEqual(source.max_row, 2)
        self.assertEqual(source['A2'].value, '001')
        self.assertEqual(wb['目标']['C5'].value, '保留备注')

    def test_version_isolation_including_multitarget_scripts(self):
        new = 'TEMPLATE_PIPELINE_VERSION = 2\nSOURCE_SHEET_MODE = "copy"'
        old = 'def main(): pass'
        self.assertTrue(is_template_pipeline_v2(new))
        self.assertFalse(is_template_pipeline_v2(old))
        self.assertTrue(uses_original_template_sources('MULTI_WORKBOOK_SCRIPTS = ' + repr({'a.xlsx': {'code': new}})))
        self.assertFalse(uses_original_template_sources('MULTI_WORKBOOK_SCRIPTS = ' + repr({'a.xlsx': {'code': old}})))
        changed = sync_source_mode(new, 'SOURCE_SHEET_MODE: rebuild')
        self.assertFalse(uses_original_template_sources(changed))
        self.assertEqual(sync_source_mode(old, 'SOURCE_SHEET_MODE: rebuild'), old)
        self.assertEqual(_sheet_references("='跨 页''部门'!$A$1:$A$2"), ["跨 页'部门"])
        self.assertEqual(_sheet_references("='[旧文件.xlsx]明细'!$A$1"), [])

    def test_empty_target_has_layout_but_no_automatic_row_generation(self):
        target, _ = self.structures()
        region = target['sheets']['目标']['regions'][0]
        region.update(data_start_row=None, data_end_row=None, header_row_end=4)
        mapping = self.generator()._build_col_map(target)
        self.assertEqual(mapping['目标']['regions'][0]['data_start_row'], 5)
        self.assertIsNone(mapping['目标']['regions'][0]['data_end_row'])

    def test_finalizer_never_calls_legacy_repairs_for_v2(self):
        from backend.utils import output_postprocess as post
        path = self.outputs / 'final.xlsx'
        Workbook().save(path)
        book = SimpleNamespace(Worksheets=SimpleNamespace(Count=0), Settings=SimpleNamespace(),
                               CalculateFormula=Mock())
        book.Save = lambda destination: Workbook().save(destination)
        session = SimpleNamespace(open=lambda _: SimpleNamespace(workbook=book), close=lambda: None)
        with patch.dict(sys.modules, {'aspose_init': SimpleNamespace(ensure_license=lambda: None)}), \
             patch.object(post, '_OutputSession', return_value=session), \
             patch.object(post, 'normalize_key_columns_to_text') as key, \
             patch.object(post, 'restore_formats_from_template') as template, \
             patch.object(post, '_normalize_source_sheet_formats_impl') as source, \
             patch.object(post, '_apply_expected_schemas_in_workbook') as schemas, \
             patch.object(post, 'restore_template_region_format') as summary, \
             patch.object(post, '_sheet_summary_from_workbook', return_value=[]):
            for marker in ('TEMPLATE_PIPELINE_VERSION = 2', 'SOURCE_PIPELINE_VERSION = 2'):
                with self.subTest(marker=marker):
                    result = post.finalize_output_workbook(path, str(self.template), {'sheets': {}}, marker)
        for repair in (key, template, source, schemas, summary):
            repair.assert_not_called()
        self.assertEqual(book.CalculateFormula.call_count, 2)
        self.assertTrue(result['calculated'])

    def test_selected_training_name_can_resolve_confirmed_original_upload(self):
        from backend.utils.script_source_scope import selected_input_folder
        origins = source_origin(self.inputs / '考勤.xlsx', '原始明细', [self.region])
        cache = {'原始明细': {'source_file': '训练考勤', 'source_sheet': '训练明细', 'source_origins': [origins]}}
        with selected_input_folder(self.inputs, ['训练考勤.xlsx'], cache) as selected:
            self.assertEqual([p.name for p in Path(selected).iterdir()], ['训练考勤.xlsx'])
            self.assertEqual((Path(selected) / '训练考勤.xlsx').read_bytes(), (self.inputs / '考勤.xlsx').read_bytes())


if __name__ == '__main__':
    unittest.main()
