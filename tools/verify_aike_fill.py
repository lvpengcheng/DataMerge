"""Exercise provided fill function on real workbook in memory; never save user files."""
import ast
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from backend.utils.openpyxl_compat import ensure_custom_filter_compat
ensure_custom_filter_compat()
script = Path(sys.argv[1])
folder = Path(sys.argv[2])
tree = ast.parse(script.read_text(encoding='utf-8-sig'))
nodes = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'fill_template']
col_map = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
               and any(isinstance(t, ast.Name) and t.id == '_COL_MAP' for t in n.targets))
env = {'_COL_MAP': col_map}
exec(compile(ast.Module(body=nodes, type_ignores=[]), str(script), 'exec'), env)
wb = openpyxl.load_workbook(folder / '薪资报告202609.xlsx')
ws = wb['Kmart_Payroll_detail']
protected = {c.coordinate: (c.value, c.style_id) for row in ws.iter_rows(min_row=1,max_row=211,max_col=76)
             for c in row if c.row <= 3 or c.row == 211 or 32 <= c.column <= 62}
merges = set(str(m) for m in ws.merged_cells.ranges)
for filename in ['Sourcing Staff Change Report.xlsx', '拆分过程表.xlsx', '薪资日历-2026.xlsx']:
    source = openpyxl.load_workbook(folder / filename, read_only=True, data_only=True)
    for s in source:
        dest = wb.create_sheet(('源_' + s.title)[:31])
        for row in s.iter_rows(values_only=True):
            dest.append(row)
    source.close()
env['fill_template'](wb, {}, 2026, 9, 174)
assert merges == set(str(m) for m in ws.merged_cells.ranges)
for coord, value in protected.items():
    assert (ws[coord].value, ws[coord].style_id) == value, coord
for r in range(4, 211):
    assert ws.cell(r, 14).data_type == 'f'
    assert ws.cell(r, 63).data_type == 'f'  # BK: child fee, not AF insurance
    assert ws.cell(r, 71).data_type == 'f'  # BS: compensation
from openpyxl.formula import Tokenizer
missing = set()
for row in ws.iter_rows(min_row=4, max_row=210):
    for cell in row:
        if cell.data_type == 'f':
            for token in Tokenizer(cell.value).items:
                if token.subtype == 'RANGE' and '!' in token.value:
                    sheet = token.value.rsplit('!', 1)[0].strip("'").replace("''", "'")
                    if sheet.startswith('源_') and sheet not in wb.sheetnames:
                        missing.add(sheet)
assert not missing, missing
wb.close()
print('PASS: 207 employee rows filled; merged headers, insurance columns and total row preserved; no files saved.')
