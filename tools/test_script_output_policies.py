"""Regression tests for per-target inputs and explicit formatting precedence.

Run: python -m unittest tools.test_script_output_policies -v
"""
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from backend.utils.script_entry import invoke_script_main
from backend.utils.script_formats import read_explicit_formats, apply_explicit_formats, prefer_script_column_schemas
from backend.utils.script_source_scope import (
    source_files_from_rules, selected_input_folder, filter_source_data, source_scoped_generation,
)
from backend.utils.multi_target_rules import rules_for_workbook
from backend.utils.multi_workbook import execute_workbooks


class OutputPolicies(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source'
        self.source.mkdir()
        self.output = self.root / 'output'
        self.output.mkdir()
        for name in ('人员.xlsx', '考勤.xlsx'):
            wb = Workbook()
            wb.active.title = '明细'
            wb.active.append(['工号', '金额'])
            wb.active.append(['001', 12.5])
            wb.save(self.source / name)

    def run_code(self, code, **extra):
        env = dict(input_folder=str(self.source), output_folder=str(self.output), __name__='policy_test_script', **extra)
        exec(code, env)
        result = invoke_script_main(env['main'], env)
        return env, result

    def test_omission_retains_all_files_and_missing_selection_fails(self):
        with selected_input_folder(str(self.source), None) as path:
            self.assertEqual(path, str(self.source))
        with self.assertRaisesRegex(ValueError, '缺失'):
            with selected_input_folder(self.source, ['缺失.xlsx']):
                self.fail('must fail before executing script')
        with self.assertRaises(ValueError):
            with selected_input_folder(self.source, ['../人员.xlsx']):
                pass

    def test_human_rules_are_scoped_per_target(self):
        rules = ('公共规则\n目标文件：工资.xlsx\n使用源文件：\n- 人员.xlsx\n- 考勤.xlsx\n'
                 '计算工资。\n目标文件：名单.xlsx\n使用源文件：["人员.xlsx"]\n生成名单。')
        self.assertEqual(source_files_from_rules(rules_for_workbook(rules, '名单.xlsx')), ['人员.xlsx'])
        self.assertEqual(source_files_from_rules(rules_for_workbook(rules, '工资.xlsx')), ['人员.xlsx', '考勤.xlsx'])
        self.assertIsNone(source_files_from_rules('规则仅提及人员.xlsx，不指定来源范围'))
        self.assertIsNone(source_files_from_rules('SOURCE_FILES: null'))

    def test_generation_sees_only_selected_files_and_persists_scope(self):
        class Generator:
            @source_scoped_generation
            def generate_code(self, input_folder, rules_content, expected_structure=None):
                self.seen = sorted(p.name for p in Path(input_folder).iterdir())
                return 'def main(): pass', 'response'
        generator = Generator()
        code, _ = generator.generate_code(str(self.source), '使用源文件：["人员.xlsx"]')
        self.assertEqual(generator.seen, ['人员.xlsx'])
        self.assertIn("SOURCE_FILES = ['人员.xlsx']", code)
        self.assertEqual(len(list(self.source.iterdir())), 2)

    def test_correction_preserves_and_can_clear_source_selection(self):
        class Generator:
            @source_scoped_generation
            def correct(self, original_code, rules_content):
                return 'def main(): pass'
        g = Generator()
        original = "SOURCE_FILES = ['人员.xlsx']\ndef main(): pass"
        self.assertIn("SOURCE_FILES = ['人员.xlsx']", g.correct(original, '修正金额'))
        self.assertIn('SOURCE_FILES = None', g.correct(original, 'SOURCE_FILES: null'))

    def test_preloaded_data_filtered_and_keys_match_selected_inputs(self):
        a = {'source_file': '人员', 'source_sheet': '明细', 'df': 'people'}
        b = {'source_file': '考勤', 'source_sheet': '明细', 'df': 'attendance'}
        data = {'人员_明细': a, 'alias': a, '考勤_明细': b}
        self.assertIs(filter_source_data(data, None), data)
        self.assertEqual(filter_source_data(data, ['人员.xlsx']), {'明细': a})
        with self.assertRaisesRegex(ValueError, '没有可读取'):
            filter_source_data(data, ['不存在.xlsx'])
        dotted = {'source_file': '员工.2026', 'source_sheet': '明细'}
        self.assertEqual(filter_source_data({'明细': dotted}, ['员工.2026.xlsx']), {'明细': dotted})

    def test_runtime_filters_cached_loader_even_if_cache_contains_all(self):
        code = '''
SOURCE_FILES = ['人员.xlsx']
def load_source_data():
    return cached
def main():
    from pathlib import Path
    return sorted(Path(input_folder).glob('*.xlsx')), load_source_data()
'''
        cached = {'人员_明细': {'source_file': '人员', 'source_sheet': '明细'},
                  '考勤_明细': {'source_file': '考勤', 'source_sheet': '明细'}}
        env, (files, data) = self.run_code(code, cached=cached)
        self.assertEqual([p.name for p in files], ['人员.xlsx'])
        self.assertEqual(list(data), ['明细'])
        self.assertEqual(env['input_folder'], str(self.source))
        self.assertEqual(len(cached), 2)

    def test_explicit_general_survives_both_script_default_helpers(self):
        self.run_code('''
from openpyxl import Workbook
from datetime import datetime
def _restore_number_formats(wb, snapshot):
    wb.active['A2'].number_format = '0.00'
def apply_target_column_schemas(wb):
    wb.active['A2'].number_format = '@'
    wb.active['B2'].number_format = '0.000'
def main():
    wb = Workbook()
    ws = wb.active
    ws.title = '结果'
    ws['A2'] = 12.5
    ws['A2'].number_format = 'General'
    ws['B2'] = 12.5
    ws['C2'] = datetime(2026, 1, 1)
    _restore_number_formats(wb, {})
    apply_target_column_schemas(wb)
    wb.save(output_folder + '/result.xlsx')
''')
        path = self.output / 'result.xlsx'
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['A2'].number_format, 'General')
        self.assertEqual(wb.active['B2'].number_format, '0.000')
        # Implicit datetime formatting and schema defaults must not become overrides.
        self.assertEqual(read_explicit_formats(path), {'结果': {'A2': 'General'}})

    def test_inherited_row_formats_are_not_explicit_script_overrides(self):
        template = self.root / 'template.xlsx'
        wb = Workbook()
        wb.active['H6'] = datetime(2023, 6, 1)
        wb.active['H6'].number_format = 'mm-dd-yy'
        wb.save(template)
        wb.close()
        self.run_code('''
from openpyxl import load_workbook
def main():
    wb = load_workbook(template_path)
    ws = wb.active
    ws['H7'].value = ws['H6'].value
    ws['H7'].number_format = 'General'
    ws['H7'].number_format = ws['H6'].number_format
    snapshot = {'number_format': ws['H6'].number_format}
    ws['H8'].value = ws['H6'].value
    ws['H8'].number_format = snapshot['number_format']
    ws['H9'].value = ws['H6'].value
    custom_format = 'yyyy-mm-dd'
    ws['H9'].number_format = custom_format
    wb.save(output_folder + '/inherited.xlsx')
''', template_path=str(template))
        self.assertEqual(read_explicit_formats(self.output / 'inherited.xlsx'),
                         {'Sheet': {'H9': 'yyyy-mm-dd'}})

    def test_moved_cells_follow_new_coordinates_and_no_cross_run_leak(self):
        self.run_code('''
from openpyxl import Workbook
def main():
    wb = Workbook()
    ws = wb.active
    ws['A2'] = 12
    ws['A2'].number_format = '0.0000'
    ws.insert_rows(2)
    wb.save(output_folder + '/moved.xlsx')
''')
        self.assertEqual(read_explicit_formats(self.output / 'moved.xlsx'), {'Sheet': {'A3': '0.0000'}})
        self.run_code('''
from openpyxl import Workbook
def main():
    wb = Workbook()
    wb.active['A3'] = 12
    wb.save(output_folder + '/plain.xlsx')
''')
        self.assertEqual(read_explicit_formats(self.output / 'plain.xlsx'), {})

    def test_column_format_and_numeric_values_survive_old_text_schema(self):
        self.run_code('''
from openpyxl import Workbook
def apply_target_column_schemas(wb):
    for cell in wb.active['A']:
        cell.value = str(cell.value)
        cell.number_format = '@'
def main():
    wb = Workbook()
    ws = wb.active
    ws.append(['金额'])
    ws.append([12.5])
    ws.column_dimensions['A'].number_format = '0.00'
    apply_target_column_schemas(wb)
    wb.save(output_folder + '/column.xlsx')
''')
        path = self.output / 'column.xlsx'
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['A2'].value, 12.5)
        self.assertEqual(wb.active['A2'].number_format, '0.00')
        self.assertEqual(read_explicit_formats(path)['Sheet']['A2'], '0.00')

    def test_script_schema_edits_override_saved_schema_without_mutating_it(self):
        original = {'sheets': {'结果': {'column_schemas': {'金额': {'field_type': 'decimal', 'number_format': '0.00'}}}}}
        updated = prefer_script_column_schemas(original,
            "_TARGET_COLUMN_SCHEMAS = {'结果': {'金额': {'field_type': 'decimal', 'number_format': 'General'}}}")
        self.assertEqual(updated['sheets']['结果']['column_schemas']['金额']['number_format'], 'General')
        self.assertEqual(original['sheets']['结果']['column_schemas']['金额']['number_format'], '0.00')

    def test_delivery_applies_overrides_to_renamed_sheet(self):
        # Minimal Aspose boundary double: verify coordinates and format precedence
        # without requiring the machine's native .NET runtime.
        class Cell:
            def __init__(self):
                self.style = SimpleNamespace(Custom='@')
            def GetStyle(self):
                return self.style
            def SetStyle(self, style):
                self.style = style
        cell = Cell()
        wb = SimpleNamespace(Worksheets={'新名称': SimpleNamespace(Cells={(1, 0): cell})})
        apply_explicit_formats(wb, {'旧名称': {'A2': 'General'}}, {'旧名称': '新名称'})
        self.assertEqual(cell.style.Custom, 'General')

    def test_finalization_keeps_override_manifest_on_repeated_delivery(self):
        import sys
        from backend.utils import output_postprocess as post
        self.run_code('''
from openpyxl import Workbook
def main():
    wb = Workbook()
    wb.active['A1'] = 1
    wb.active['A1'].number_format = '0.0000'
    wb.save(output_folder + '/final.xlsx')
''')
        path = self.output / 'final.xlsx'
        book = SimpleNamespace(Worksheets=SimpleNamespace(Count=0), Settings=SimpleNamespace())
        book.CalculateFormula = lambda: None
        def save(destination):
            wb = Workbook()
            wb.active['A1'] = 1
            wb.save(destination)
        book.Save = save
        session = SimpleNamespace(open=lambda _: SimpleNamespace(workbook=book), close=lambda: None)
        with patch.dict(sys.modules, {'aspose_init': SimpleNamespace(ensure_license=lambda: None)}), \
             patch.object(post, '_OutputSession', return_value=session), \
             patch.object(post, 'normalize_key_columns_to_text'), \
             patch.object(post, '_normalize_source_sheet_formats_impl'), \
             patch.object(post, '_sheet_summary_from_workbook', return_value=[]), \
             patch('backend.utils.script_formats.apply_explicit_formats') as apply:
            post.finalize_output_workbook(path)
            post.finalize_output_workbook(path)
        self.assertEqual(apply.call_count, 2)
        self.assertEqual(read_explicit_formats(path), {'Sheet': {'A1': '0.0000'}})

    def test_multitarget_outputs_have_independent_inputs_and_format_manifests(self):
        code = '''
from pathlib import Path
from openpyxl import Workbook
def main():
    wb = Workbook()
    ws = wb.active
    ws.title = '结果'
    ws['A1'] = 1.25
    ws['A1'].number_format = '0.0000'
    for path in sorted(Path(input_folder).glob('*.xlsx')):
        wb.create_sheet('源_' + path.stem)
    wb.save(Path(output_folder) / 'output.xlsx')
'''
        targets = {'selected.xlsx': {'code': code, 'source_files': ['人员.xlsx']},
                   'all.xlsx': {'code': code}}
        with patch('backend.utils.multi_workbook.resolve_child_template', return_value=None):
            execute_workbooks(targets, {'input_folder': str(self.source), 'output_folder': str(self.output)})
        expected = {'selected.xlsx': ['结果', '源_人员'], 'all.xlsx': ['结果', '源_人员', '源_考勤']}
        for name, sheets in expected.items():
            wb = load_workbook(self.output / name)
            self.addCleanup(wb.close)
            self.assertEqual(wb.sheetnames, sheets)
            self.assertEqual(read_explicit_formats(self.output / name), {'结果': {'A1': '0.0000'}})


if __name__ == '__main__':
    unittest.main()
