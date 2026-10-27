"""Explicit date display codes must survive parsing, schema writes and export."""
import importlib.util
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.utils.column_formats import apply_openpyxl_column_format, resolve_column_format


class DateFormatPolicyTests(unittest.TestCase):
    def test_explicit_format_wins_over_stale_builtin_id(self):
        for fmt in ('yyyy-mm-dd', 'yyyy/mm/dd', 'yyyy-mm', 'yyyy-mm-dd hh:mm:ss'):
            with self.subTest(fmt=fmt):
                cell = Workbook().active['A1']
                cell.value = datetime(2023, 12, 1)
                apply_openpyxl_column_format(cell, {
                    'field_type': 'date', 'number_format': fmt,
                    'number_format_id': 14, 'is_builtin_format': True,
                })
                self.assertEqual(cell.number_format, fmt)

    def test_concrete_format_wins_over_regional_builtin_id(self):
        cell = Workbook().active['A1']
        apply_openpyxl_column_format(cell, {
            'field_type': 'date', 'number_format': 'yyyy-mm-dd',
            'number_format_id': 31, 'is_builtin_format': True,
        })
        self.assertEqual(cell.number_format, 'yyyy-mm-dd')

    def test_matching_builtin_and_explicit_short_date_are_preserved(self):
        for fmt, number_id in (('mm-dd-yy', 14), ('0.00', 2), ('h:mm:ss', 21)):
            cell = Workbook().active['A1']
            apply_openpyxl_column_format(cell, {
                'field_type': 'date' if number_id == 14 else 'time' if number_id == 21 else 'decimal',
                'number_format': fmt, 'number_format_id': number_id, 'is_builtin_format': True,
            })
            self.assertEqual(cell.number_format, fmt)
            self.assertEqual(cell._style.numFmtId, number_id)

    def test_missing_date_format_uses_existing_format_or_full_year_default(self):
        self.assertEqual(resolve_column_format({'field_type': 'date'}, 'yyyy/mm/dd'), 'yyyy/mm/dd')
        self.assertEqual(resolve_column_format({'field_type': 'date'}, 'General'), 'yyyy-mm-dd')


@unittest.skipUnless(importlib.util.find_spec('pythonnet'), 'requires application native Excel runtime')
class NativeDateFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import aspose_init
        if not aspose_init.init_aspose():
            raise RuntimeError('native Excel runtime could not initialize')

    def test_parser_retains_custom_dates_and_localized_builtin_display(self):
        from Aspose.Cells import Workbook as NativeWorkbook
        from excel_parser import IntelligentExcelParser
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'example.xlsx'
            wb = Workbook()
            formats = ('yyyy-mm-dd', 'yyyy/mm/dd', 'mm-dd-yy')
            for column, fmt in enumerate(formats, 1):
                wb.active.cell(1, column, f'日期{column}')
                wb.active.cell(2, column, datetime(2023, 12, 1)).number_format = fmt
            wb.save(path)
            wb.close()
            native = NativeWorkbook(str(path))
            try:
                localized = str(native.Worksheets[0].Cells[1, 2].GetStyle().CultureCustom)
            finally:
                native.Dispose()
            regions = IntelligentExcelParser().parse_excel_file(str(path), raise_errors=True)[0].regions
            schemas = regions[0].column_schemas
            self.assertEqual(schemas['A']['number_format'], 'yyyy-mm-dd')
            self.assertEqual(schemas['B']['number_format'], 'yyyy/mm/dd')
            self.assertEqual(schemas['C']['number_format'], localized or 'mm-dd-yy')
            # Rebuilding a workbook must use the concrete parsed display code.
            rebuilt = Workbook()
            for index, letter in enumerate(('A', 'B', 'C'), 1):
                cell = rebuilt.active.cell(1, index, datetime(2023, 12, 1))
                apply_openpyxl_column_format(cell, schemas[letter])
                self.assertEqual(cell.number_format, schemas[letter]['number_format'])
            rebuilt.close()

    def test_native_save_and_values_export_keep_explicit_date_format(self):
        from Aspose.Cells import Workbook as NativeWorkbook
        from backend.utils.column_formats import apply_aspose_column_format
        from backend.utils.output_postprocess import make_values_only_copy, finalize_output_workbook
        with tempfile.TemporaryDirectory() as directory:
            source, values = Path(directory) / 'original.xlsx', Path(directory) / 'values.xlsx'
            wb = Workbook()
            wb.active.title = '结果'
            wb.active['A1'] = '日期'
            wb.active['A2'] = '=DATE(2023,12,1)'
            wb.save(source)
            wb.close()
            native = NativeWorkbook(str(source))
            try:
                cell = native.Worksheets[0].Cells['A2']
                style = cell.GetStyle()
                apply_aspose_column_format(style, {
                    'field_type': 'date', 'number_format': 'yyyy-mm-dd',
                    'number_format_id': 14, 'is_builtin_format': True,
                })
                cell.SetStyle(style)
                native.CalculateFormula()
                native.Save(str(source))
            finally:
                native.Dispose()
            finalize_output_workbook(str(source), expected_structure={
                'sheets': {'结果': {'column_schemas': {'日期': {
                    'field_type': 'date', 'number_format': 'yyyy-mm-dd',
                    'number_format_id': 14, 'is_builtin_format': True,
                }}}},
            })
            self.assertTrue(make_values_only_copy(str(source), str(values), recalculate=False))
            for path in (source, values):
                with self.subTest(path=path.name):
                    result = load_workbook(path, data_only=False)
                    try:
                        self.assertEqual(result['结果']['A2'].number_format, 'yyyy-mm-dd')
                        self.assertEqual(result['结果']['A2'].data_type, 'f' if path == source else 'd')
                    finally:
                        result.close()

    def test_legacy_mixed_template_and_inserted_rows_share_declared_date_format(self):
        from backend.utils.output_postprocess import finalize_output_workbook
        from backend.utils.script_formats import write_explicit_formats
        with tempfile.TemporaryDirectory() as directory:
            template, output = Path(directory) / 'template.xlsx', Path(directory) / 'result.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = '账单明细'
            ws['H5'], ws['J5'] = '入职日期', '工资总额'
            for row, fmt in ((6, 'yyyy-m-d'), (7, 'mm-dd-yy')):
                ws.cell(row, 8, datetime(2023, 6, 1)).number_format = fmt
                ws.cell(row, 10, 1200).number_format = '0.00'
            ws['A8'], ws['J8'] = '合计', '=SUM(J6:J7)'
            ws['J8'].number_format = '#,##0.00'
            ws['H9'] = '=COUNT(H6:H7)'
            ws['H9'].number_format = '0'
            wb.save(template)
            ws['A8'], ws['J8'] = None, None
            ws['H9'] = None
            for row in range(8, 11):
                ws.cell(row, 8, datetime(2024, 8, 19)).number_format = 'mm-dd-yy'
                ws.cell(row, 10, 1200).number_format = '0.00'
            ws['A11'], ws['J11'] = '合计', '=SUM(J6:J10)'
            ws['H8'] = '=DATE(2024,8,19)'
            ws['H12'] = '=COUNT(H6:H10)'
            ws['H12'].number_format = '0'
            ws['H9'].number_format = 'yyyy/mm/dd'  # explicit script setting remains authoritative
            wb.save(output)
            wb.close()
            write_explicit_formats(output, {'账单明细': {'H9': 'yyyy/mm/dd'}})
            schemas = {'sheets': {'账单明细': {'column_schemas': {
                '入职日期': {'field_type': 'date', 'number_format': 'yyyy-mm-dd'},
            }}}}
            code = "_COL_MAP = {'账单明细': {'regions': [{'data_start_row': 6, 'data_end_row': 7}]}}"
            finalize_output_workbook(output, template, schemas, code)
            result = load_workbook(output, data_only=False)
            try:
                ws = result['账单明细']
                self.assertEqual([ws.cell(r, 8).number_format for r in range(6, 11)],
                                 ['yyyy-mm-dd', 'yyyy-mm-dd', 'yyyy-mm-dd', 'yyyy/mm/dd', 'yyyy-mm-dd'])
                self.assertEqual(ws['H6'].value, datetime(2023, 6, 1))
                self.assertEqual(ws['H10'].value, datetime(2024, 8, 19))
                self.assertEqual(ws['H8'].value, '=DATE(2024,8,19)')
                self.assertEqual(ws['J11'].value, '=SUM(J6:J10)')
                self.assertEqual(ws['J11'].number_format, '#,##0.00')
                self.assertEqual(ws['J6'].value, 1200)
                self.assertEqual(ws['H12'].value, '=COUNT(H6:H10)')
                self.assertEqual(ws['H12'].number_format, '0')
            finally:
                result.close()


if __name__ == '__main__':
    unittest.main()
