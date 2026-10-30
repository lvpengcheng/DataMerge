import json
import uuid
from types import SimpleNamespace

import pytest

from backend.api import tools
from backend.utils.integrate_formula_ai import compile_generated_formula
from backend.utils.integrate_engine import eval_source_expr_cross
from backend.utils.merge_engine import compute_header_fingerprint

pytest_plugins = ('backend.tests.test_table_split_api',)


class Provider:
    def __init__(self,*responses):
        self.responses=list(responses);self.calls=[]

    def chat(self,messages,**kwargs):
        self.calls.append(json.loads(json.dumps(messages)))
        return self.responses[min(len(self.calls)-1,len(self.responses)-1)]


@pytest.fixture
def formula_session(api,monkeypatch):
    client,state,_=api
    for action in ('create','edit','apply'):
        state.user.role.permissions['tools.data_integrate.'+action]=True
    session=uuid.uuid4().hex
    columns={'template.xlsx':['员工号','发薪金额','银行卡号'],
             'source.xlsx':['员工号','其他','金额','银行账号'],
             'excluded.xlsx':['员工号','未参与金额']}
    files=[{'name':name,'physical_file':name,'sheet':'Sheet1','columns':cols,
        'fingerprint':compute_header_fingerprint(cols),'suggested_key':'员工号',
        'sample_values':['PRIVATE-EMPLOYEE-123'],'header_start':1,'header_end':1} for name,cols in columns.items()]
    files[0]['physical_file']='发薪结果模板.xlsx';files[0]['sheet']='发薪明细'
    files[1]['physical_file']='员工信息.xlsx';files[1]['sheet']='在职'
    files[1]['field_context']={'银行账号':{'label':'银行账号','header_parts':['个人信息','发放账户','银行账号'],
        'instructions':['用于发薪的银行卡'],
        'profile':{'sampled_count':50,'type_counts':{'text':48,'empty':2},'nonempty_ratio':0.96,
                   'sample_values':['PRIVATE-BANK-987654321']}}}
    for workflow in ('integrate','split'):
        root=tools._workflow_session_dir(session,workflow,state.user);root.mkdir(parents=True,exist_ok=True)
        (root/'_meta.json').write_text(json.dumps({'owner_id':1,'files':files}),encoding='utf-8')
    request={'session_id':session,'main_file':'template.xlsx','default_file':'source.xlsx','target_col':'银行卡号',
        'description':'其他列等于转移取银行账号，否则也取银行账号','active_files':['template.xlsx','source.xlsx'],
        'current_formula':'source.xlsx.银行账号','selected_tokens':['银行账号'],'key_map':{'source.xlsx':'员工号'}}
    from backend.ai_engine.ai_provider import AIProviderFactory
    provider=Provider(json.dumps({'formula':'IF(F1="转移",F3,F3)','explanation':'按标记判断并返回银行卡'}))
    monkeypatch.setattr(AIProviderFactory,'create_provider',lambda:provider)
    return SimpleNamespace(client=client,state=state,request=request,provider=provider)


@pytest.mark.parametrize('workflow',['integrate','table-split'])
def test_both_routes_generate_runtime_safe_text_formulas_without_sending_employee_data(formula_session,workflow):
    s=formula_session
    response=s.client.post('/api/tools/'+workflow+'/generate-formula',json=s.request)
    assert response.status_code==200,response.text
    result=response.json()
    indexes={'source.xlsx':{'cols':['其他','银行账号'],'rows':{'one':[{'其他':'转移','银行账号':'00123456789012345678'}]}}}
    assert eval_source_expr_cross(result['formula'],'source.xlsx',indexes,'one')=='00123456789012345678'
    assert result['text_coercion_applied']
    payload=json.loads(s.provider.calls[0][1]['content'])
    assert 'PRIVATE-EMPLOYEE' not in json.dumps(payload)
    assert 'excluded.xlsx' not in json.dumps(payload)
    assert payload['current_formula']=='F3'
    assert next(f for f in payload['fields'] if f['column']=='银行账号')['selected']
    assert payload['tables'][0]['key_column']=='员工号'
    assert payload['tables'][0]['key_field_id']=='F0'
    assert payload['tables'][0]['physical_file']=='员工信息.xlsx'
    assert payload['tables'][0]['sheet']=='在职'
    assert payload['target']['physical_file']=='发薪结果模板.xlsx'
    assert payload['target']['sheet']=='发薪明细'
    bank=next(f for f in payload['fields'] if f['column']=='银行账号')
    assert bank['header_parts']==['个人信息','发放账户','银行账号']
    assert bank['profile']['type_counts']=={'text':48,'empty':2}
    assert 'PRIVATE-BANK' not in json.dumps(payload)


def test_generation_repairs_unsupported_excel_functions_using_the_interpreter_spec(formula_session):
    s=formula_session;s.request['target_col']='发薪金额'
    s.provider.responses=[json.dumps({'formula':'SUMIF(F0,"001",F2)'}),json.dumps({'formula':'ROUND(F2*0.1,2)'})]
    response=s.client.post('/api/tools/integrate/generate-formula',json=s.request)
    assert response.status_code==200 and len(s.provider.calls)==2,response.text
    assert 'ROUND' in response.json()['formula'] and 'SUMIF' not in response.json()['formula']


@pytest.mark.parametrize('formula',['VLOOKUP(F0,F1,2,0)','F999','SUM(F2: F3)','F2**2','IF(F0,1)',
    '__import__("os").system("echo bad")','F0[0]'])
def test_unknown_references_and_non_interpreter_syntax_never_reach_the_ui(formula_session,formula):
    s=formula_session;s.provider.responses=[json.dumps({'formula':formula})]
    response=s.client.post('/api/tools/integrate/generate-formula',json=s.request)
    assert response.status_code==422,response.text
    assert 'formula' not in response.json()


def test_unsupported_business_rules_request_clarification(formula_session):
    s=formula_session
    s.provider.responses=[json.dumps({'formula':'','clarification':'同一员工的所有行不能逐行筛选，请先准备辅助列'})]
    response=s.client.post('/api/tools/table-split/generate-formula?stream=true',json=s.request)
    assert response.status_code==200 and response.json()['needs_clarification']
    assert response.json()['formula'] is None


def test_session_ownership_scope_and_permissions_are_checked_before_ai(formula_session):
    s=formula_session
    for update in ({'active_files':['template.xlsx','missing.xlsx']},{'target_col':'不存在'},
                   {'key_map':{'source.xlsx':'不存在'}}):
        assert s.client.post('/api/tools/integrate/generate-formula',json={**s.request,**update}).status_code==400
    s.state.user.id=2
    assert s.client.post('/api/tools/integrate/generate-formula',json=s.request).status_code==403
    s.state.user.id=1
    for action in ('create','edit','apply'):s.state.user.role.permissions['tools.data_integrate.'+action]=False
    assert s.client.post('/api/tools/integrate/generate-formula',json=s.request).status_code==403
    assert not s.provider.calls


def test_provider_failure_does_not_leak_internal_details(formula_session,monkeypatch):
    def fail(*args,**kwargs):raise RuntimeError('PRIVATE-API-KEY')
    monkeypatch.setattr(formula_session.provider,'chat',fail)
    response=formula_session.client.post('/api/tools/integrate/generate-formula',json=formula_session.request)
    assert response.status_code==503
    assert 'PRIVATE-API-KEY' not in response.text


def test_description_and_exact_formula_survive_scheme_save_export_import(formula_session):
    s=formula_session
    pair={'a_col':'银行卡号','source_file':'source.xlsx','source_expr':'银行账号&""',
          'formula_description':'按来源返回银行卡号'*700,'formula_expected_result':'转移员工返回银行卡而不是0','formula_exact':True}
    saved=s.client.post('/api/tools/integrate/scheme/save',json={'session_id':s.request['session_id'],
        'name':'描述公式','main_file':'template.xlsx','key_map':{'template.xlsx':'员工号','source.xlsx':'员工号'},
        'overwrite_pairs':[pair]})
    assert saved.status_code==200,saved.text
    exported=s.client.get('/api/tools/integrate/scheme/'+str(saved.json()['id'])+'/export')
    p=exported.json()['config']['overwrite_pairs'][0]
    assert p['formula_description']==pair['formula_description'] and p['formula_exact']
    assert p['formula_expected_result']==pair['formula_expected_result']
    imported=s.client.post('/api/tools/integrate/scheme/import',data={'name':'导入描述公式'},
        files={'file':('方案.json',exported.content,'application/json')})
    assert imported.status_code==200,imported.text


def test_long_description_expected_results_and_output_context_are_available_to_ai(formula_session):
    s=formula_session
    request={**s.request,'description':'完整规则说明。'*1000,'expected_result':'有转移标记时必须得到原银行卡号',
        'output_file':'山东发薪结果.xlsx','output_sheet':'工资明细','date_key_mode':'yearmonthday',
        'union_source_keys':True,'data_scope':{'mode':'source','source_file':'source.xlsx'}}
    response=s.client.post('/api/tools/integrate/generate-formula',json=request)
    assert response.status_code==200,response.text
    payload=json.loads(s.provider.calls[0][1]['content'])
    assert payload['description']==request['description']
    assert payload['expected_result']==request['expected_result']
    assert payload['output']['sheet']=='工资明细'
    assert payload['execution']['data_scope']['source_file']=='source.xlsx'
    assert '预期结果' in s.provider.calls[0][0]['content']
    too_long=s.client.post('/api/tools/integrate/generate-formula',json={**request,'description':'字'*20001})
    assert too_long.status_code==400


def test_serializer_preserves_literals_rounding_and_does_not_turn_numeric_branches_into_text():
    fields=[{'id':'F0','file':'C.xlsx','column':'状态'}, {'id':'F1','file':'C.xlsx','column':'金额'}]
    expression,_,_=compile_generated_formula('IF((F0&"")="a""b",ROUND(F1*0.1,2),0)',fields,'C.xlsx','发薪金额')
    indexes={'C.xlsx':{'cols':['状态','金额'],'rows':{'one':[{'状态':'a"b','金额':123.45}]}}}
    assert eval_source_expr_cross(expression,'C.xlsx',indexes,'one')==12.35


def test_single_workbook_split_can_generate_and_run_a_formula_on_its_original_fields(formula_session):
    import io
    import zipfile
    from openpyxl import Workbook,load_workbook
    s=formula_session
    book=Workbook();ws=book.active
    ws.append(['工号','金额','工作地']);ws.append(['001',10,'上海']);ws.append(['002',20,'上海'])
    stream=io.BytesIO();book.save(stream)
    analyzed=s.client.post('/api/tools/table-split/analyze',files={'files':('input.xlsx',stream.getvalue(),'application/octet-stream')})
    assert analyzed.status_code==200,analyzed.text
    s.provider.responses=[json.dumps({'formula':'F1*2','explanation':'按原金额加倍'})]
    response=s.client.post('/api/tools/table-split/generate-formula',json={
        'session_id':analyzed.json()['session_id'],'main_file':'input.xlsx','default_file':'input.xlsx',
        'target_col':'金额','description':'金额乘以2','active_files':['input.xlsx'],'include_main':True})
    assert response.status_code==200,response.text
    executed=s.client.post('/api/tools/table-split/execute',json={
        'session_id':analyzed.json()['session_id'],'main_file':'input.xlsx','key_map':{'input.xlsx':'工号'},
        'overwrite_pairs':[{'a_col':'金额','source_file':'input.xlsx','source_expr':response.json()['formula']}],
        'split_config':{'mode':'value','column':'工作地'}})
    assert executed.status_code==200,executed.text
    with zipfile.ZipFile(io.BytesIO(executed.content)) as archive:
        result=load_workbook(io.BytesIO(archive.read(archive.namelist()[0])),data_only=True).active
        assert [result.cell(row,2).value for row in (2,3)]==[20,40]
