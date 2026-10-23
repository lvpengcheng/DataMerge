"""Selected-sheet isolation, renamed role confirmation and AI pre-matching."""
import io
import json
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from backend.api import tools
from backend.tests.test_table_split_api import api


def workbook(names=('模板', '工资', '不参与')):
    book = Workbook()
    book.remove(book.active)
    for index, name in enumerate(names):
        sheet = book.create_sheet(name)
        sheet.append(['工号', '金额', '工作地'])
        sheet.append(['001', (index + 1) * 100, '上海'])
        sheet.append(['002', (index + 1) * 200, '北京'])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def inline_all_sheet_parser(monkeypatch):
    async def parse(path, **kwargs):
        return tools._parse_file_all_sheets_impl(path, **kwargs)
    monkeypatch.setattr(tools, '_parse_file_all_sheets_fresh', parse)


def upload(client, tool, names=('模板', '工资', '不参与'), filename='多表.xlsx'):
    result = client.post(f'/api/tools/{tool}/analyze', data={'parse_all_sheets': 'true'},
        files={'files': (filename, workbook(names), 'application/octet-stream')})
    assert result.status_code == 200, result.text
    return result.json()


def select(client, tool, analyzed, sheets):
    selected = [f['name'] for f in analyzed['files'] if f['sheet'] in sheets]
    response = client.post(f'/api/tools/{tool}/select-sheets', json={
        'session_id': analyzed['session_id'], 'selected_tables': selected})
    assert response.status_code == 200, response.text
    return response.json()


def allow_integrate(state):
    for action in ('', '.create', '.apply', '.edit', '.delete'):
        state.user.role.permissions['tools.data_integrate' + action] = True


@pytest.mark.parametrize('tool', ['integrate', 'table-split'])
def test_only_chosen_sheets_are_available_and_unconfirmed_selection_cannot_execute(api, tool):
    client, state, _ = api
    allow_integrate(state)
    analyzed = upload(client, tool)
    assert analyzed['sheet_selection_pending']
    main, source, excluded = [f['name'] for f in analyzed['files']]
    request = {'session_id': analyzed['session_id'], 'main_file': main,
        'key_map': {main: '工号', source: '工号'},
        'overwrite_pairs': [{'a_col': '金额', 'source_file': source, 'source_expr': '金额'}]}
    if tool == 'table-split':
        request['split_config'] = {'mode': 'value', 'column': '工作地'}
    blocked = client.post(f'/api/tools/{tool}/execute', json=request)
    assert blocked.status_code == 400
    selected = select(client, tool, analyzed, ['模板', '工资'])
    assert [f['name'] for f in selected['files']] == [main, source]
    matched = client.post(f'/api/tools/{tool}/match', json={
        'session_id': analyzed['session_id'], 'main_file': main, 'target_cols': ['金额'],
        'source_cols': [{'file': excluded, 'col': '金额'}]})
    assert matched.status_code == 200 and matched.json()['pairs'] == []
    response = client.post(f'/api/tools/{tool}/execute', json=request)
    assert response.status_code == 200, response.text


def test_invalid_or_other_users_sheet_selection_is_rejected(api):
    client, state, _ = api
    analyzed = upload(client, 'table-split')
    path = '/api/tools/table-split/select-sheets'
    for selected in ([], ['unuploaded.xlsx'], [analyzed['files'][0]['name']] * 2):
        assert client.post(path, json={'session_id': analyzed['session_id'], 'selected_tables': selected}).status_code == 400
    state.user.id = 2
    assert client.post(path, json={'session_id': analyzed['session_id'],
        'selected_tables': [analyzed['files'][0]['name']]}).status_code == 403


@pytest.mark.parametrize('tool', ['integrate', 'table-split'])
def test_renamed_sheets_require_confirmation_and_chosen_roles_control_values(api, tool):
    client, state, _ = api
    allow_integrate(state)
    old = select(client, tool, upload(client, tool), ['模板', '工资'])
    main, source = [f['name'] for f in old['files']]
    request = {'session_id': old['session_id'], 'main_file': main, 'name': 'Sheet对应方案',
        'key_map': {main: '工号', source: '工号'},
        'overwrite_pairs': [{'a_col': '金额', 'source_file': source, 'source_expr': '金额'}],
        'data_scope': {'mode': 'source', 'source_file': source}, 'union_source_keys': True}
    if tool == 'table-split':
        request['split_config'] = {'mode': 'value', 'column': '工作地'}
    saved = client.post(f'/api/tools/{tool}/scheme/save', json=request)
    assert saved.status_code == 200, saved.text
    scheme_id = saved.json()['id']
    current = select(client, tool, upload(client, tool, ('新模板', '新工资', '不参与'), '本月.xlsx'), ['新模板', '新工资'])
    url = f'/api/tools/{tool}/scheme/apply-validate'
    check = client.post(url, json={'session_id': current['session_id'], 'scheme_id': scheme_id})
    assert check.status_code == 200, check.text
    info = check.json()
    assert info['sheet_mapping']['required'] is True
    assert info['sheet_mapping']['roles'][0]['saved_sheet'] == '模板'
    chosen = [f['name'] for f in reversed(current['files'])]  # exercise manual role swap
    result = client.post(url, json={'session_id': current['session_id'], 'scheme_id': scheme_id, 'role_files': chosen})
    assert result.status_code == 200 and result.json()['ok'], result.text
    scheme = result.json()['resolved_scheme']
    assert scheme['main_file'] == chosen[0]
    assert scheme['config']['data_scope']['source_file'] == chosen[1]
    assert result.json()['sheet_mapping'] is None
    execution = {'session_id': current['session_id'], 'main_file': chosen[0],
        'key_map': {name: '工号' for name in chosen}, 'union_source_keys': True,
        'data_scope': scheme['config']['data_scope'],
        'overwrite_pairs': [{'a_col': p['a_col'], 'source_file': chosen[p['source_role']],
                             'source_expr': p['source_expr']} for p in scheme['config']['overwrite_pairs']]}
    if tool == 'table-split':
        execution['split_config'] = scheme['config']['split_config']
    output = client.post(f'/api/tools/{tool}/execute', json=execution)
    assert output.status_code == 200, output.text
    contents = [output.content]
    if tool == 'table-split':
        with zipfile.ZipFile(io.BytesIO(output.content)) as archive:
            contents = [archive.read(name) for name in archive.namelist() if name.endswith('.xlsx')]
    amounts = []
    for content in contents:
        sheet = load_workbook(io.BytesIO(content), data_only=True)['新工资']
        amounts.extend(row[1].value for row in sheet.iter_rows(min_row=2) if row[0].value)
    assert sorted(amounts) == [100, 200]  # values from the user-chosen new source
    bad = client.post(url, json={'session_id': current['session_id'], 'scheme_id': scheme_id,
        'role_files': [chosen[0], chosen[0]]})
    assert bad.status_code == 200 and bad.json()['ok'] is False


def test_all_sheet_catalog_lists_empty_visible_sheets_and_can_skip_a_workbook(api):
    client, _, _ = api
    empty = Workbook()
    empty.active.title = '空白页'
    buffer = io.BytesIO()
    empty.save(buffer)
    result = client.post('/api/tools/table-split/analyze', data={'parse_all_sheets': 'true'}, files=[
        ('files', ('有数据.xlsx', workbook(), 'application/octet-stream')),
        ('files', ('空白.xlsx', buffer.getvalue(), 'application/octet-stream'))])
    assert result.status_code == 200, result.text
    analyzed = result.json()
    assert next(book for book in analyzed['workbooks'] if book['file'] == '空白.xlsx')['sheets'] == [
        {'sheet': '空白页', 'selectable': False}]
    selected = select(client, 'table-split', analyzed, ['模板'])
    assert len(selected['files']) == 1 and selected['files'][0]['physical_file'] == '有数据.xlsx'


def test_ai_endpoint_falls_back_without_applying_any_mapping(api, monkeypatch):
    from backend.ai_engine.ai_provider import AIProviderFactory
    client, _, _ = api
    selected = select(client, 'table-split', upload(client, 'table-split'), ['模板', '工资'])
    main, source = [f['name'] for f in selected['files']]
    saved = client.post('/api/tools/table-split/scheme/save', json={
        'session_id': selected['session_id'], 'name': 'AI对应方案', 'main_file': main,
        'key_map': {main: '工号', source: '工号'}, 'overwrite_pairs': [],
        'split_config': {'mode': 'value', 'column': '工作地'}})
    assert saved.status_code == 200, saved.text
    def fail(*args, **kwargs):
        raise RuntimeError('PRIVATE-ERROR')
    monkeypatch.setattr(AIProviderFactory, 'create_provider', fail)
    result = client.post('/api/tools/table-split/suggest-sheet-mapping', json={
        'session_id': selected['session_id'], 'scheme_id': saved.json()['id']})
    assert result.status_code == 200 and result.json()['ai_suggested'] is False
    assert 'PRIVATE-ERROR' not in result.text
    assert set(result.json()['role_files']) == {main, source}


def test_saved_header_ranges_are_restored_only_after_roles_are_confirmed(api):
    client, state, _ = api
    selected = select(client, 'table-split', upload(client, 'table-split'), ['模板', '工资'])
    main, source = [f['name'] for f in selected['files']]
    saved = client.post('/api/tools/table-split/scheme/save', json={
        'session_id': selected['session_id'], 'name': '列头与角色方案', 'main_file': main,
        'key_map': {main: '工号', source: '工号'}, 'overwrite_pairs': [],
        'split_config': {'mode': 'value', 'column': '工作地'}})
    assert saved.status_code == 200, saved.text
    current = select(client, 'table-split', upload(client, 'table-split', ('新模板', '新工资', '不参与')),
                     ['新模板', '新工资'])
    directory = tools._workflow_session_dir(current['session_id'], 'split', state.user)
    path = directory / '_meta.json'
    meta = json.loads(path.read_text(encoding='utf-8'))
    meta['files'][0]['header_start'] = 2
    meta['files'][0]['header_end'] = 2
    path.write_text(json.dumps(meta, ensure_ascii=False), encoding='utf-8')
    request = {'session_id': current['session_id'], 'scheme_id': saved.json()['id']}
    url = '/api/tools/table-split/scheme/apply-validate'
    initial = client.post(url, json=request)
    assert initial.status_code == 200 and initial.json()['ok'], initial.text
    assert json.loads(path.read_text(encoding='utf-8'))['files'][0]['header_start'] == 2
    confirmed = client.post(url, json={**request, 'role_files': [f['name'] for f in current['files']]})
    assert confirmed.status_code == 200 and confirmed.json()['ok'], confirmed.text
    assert confirmed.json()['files'][0]['header_start'] == 1
    assert confirmed.json()['files'][0]['header_manual'] is True


def test_ai_suggestions_are_complete_eligible_and_do_not_include_business_rows():
    from backend.utils.merge_engine import compute_header_fingerprint
    columns = ['工号', '金额']
    fingerprint = compute_header_fingerprint(columns)
    cfg = {'roles': [{'file': 'old__模板.xlsx', 'fp': fingerprint, 'sheet': '模板'},
                     {'file': 'old__工资.xlsx', 'fp': fingerprint, 'sheet': '工资'}],
           'cols_by_fp': {fingerprint: columns}, 'key_map_by_role': {'0': '工号', '1': '工号'}}
    uploaded = [{'name': name, 'sheet': sheet, 'columns': columns, 'fingerprint': fingerprint}
                for name, sheet in [('new__新模板.xlsx', '新模板'), ('new__新工资.xlsx', '新工资')]]
    match = tools._match_scheme_config(cfg, uploaded)
    class Provider:
        def chat(self, messages, **kwargs):
            self.messages = messages
            return self.output
    provider = Provider()
    provider.output = json.dumps({'matches': [{'role_id': 'R0', 'table_id': 'U1'}, {'role_id': 'R1', 'table_id': 'U0'}]})
    assert tools._suggest_sheet_roles(cfg, uploaded, match, provider) == [uploaded[1]['name'], uploaded[0]['name']]
    payload = json.loads(provider.messages[1]['content'])
    assert set(payload) == {'roles', 'tables'}
    assert not any('rows' in f for f in payload['tables'])
    for invalid in ([{'role_id': 'R0', 'table_id': 'U0'}],
        [{'role_id': 'R0', 'table_id': 'U0'}, {'role_id': 'R1', 'table_id': 'U0'}],
        [{'role_id': 'R0', 'table_id': 'U999'}, {'role_id': 'R1', 'table_id': 'U0'}]):
        provider.output = json.dumps({'matches': invalid})
        with pytest.raises(ValueError):
            tools._suggest_sheet_roles(cfg, uploaded, match, provider)
