"""Key-set rules and template preservation for integration and splitting."""
import io
import json
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from backend.utils.data_scope import select_scope_keys, translate_data_scope, validate_data_scope
from backend.api.tools import _replace_expr_tokens, _required_cols_by_role, _resolved_scheme_config
from backend.tests.test_table_split_api import api  # shared isolated SQLite/API fixture


def tables():
    return {name: {'df': pd.DataFrame(rows, columns=['工号', '状态'])} for name, rows in {
        'A.xlsx': [('1', '在职'), ('9', '在职')],
        'B.xlsx': [('1', '在职'), ('2', '在职'), ('3', '离职'), ('2', '在职'), ('', '在职')],
        'C.xlsx': [('2', '在职'), ('4', '在职')],
    }.items()}


@pytest.mark.parametrize('config,expected', [
    ({'mode': 'source', 'source_file': 'B.xlsx'}, ['1', '2', '3']),
    ({'mode': 'source', 'source_file': 'B.xlsx', 'filter_expression': '[状态]="在职"'}, ['1', '2']),
    ({'mode': 'expression', 'expression': 'UNION("B.xlsx","C.xlsx")'}, ['1', '2', '3', '4']),
    ({'mode': 'expression', 'expression': 'INTERSECT("B.xlsx","C.xlsx")'}, ['2']),
    ({'mode': 'expression', 'expression': 'EXCEPT("B.xlsx","C.xlsx")'}, ['1', '3']),
    ({'mode': 'expression', 'expression': 'EXCEPT(UNION("B.xlsx","C.xlsx"),"A.xlsx")'}, ['2', '3', '4']),
    ({'mode': 'expression', 'expression': 'FILTER("B.xlsx", "[状态]=\'在职\'")'}, ['1', '2']),
    ({'mode': 'expression', 'expression': 'EXCEPT("B.xlsx","B.xlsx")'}, []),
])
def test_scope_set_rules(config, expected):
    parsed = tables()
    rows = select_scope_keys(config, parsed, {name: '工号' for name in parsed})
    assert [row['normalized'] for row in rows] == expected


def test_old_schemes_have_no_restriction():
    assert select_scope_keys(None, tables(), {}) is None
    assert select_scope_keys({'mode': 'all'}, tables(), {}) is None


@pytest.mark.parametrize('config', [
    {'mode': 'source', 'source_file': '不存在.xlsx'},
    {'mode': 'source', 'source_file': 'B.xlsx', 'filter_expression': '[不存在]="在职"'},
    {'mode': 'expression', 'expression': '__import__("os").system("whoami")'},
    {'mode': 'expression', 'expression': 'FILTER("B.xlsx", "状态")'},
    {'mode': 'expression', 'expression': 'EXCEPT("B.xlsx")'},
])
def test_invalid_rules_fail_explicitly(config):
    with pytest.raises(ValueError):
        select_scope_keys(config, tables(), {name: '工号' for name in tables()})


def test_saved_scope_renames_files_and_filter_columns_without_changing_literals():
    config = {'mode': 'expression', 'expression': 'EXCEPT(FILTER("B.xlsx", "[状态]=\'状态\'"),"C.xlsx")'}
    translated = translate_data_scope(config, {'B.xlsx': '新工资.xlsx', 'C.xlsx': '新清单.xlsx'},
        {'B.xlsx': {'状态': '员工状态'}}, _replace_expr_tokens)
    assert '新工资.xlsx' in translated['expression'] and '新清单.xlsx' in translated['expression']
    # Single-quoted text in the condition must remain literal as well.
    assert "[员工状态]='状态'" in translated['expression']


def book(keys, source=False):
    wb = Workbook()
    ws = wb.active
    ws.title = '明细'
    ws.append(['工号', '工作地', '金额', '加倍', '状态'])
    for row, key in enumerate(keys, 2):
        ws.append([key, '上海' if int(key) % 2 else '北京', int(key) * 10,
                   f'=C{row}*2', '在职' if int(key) % 2 else '离职'])
        ws.cell(row, 3).number_format = '#,##0.00'
    if not source:
        ws.append(['合计', None, f'=SUM(C2:C{len(keys)+1})', f'=SUM(D2:D{len(keys)+1})'])
        ws.append(['备注：请保留'])
        wb.create_sheet('其他页')['A1'] = '保留内容'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def uploaded(client, tool):
    data = client.post(f'/api/tools/{tool}/analyze', files=[
        ('files', ('A.xlsx', book([str(i) for i in range(1, 20)] + ['99']), 'application/octet-stream')),
        ('files', ('B.xlsx', book([str(i) for i in range(1, 31)], True), 'application/octet-stream')),
        ('files', ('C.xlsx', book([str(i) for i in range(25, 34)] + ['100'], True), 'application/octet-stream')),
    ])
    assert data.status_code == 200, data.text
    return {'session_id': data.json()['session_id'], 'main_file': 'A.xlsx',
        'key_map': {'A.xlsx': '工号', 'B.xlsx': '工号', 'C.xlsx': '工号'},
        'overwrite_pairs': [{'a_col': col, 'source_file': 'B.xlsx', 'source_col': col}
                            for col in ('工作地', '金额', '状态')],
        'union_source_keys': True,
        'data_scope': {'mode': 'source', 'source_file': 'B.xlsx'}}


@pytest.mark.parametrize('tool', ['integrate', 'table-split'])
def test_b_controls_output_a_supplies_template_and_schemes_preserve_scope(api, tool):
    client, state, _ = api
    for action in ('', '.create', '.apply', '.edit', '.delete'):
        state.user.role.permissions['tools.data_integrate' + action] = True
    request = uploaded(client, tool)
    if tool == 'table-split':
        request['split_config'] = {'mode': 'value', 'column': '工作地', 'data_row_start': 2, 'data_row_end': 21}
    saved = client.post(f'/api/tools/{tool}/scheme/save', json={**request, 'name': '以B为准'})
    assert saved.status_code == 200, saved.text
    schemes = client.get(f'/api/tools/{tool}/schemes').json()
    assert schemes[0]['config']['data_scope'] == request['data_scope']
    exported = client.get(f'/api/tools/{tool}/scheme/{saved.json()["id"]}/export')
    assert exported.status_code == 200
    assert exported.json()['config']['data_scope'] == request['data_scope']
    imported = client.post(f'/api/tools/{tool}/scheme/import', files={
        'file': ('scheme.json', exported.content, 'application/json')}, data={'name': '导入以B为准'})
    assert imported.status_code == 200, imported.text
    response = client.post(f'/api/tools/{tool}/execute', json=request)
    assert response.status_code == 200, response.text
    outputs = [response.content]
    if tool == 'table-split':
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            outputs = [archive.read(name) for name in archive.namelist() if name.endswith('.xlsx')]
        assert len(outputs) == 2
    seen = []
    for output in outputs:
        wb = load_workbook(io.BytesIO(output))
        ws = wb['明细']
        cached = load_workbook(io.BytesIO(output), data_only=True)['明细']
        assert wb['其他页']['A1'].value == '保留内容'
        footer = next(row for row in ws.iter_rows() if row[0].value == '合计')
        assert ws.cell(footer[0].row + 1, 1).value == '备注：请保留'
        assert footer[2].value.startswith('=SUM(') and '#REF!' not in footer[2].value
        expected_total = sum(ws.cell(row, 3).value for row in range(2, footer[0].row))
        assert cached.cell(footer[0].row, 3).value == expected_total
        assert cached.cell(footer[0].row, 4).value == expected_total * 2
        for row in ws.iter_rows(min_row=2, max_row=footer[0].row - 1):
            key = str(row[0].value)
            seen.append(key)
            assert row[2].value == int(key) * 10
            assert row[2].number_format == '#,##0.00'
            assert row[3].value.startswith('=C')
    assert len(seen) == 30 and set(seen) == {str(i) for i in range(1, 31)}


def test_no_result_scope_returns_actionable_error(api):
    client, _, _ = api
    request = uploaded(client, 'table-split')
    request.update(split_config={'mode': 'value', 'column': '工作地'},
        data_scope={'mode': 'expression', 'expression': 'EXCEPT("B.xlsx","B.xlsx")'})
    result = client.post('/api/tools/table-split/execute', json=request)
    assert result.status_code >= 400
    assert '没有有效主键' in result.text


@pytest.mark.parametrize('tool', ['integrate', 'table-split'])
@pytest.mark.parametrize('expression,expected', [
    ('INTERSECT("B.xlsx","C.xlsx")', set(range(25, 31))),
    ('EXCEPT("B.xlsx","C.xlsx")', set(range(1, 25))),
    ('FILTER("B.xlsx", "[状态]=\'在职\'")', set(range(1, 31, 2))),
])
def test_expression_scope_controls_delivered_rows(api, tool, expression, expected):
    client, _, _ = api
    request = uploaded(client, tool)
    request['data_scope'] = {'mode': 'expression', 'expression': expression}
    if tool == 'table-split':
        # Explicit union mode still enables scope when the old split expansion
        # checkbox is off; either enabling control must work consistently.
        request['split_config'] = {'mode': 'value', 'column': '工作地', 'include_source_keys': False}
    response = client.post(f'/api/tools/{tool}/execute', json=request)
    assert response.status_code == 200, response.text
    outputs = [response.content]
    if tool == 'table-split':
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            outputs = [archive.read(name) for name in archive.namelist() if name.endswith('.xlsx')]
    keys = []
    for output in outputs:
        ws = load_workbook(io.BytesIO(output))['明细']
        footer_row = next(row[0].row for row in ws.iter_rows() if row[0].value == '合计')
        assert '#REF!' not in ws.cell(footer_row, 3).value
        assert ws.cell(footer_row + 1, 1).value == '备注：请保留'
        keys.extend(int(ws.cell(row, 1).value) for row in range(2, footer_row))
    assert len(keys) == len(expected) and set(keys) == expected


def test_scope_filter_columns_are_required_and_translated_when_reusing_scheme():
    config = {'main_fp': 'a', 'roles': [{'file': 'A.xlsx', 'fp': 'a'}, {'file': 'B.xlsx', 'fp': 'b'}],
        'cols_by_fp': {'a': ['工号'], 'b': ['工号', '状态']},
        'key_map_by_role': {'0': '工号', '1': '工号'},
        'data_scope': {'mode': 'source', 'source_file': 'B.xlsx', 'filter_expression': '[状态]="在职"'}}
    assert '状态' in _required_cols_by_role(config)[1]
    resolved = _resolved_scheme_config(config, {'roles': config['roles'],
        'role_files': ['新模板.xlsx', '新花名册.xlsx'],
        'column_maps_by_role': {0: {'工号': '员工编号'}, 1: {'工号': '编号', '状态': '人员状态'}}})
    assert resolved['data_scope']['source_file'] == '新花名册.xlsx'
    assert resolved['data_scope']['filter_expression'] == '[人员状态]="在职"'


def test_invalid_scope_is_rejected_at_save_and_execute(api):
    client, _, _ = api
    request = uploaded(client, 'table-split')
    request['split_config'] = {'mode': 'value', 'column': '工作地'}
    request['data_scope'] = {'mode': 'source', 'source_file': '不存在.xlsx'}
    for endpoint, payload in [('execute', request), ('scheme/save', {**request, 'name': '非法范围'})]:
        response = client.post(f'/api/tools/table-split/{endpoint}', json=payload)
        assert response.status_code == 400
        assert '不存在的表' in response.text
