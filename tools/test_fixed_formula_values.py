"""Fixed formulas must work without source rows, including saved schemes."""
import io
import zipfile

import pytest
from openpyxl import load_workbook

from backend.utils.integrate_engine import eval_source_expr, eval_source_expr_cross, resolve_overwrites
from backend.api.tools import _required_cols_by_role
from backend.tests.test_table_split_api import api, analyze, payload


@pytest.mark.parametrize('expression,expected', [
    ('"aaaa"', 'aaaa'), ('="aaaa"', 'aaaa'), ('"0012"', '0012'),
    ('"姓名"', '姓名'), ('"说""你好"""', '说"你好"'),
    ('"A"&"B"', 'AB'), ('42', 42), ('=ROUND(1.235,2)', 1.24),
    ('IF(1=1,"上海","其他")', '上海'), ('""', ''),
])
def test_constant_formulas_do_not_need_source_rows(expression, expected):
    sources = {'B.xlsx': {'cols': ['姓名'], 'rows': {}}}
    assert eval_source_expr_cross(expression, 'B.xlsx', sources, 'missing') == expected
    assert eval_source_expr(expression, [], ['姓名']) == expected
    assert eval_source_expr(expression, [{'姓名': '张三'}], ['姓名']) == expected


@pytest.mark.parametrize('expression', ['未知列', '姓名&未知列', '__import__("os")', '姓名&"x"'])
def test_constant_fix_keeps_invalid_or_fully_missing_references_empty(expression):
    assert eval_source_expr_cross(expression, 'B.xlsx', {'B.xlsx': {'cols': ['姓名'], 'rows': {}}}, 'missing') is None


def test_constant_strings_are_not_required_source_columns_in_saved_schemes():
    cfg = {'main_fp': 'a', 'roles': [{'fp': 'a', 'file': 'A.xlsx'}, {'fp': 'b', 'file': 'B.xlsx'}],
        'cols_by_fp': {'a': ['工号', '备注'], 'b': ['工号', 'aaaa']},
        'key_map_by_role': {'0': '工号', '1': '工号'},
        'overwrite_pairs': [{'a_col': '备注', 'source_role': 1, 'source_expr': '"aaaa"'}]}
    assert _required_cols_by_role(cfg)[1] == ['工号']
    cfg['overwrite_pairs'][0]['source_expr'] = '"固定内容"'
    assert _required_cols_by_role(cfg)[1] == ['工号']
    cfg['required_cols_by_role'] = {'0': ['工号', '备注'], '1': ['工号', '"aaaa"']}
    assert _required_cols_by_role(cfg)[1] == ['工号']


def test_fixed_value_overwrites_even_when_default_source_has_no_matching_key():
    assert resolve_overwrites('missing', [{'a_col': '备注', 'source_file': 'B.xlsx', 'source_expr': '"aaaa"'}],
        {'B.xlsx': {'cols': ['姓名'], 'rows': {}}}) == {'备注': 'aaaa'}


@pytest.mark.parametrize('tool', ['integrate', 'table-split'])
def test_constant_value_is_written_and_scheme_can_be_reapplied(api, tool):
    client, state, _ = api
    for action in ('', '.create', '.apply', '.edit', '.delete'):
        state.user.role.permissions['tools.data_integrate' + action] = True
    session = analyze(client)
    if tool == 'integrate':
        # Integration and splitting use different session roots.
        from backend.tests.test_table_split_api import workbook_bytes
        response = client.post('/api/tools/integrate/analyze', files=[
            ('files', ('template.xlsx', workbook_bytes(), 'application/octet-stream')),
            ('files', ('source.xlsx', workbook_bytes(True), 'application/octet-stream'))])
        assert response.status_code == 200, response.text
        session = response.json()
    request = payload(session)
    request['overwrite_pairs'] = [{'a_col': '工作地', 'source_file': 'source.xlsx', 'source_expr': '"aaaa"'}]
    if tool == 'integrate':
        request.pop('split_config')
    saved = client.post(f'/api/tools/{tool}/scheme/save', json={**request, 'name': '固定文本方案'})
    assert saved.status_code == 200, saved.text
    validation = client.post(f'/api/tools/{tool}/scheme/apply-validate', json={
        'session_id': session['session_id'], 'scheme_id': saved.json()['id']})
    assert validation.status_code == 200 and validation.json()['ok'], validation.text
    response = client.post(f'/api/tools/{tool}/execute', json=request)
    assert response.status_code == 200, response.text
    outputs = [response.content]
    if tool == 'table-split':
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            outputs = [archive.read(name) for name in archive.namelist() if name.endswith('.xlsx')]
    count = 0
    for content in outputs:
        ws = load_workbook(io.BytesIO(content), data_only=True).active
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value or '').isdigit():
                assert row[1].value == 'aaaa'
                count += 1
    assert count == 6
