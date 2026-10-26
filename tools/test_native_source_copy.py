"""Integration tests of the actual native source-copy helper across workbooks.

Requires the application's pythonnet/.NET runtime and bundled Aspose assembly.
Run separately from the dependency-light test suite.
"""
import copy
import importlib.util
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904


@unittest.skipUnless(importlib.util.find_spec('pythonnet'), 'requires application native Excel runtime')
class NativeSourceCopyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import aspose_init
        if not aspose_init.init_aspose():
            raise RuntimeError('native Excel runtime could not initialize')

    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.output = self.root / 'result.xlsx'

    def assert_sheet_equal(self, original, copied):
        self.assertEqual((original.max_row, original.max_column), (copied.max_row, copied.max_column))
        original_values = {c.coordinate: c.value for c in original._cells.values() if c.value is not None}
        copied_values = {c.coordinate: c.value for c in copied._cells.values() if c.value is not None}
        self.assertEqual(original_values, copied_values)
        self.assertEqual(str(original.merged_cells), str(copied.merged_cells))
        for cell in original._cells.values():
            result = copied[cell.coordinate]
            self.assertEqual(cell.number_format, result.number_format, cell.coordinate)
            self.assertEqual(cell.data_type, result.data_type, cell.coordinate)
            self.assertEqual(copy.copy(cell.font), copy.copy(result.font), cell.coordinate)
            self.assertEqual(cell.fill.patternType, result.fill.patternType, cell.coordinate)
            if cell.fill.patternType == 'solid':
                # Excel ignores alpha for cell fills; native serialization makes it opaque.
                self.assertEqual(cell.fill.fgColor.rgb[-6:], result.fill.fgColor.rgb[-6:], cell.coordinate)
        for key, dim in original.row_dimensions.items():
            self.assertEqual(dim.height or original.sheet_format.defaultRowHeight,
                             copied.row_dimensions[key].height or copied.sheet_format.defaultRowHeight, key)
            self.assertEqual(dim.hidden, copied.row_dimensions[key].hidden, key)
        for key, dim in original.column_dimensions.items():
            self.assertAlmostEqual(dim.width, copied.column_dimensions[key].width, places=2)
            self.assertEqual(dim.hidden, copied.column_dimensions[key].hidden, key)

    def test_multiple_files_keep_distinct_shared_strings_styles_and_coordinates(self):
        from backend.utils.template_pipeline import copy_original_source_sheets
        from backend.utils.output_postprocess import finalize_output_workbook
        plan = []
        target = Workbook()
        target.active.title = '目标'
        target.active['A1'] = '目标已有内容'
        for index in range(4):
            book = Workbook()
            source = book.active
            source.title = f'明细{index}'
            # Each file has different strings/styles at the same local indices.
            source.merge_cells('A1:H1')
            source['A1'] = f'独立文件{index}的标题'
            source['C7'], source['F7'] = f'编号{index}', f'字段{index}'
            source['C8'] = f'{index:06d}'
            source['F8'] = datetime(2026, 10, 23) if index == 0 else index * 100.125
            source['F8'].number_format = 'yyyy/mm/dd' if index == 0 else '0.' + '0' * index
            source['F8'].fill = PatternFill('solid', fgColor=f'{index + 1:02X}3355')
            source['C8'].font = Font(name='Arial', family=2, bold=index % 2 == 0, size=10 + index)
            source['Z50'] = f'数据区域以外的尾注{index}'
            source['Y49'].number_format = f'0.{"0" * (index + 1)}'
            source.row_dimensions[8].height = 25 + index
            source.row_dimensions[12].hidden = True
            source.column_dimensions['E'].width = 20 + index
            source.column_dimensions['D'].hidden = True
            path = self.root / f'原始{index}.xlsx'
            book.save(path)
            book.close()
            name = '源_' + source.title
            target.create_sheet(name)
            plan.append({'file_path': str(path), 'sheet_name': source.title,
                         'destination': name, 'placeholder': True})
        target.save(self.output)
        target.close()
        before = {p['file_path']: Path(p['file_path']).read_bytes() for p in plan}
        copy_original_source_sheets(self.output, plan)
        # Both copies and the final calculation/save must preserve sources.
        for finalized in (False, True):
            if finalized:
                finalize_output_workbook(self.output, script_code='TEMPLATE_PIPELINE_VERSION = 2')
            with self.subTest(finalized=finalized):
                result = load_workbook(self.output)
                try:
                    self.assertEqual(result['目标']['A1'].value, '目标已有内容')
                    for spec in plan:
                        original = load_workbook(spec['file_path'])
                        try:
                            self.assert_sheet_equal(original[spec['sheet_name']], result[spec['destination']])
                        finally:
                            original.close()
                finally:
                    result.close()
        self.assertTrue(all(Path(path).read_bytes() == data for path, data in before.items()))

    def test_native_helper_preserves_dependencies_charts_and_different_date_systems(self):
        from backend.utils.template_pipeline import copy_original_source_sheets
        from backend.utils.output_postprocess import finalize_output_workbook
        source_path = self.root / '原始.xlsx'
        book = Workbook()
        book.epoch = CALENDAR_MAC_1904
        source = book.active
        source.title = '原始明细'
        source['C7'], source['F7'] = '编号', '金额'
        source['C8'], source['F8'] = '001', 20.125
        source['F8'].number_format = '0.0000'
        source['E8'] = datetime(2026, 10, 23)
        support = book.create_sheet('辅助')
        support.sheet_state = 'hidden'
        support['A1'], support['A2'] = 2, 4
        source['G8'] = '=辅助!A1+1'
        chart = BarChart()
        chart.add_data(Reference(support, min_col=1, min_row=1, max_row=2))
        source.add_chart(chart, 'A13')
        book.save(source_path)
        book.close()
        target = Workbook()
        target.active.title = '目标'
        target.create_sheet('源_明细')
        target.active['A1'] = "=INDEX('源_明细'!F:F,MATCH(\"001\",'源_明细'!C:C,0))"
        target.active['B1'] = "='源_明细'!G8"
        target.save(self.output)
        target.close()
        copied = copy_original_source_sheets(self.output, [
            {'file_path': str(source_path), 'sheet_name': '原始明细',
             'destination': '源_明细', 'placeholder': True}])
        self.assertEqual(len(copied), 2)
        finalize_output_workbook(self.output, script_code='TEMPLATE_PIPELINE_VERSION = 2')
        result = load_workbook(self.output)
        values = load_workbook(self.output, data_only=True)
        try:
            self.assertEqual(values['目标']['A1'].value, 20.125)
            self.assertEqual(values['目标']['B1'].value, 3)
            self.assertEqual(result['源_明细']['E8'].value, datetime(2026, 10, 23))
            self.assertEqual(len(result['源_明细']._charts), 1)
            dependency = next(name for name in copied if name != '源_明细')
            self.assertEqual(result[dependency].sheet_state, 'hidden')
            self.assertIn(dependency, result['源_明细']['G8'].value)
            self.assertIn(dependency, result['源_明细']._charts[0].ser[0].val.numRef.f)
            self.assertEqual(len(result._external_links), 0)
        finally:
            result.close()
            values.close()


if __name__ == '__main__':
    unittest.main()
