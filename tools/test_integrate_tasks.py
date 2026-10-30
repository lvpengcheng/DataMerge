"""Real workbook/API coverage for independent targets and saved role translation."""
import copy
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from backend.api import tools
pytest_plugins = ('backend.tests.test_table_split_api',)


def workbook(source=False):
    book = Workbook()
    book.remove(book.active)
    for name, amount, extra in [('工资', 100, '002'), ('奖金', 700, '003')]:
        sheet = book.create_sheet(name)
        sheet.append(['工号', '金额', '加倍'])
        sheet.append(['001', amount if source else 0, '=B2*2'])
        sheet['B2'].fill = PatternFill('solid', fgColor='FFFF00')
        if source:
            sheet.append([extra, amount * 2, '=B3*2'])
    if not source:
        book.create_sheet('汇总')['A1'] = '=工资!B2+奖金!B2'
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture
def session(api, monkeypatch):
    client, state, _ = api
    for suffix in ('', '.create', '.edit', '.apply', '.delete'):
        state.user.role.permissions['tools.data_integrate'+suffix] = True
    async def parse(path, **kwargs):
        return tools._parse_file_all_sheets_impl(path, **kwargs)
    monkeypatch.setattr(tools, '_parse_file_all_sheets_fresh', parse)
    response = client.post('/api/tools/integrate/analyze', data={'parse_all_sheets':'true'}, files=[
        ('files', ('模板.xlsx', workbook(), 'application/octet-stream')),
        ('files', ('来源.xlsx', workbook(True), 'application/octet-stream'))])
    assert response.status_code == 200, response.text
    analyzed = response.json()
    selected = [f for f in analyzed['files'] if f['sheet'] != '汇总']
    response = client.post('/api/tools/integrate/select-sheets', json={
        'session_id':analyzed['session_id'], 'selected_tables':[f['name'] for f in selected]})
    assert response.status_code == 200, response.text
    tasks = []
    for sheet, multiplier in [('工资', 2), ('奖金', 3)]:
        main = next(f['name'] for f in selected if f['physical_file'] == '模板.xlsx' and f['sheet'] == sheet)
        source = next(f['name'] for f in selected if f['physical_file'] == '来源.xlsx' and f['sheet'] == sheet)
        tasks.append({'main_file':main, 'active_files':[main,source], 'key_map':{main:'工号',source:'工号'},
            'overwrite_pairs':[{'a_col':'金额','source_file':source,'source_expr':f'金额*{multiplier}'}],
            'compare_pairs':[], 'output_mode':1, 'union_source_keys':True,
            'output_file':'结果.xlsx', 'output_sheet':sheet+'明细'})
    request = {'session_id':analyzed['session_id'], **tasks[0], 'output_tasks':tasks}
    return client, state, request, selected


def test_multiple_targets_preserve_formulas_styles_and_source_isolation(session):
    client, _, request, _ = session
    response = client.post('/api/tools/integrate/execute', json=request)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), data_only=True)
    assert {row[0]:row[1] for row in list(book['工资明细'].values)[1:]} == {'001':200,'002':400}
    assert {row[0]:row[1] for row in list(book['奖金明细'].values)[1:]} == {'001':2100,'003':4200}
    assert {book['工资明细'].cell(i,1).value for i in (2,3)} == {'001','002'}
    assert {book['奖金明细'].cell(i,1).value for i in (2,3)} == {'001','003'}
    assert book['汇总']['A1'].value == 2300
    formulas = load_workbook(io.BytesIO(response.content))
    assert formulas['工资明细']['C2'].value == '=B2*2'
    assert '工资明细' in formulas['汇总']['A1'].value
    assert formulas['工资明细']['B2'].fill.fgColor.rgb.endswith('FFFF00')


def test_multiple_workbooks_and_heartbeat_download(session):
    client, _, request, _ = session
    request['output_tasks'][1]['output_file'] = '奖金.xlsx'
    response = client.post('/api/tools/integrate/execute?stream=true', json=request)
    assert response.status_code == 200, response.text
    download = client.get(response.json()['download_url'])
    assert download.status_code == 200
    with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
        assert set(archive.namelist()) == {'结果.xlsx', '奖金.xlsx'}
        book = load_workbook(io.BytesIO(archive.read('奖金.xlsx')),data_only=True)
        assert {r[0]:r[1] for r in list(book['奖金明细'].values)[1:]}['001'] == 2100


def test_scheme_save_export_import_and_apply_translate_all_target_roles(session):
    client, _, request, selected = session
    saved = client.post('/api/tools/integrate/scheme/save', json={**request,'name':'多目标'})
    assert saved.status_code == 200, saved.text
    sid = saved.json()['id']
    exported = client.get(f'/api/tools/integrate/scheme/{sid}/export')
    assert len(exported.json()['config']['output_tasks']) == 2
    imported = client.post('/api/tools/integrate/scheme/import',data={'name':'导入多目标'},
        files={'file':('方案.json',exported.content,'application/json')})
    assert imported.status_code == 200, imported.text
    roles = exported.json()['config']['roles']
    automatic = client.post('/api/tools/integrate/analyze', data={'scheme_id':sid}, files=[
        ('files', ('模板.xlsx',workbook(),'application/octet-stream')),
        ('files', ('来源.xlsx',workbook(True),'application/octet-stream'))])
    assert automatic.status_code == 200, automatic.text
    assert automatic.json()['parse_all_sheets'] is True
    assert automatic.json()['sheet_selection_pending'] is False
    assert automatic.json()['matched_schemes']
    validation = client.post('/api/tools/integrate/scheme/apply-validate',json={
        'session_id':request['session_id'],'scheme_id':sid,'role_files':[r['file'] for r in roles]})
    assert validation.status_code == 200 and validation.json()['ok'], validation.text
    resolved = validation.json()['resolved_scheme']['config']['output_tasks']
    assert [t['main_file'] for t in resolved] == [t['main_file'] for t in request['output_tasks']]
    assert all(set(t['key_map']) == set(t['active_files']) for t in resolved)
    # Across-month renaming must translate each task's target, keys and formula prefixes.
    cfg = exported.json()['config']
    uploaded = [{**f,'name':'new_'+f['name']} for f in selected]
    match = tools._match_scheme_config(cfg, uploaded, ['new_'+r['file'] for r in roles])
    assert match
    result = tools._resolved_scheme_config(cfg,match)['output_tasks']
    assert all(t['main_file'].startswith('new_') for t in result)
    assert all(p['source_file'].startswith('new_') for t in result for p in t['overwrite_pairs'])


@pytest.mark.parametrize('change', ['duplicate','path','excluded','missing-column','empty','kept-sheet','same-template-sheet'])
def test_invalid_targets_are_rejected_before_writing(session,change):
    client, _, original, _ = session
    request = copy.deepcopy(original)
    tasks = request['output_tasks']
    if change == 'duplicate': tasks[1]['output_sheet'] = tasks[0]['output_sheet']
    if change == 'path': tasks[1]['output_file'] = '../unsafe.xlsx'
    if change == 'excluded': tasks[1]['overwrite_pairs'][0]['source_expr'] = tasks[0]['active_files'][1]+'.金额'
    if change == 'missing-column': tasks[1]['key_map'][tasks[1]['main_file']] = '不存在'
    if change == 'empty': tasks[1]['overwrite_pairs'] = []
    if change == 'kept-sheet': tasks[1]['output_sheet'] = '汇总'
    if change == 'same-template-sheet': tasks[1]['main_file'] = tasks[0]['main_file']
    response = client.post('/api/tools/integrate/execute',json=request)
    assert response.status_code in (400,500), response.text
    assert 'download_url' not in response.text
