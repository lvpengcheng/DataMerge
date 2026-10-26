"""Formula scripts must retain original source coordinates and use native copying."""
import ast
import importlib
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from backend.utils.script_entry import invoke_script_main
from backend.utils.template_pipeline import (
    source_origin, prepare_formula_source_sheets, uses_original_template_sources,
    is_source_pipeline_v2, sync_source_mode,
)
from backend.ai_engine.excel_formula_builder import ExcelFormulaBuilder
from backend.ai_engine.prompt_generator import PromptGenerator


class FormulaSourceCopyTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.inputs, self.outputs = self.root / 'in', self.root / 'out'
        self.inputs.mkdir()
        self.outputs.mkdir()
        self.path = self.inputs / '考勤.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '原始明细'
        ws.merge_cells('A1:H1')
        ws['A1'] = '原始标题'
        ws['C7'], ws['F7'] = '工号', '天数'
        ws['C8'], ws['F8'] = '001', 20
        ws['C9'], ws['F9'] = '002', 22
        ws['Z50'] = '解析区域以外也必须保留'
        wb.save(self.path)
        wb.close()
        self.region = SimpleNamespace(head_data={'工号': 'C', '天数': 'F'},
            head_row_start=7, head_row_end=7, data_row_start=8, data_row_end=9,
            column_schemas={}, data=[{'C': '001', 'F': 20}, {'C': '002', 'F': 22}])
        self.entry = {'df': pd.DataFrame({'工号': ['001', '002'], '天数': [20, 22]}),
            'columns': ['工号', '天数'], 'source_file': self.path.name,
            'source_sheet': '原始明细', 'source_origins': [source_origin(self.path, '原始明细', [self.region])]}

    def generator(self, mode='copy'):
        provider = SimpleNamespace(BaseAIProvider=object, AIProviderFactory=object)
        with patch.dict(sys.modules, {'backend.ai_engine.ai_provider': provider}):
            cls = importlib.import_module('backend.ai_engine.formula_code_generator').FormulaCodeGenerator
        generator = object.__new__(cls)
        generator._source_sheet_mode = mode
        generator._source_pipeline_v2 = True
        return generator

    def build(self, mode='copy', fill=None):
        return self.generator(mode)._build_complete_code(fill or '''
def fill_result_sheets(wb, source_sheets, salary_year, salary_month, monthly_hours):
    source = next(iter(source_sheets.values()))
    ws = wb.create_sheet('结果')
    ws['A1'], ws['B1'] = '工号', '天数'
    for i, identifier in enumerate(source['df']['工号'], 2):
        ws.cell(i, 1, identifier)
        key_col = source['column_letters']['工号']
        value_col = source['column_letters']['天数']
        start, end = source['data_start_row'], source['data_end_row']
        sheet = source['ws'].title
        ws.cell(i, 2, f"=INDEX('{sheet}'!{value_col}{start}:{value_col}{end},MATCH(A{i},'{sheet}'!{key_col}{start}:{key_col}{end},0))")
    ws['B2'].number_format = '0.0000'
''')

    def environment(self, code):
        env = {'__name__': 'test_formula_source', 'input_folder': str(self.inputs),
               'output_folder': str(self.outputs), '_pre_loaded_source_data': {'原始明细': self.entry}}
        with patch.dict(sys.modules, {'excel_parser': SimpleNamespace(IntelligentExcelParser=object)}):
            exec(code, env)
        return env

    def test_generated_script_uses_runtime_coordinates_and_never_rewrites_sources(self):
        code = self.build()
        self.assertTrue(is_source_pipeline_v2(code))
        self.assertTrue(uses_original_template_sources(code))
        original = self.path.read_bytes()
        env = self.environment(code)
        # Simulate an uploaded source whose coordinates differ from training.
        layout = self.entry['source_origins'][0]['regions'][0]
        layout['data_start_row'], layout['data_end_row'] = 13, 14
        layout['columns'][0]['letter'], layout['columns'][1]['letter'] = 'E', 'H'
        native_copy = env['copy_original_source_sheets'] = Mock()
        env['write_source_sheets'] = Mock(side_effect=AssertionError('must not rewrite DataFrames'))
        invoke_script_main(env['main'], env)
        output = self.outputs / '薪资汇总表.xlsx'
        wb = load_workbook(output)
        self.addCleanup(wb.close)
        self.assertEqual(wb['结果']['B2'].value,
            "=INDEX('原始明细'!H13:H14,MATCH(A2,'原始明细'!E13:E14,0))")
        self.assertEqual(wb['结果']['B2'].number_format, '0.0000')
        self.assertEqual(wb['原始明细'].max_row, 1)
        self.assertEqual(len(wb['原始明细']._cells), 0)  # only a reserved sheet, no row writing
        native_copy.assert_called_once()
        plan = native_copy.call_args.args[1]
        self.assertEqual(plan[0]['file_path'], str(self.path.resolve()))
        self.assertEqual(plan[0]['sheet_name'], '原始明细')
        self.assertTrue(plan[0]['placeholder'])
        self.assertEqual(self.path.read_bytes(), original)

    def test_rebuild_remains_available_only_when_explicit(self):
        code = self.build('rebuild', '''
def fill_result_sheets(wb, source_sheets, salary_year, salary_month, monthly_hours):
    wb.create_sheet('结果')['A1'] = '重建'
''')
        self.assertFalse(uses_original_template_sources(code))
        env = self.environment(code)
        native_copy = env['copy_original_source_sheets'] = Mock()
        invoke_script_main(env['main'], env)
        wb = load_workbook(self.outputs / '薪资汇总表.xlsx')
        self.addCleanup(wb.close)
        self.assertEqual(wb['原始明细']['A1'].value, '工号')
        self.assertEqual(wb['原始明细']['A2'].value, '001')
        native_copy.assert_not_called()

    def test_copy_rejects_silent_cleaning(self):
        env = self.environment(self.build())
        env['clean_source_data'] = lambda source_data: source_data
        with self.assertRaisesRegex(ValueError, 'rebuild'):
            invoke_script_main(env['main'], env)

    def test_aliases_and_reserved_names_do_not_duplicate_or_overwrite_sheets(self):
        wb = Workbook()
        wb.active.title = '原始明细'
        sources, plan = prepare_formula_source_sheets(wb,
            {'原始明细': self.entry, '别名': self.entry}, self.inputs)
        self.assertEqual(len(plan), 1)
        self.assertIs(sources['原始明细'], sources['别名'])
        self.assertNotEqual(sources['原始明细']['ws'].title, '原始明细')
        self.assertEqual(sources['别名']['column_letters'], {'工号': 'C', '天数': 'F'})
        wb.close()

    def test_prompt_uses_physical_columns_for_lookup_index(self):
        builder = ExcelFormulaBuilder()
        builder.source_sheets = {'原始明细': self.entry}
        text = builder.get_source_structure_for_prompt(original_coordinates=True)
        self.assertIn('C列: 工号', text)
        self.assertIn('F列: 天数', text)
        self.assertIn('数据起始行: 8', text)
        self.assertIn('列号=6-3+1=4', text)
        legacy = builder.get_source_structure_for_prompt(original_coordinates=False)
        self.assertIn('B列: 天数', legacy)

    def test_fallback_loader_captures_original_provenance(self):
        env = self.environment(self.build())
        env.pop('_pre_loaded_source_data')
        parser = Mock()
        parser.parse_excel_file.return_value = [SimpleNamespace(sheet_name='原始明细', regions=[self.region])]
        env['IntelligentExcelParser'] = Mock(return_value=parser)
        loaded = env['load_source_data'](str(self.inputs), {})
        origin = loaded['原始明细']['source_origins'][0]
        self.assertEqual(origin['file_path'], str(self.path.resolve()))
        self.assertEqual(origin['regions'][0]['columns'][1]['letter'], 'F')
        self.assertFalse(parser.parse_excel_file.call_args.kwargs['normalize_source'])

    def test_corrections_retain_native_pipeline_and_legacy_coordinates(self):
        generator = self.generator()
        code = self.build()
        self.assertIn('column_letters', generator._restore_source_pipeline(code, ''))
        corrected = generator._build_complete_code('def fill_result_sheets(*args):\n    pass\n')
        self.assertTrue(uses_original_template_sources(corrected))
        generator._restore_source_pipeline('def main():\n    pass\n', '')
        legacy = generator._build_complete_code('def fill_result_sheets(*args):\n    pass\n')
        self.assertFalse(is_source_pipeline_v2(legacy))
        self.assertFalse(uses_original_template_sources(legacy))
        switched = sync_source_mode(code, 'SOURCE_SHEET_MODE: rebuild')
        self.assertFalse(uses_original_template_sources(switched))
        ast.parse(switched)

    def test_generation_prompts_do_not_instruct_rebuilding_in_copy_mode(self):
        generator = PromptGenerator()
        generator.source_sheet_mode = 'copy'
        structure = {'sheets': {'结果': {'headers': {'工号': 'A', '天数': 'B'}}}}
        single = generator.generate_formula_mode_prompt('真实坐标 C、F', structure, '按工号查找天数')
        multi = generator.generate_multi_step_prompts('真实坐标 C、F', structure, '按工号查找天数')['step3']
        for prompt in (single, multi):
            self.assertIn('整张原样复制', prompt)
            self.assertIn("source['column_letters']['工号']", prompt)
            self.assertNotIn('def clean_source_data', prompt)
            self.assertNotIn('清洗后的数据写入Excel', prompt)
            self.assertIn("'{sheet}'!{val_col}{start}", prompt)
            self.assertNotIn('__SOURCE_OVERVIEW__', prompt)
        generator.source_sheet_mode = 'rebuild'
        rebuilt = generator.generate_formula_mode_prompt('结构', structure, '规则')
        self.assertIn('def clean_source_data', rebuilt)

    def test_column_correction_restores_mode_from_full_script(self):
        generator = self.generator()
        generator.prompt_generator = Mock()
        generator.prompt_generator.generate_column_adjustment_prompt.return_value = 'prompt'
        generator.ai_provider = Mock()
        generator.ai_provider.generate_code.return_value = ''
        generator.training_logger = None
        result = generator.generate_column_level_correction(self.build(), {'天数': {'count': 1}}, '', '', {})
        self.assertEqual(result, (None, ''))
        self.assertEqual(generator.prompt_generator.source_sheet_mode, 'copy')
        self.assertIn('column_letters',
            generator.prompt_generator.generate_column_adjustment_prompt.call_args.kwargs['source_structure'])


if __name__ == '__main__':
    unittest.main()
