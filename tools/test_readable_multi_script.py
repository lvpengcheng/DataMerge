import unittest

from backend.utils.multi_workbook import build_script, script_targets, readable_script, script_download_response


class ReadableMultiScriptTests(unittest.TestCase):
    def setUp(self):
        self.targets = {
            '工资.xlsx': {'code': '"""说明"""\r\nimport re\r\ndef main():\r\n    return re.sub(r"\\d+", "", "123")\r\n', 'structure': {'sheets': {}}},
            '汇总.xlsx': {'code': 'def main():\n    return "汇总"', 'structure': {}},
        }

    def assert_targets_equal(self, actual):
        import ast
        self.assertEqual(set(actual), set(self.targets))
        for name, target in self.targets.items():
            self.assertEqual(ast.dump(ast.parse(actual[name]['code'])), ast.dump(ast.parse(target['code'])))
            self.assertEqual(actual[name]['structure'], target['structure'])

    def test_visible_executable_sections_compile_and_parse(self):
        import ast
        code = build_script(self.targets)
        self.assert_targets_equal(script_targets(code))
        compile(code, 'bundle.py', 'exec')
        self.assertIn('    def main():', code)
        self.assertIn('    import re', code)
        self.assertNotIn("'code':", code)
        self.assertNotIn('\\ndef main()', code)
        tree = ast.parse(code)
        metadata = next(n for n in tree.body if isinstance(n, ast.Assign) and
                        any(isinstance(t, ast.Name) and t.id == 'MULTI_WORKBOOK_SCRIPTS' for t in n.targets))
        self.assertTrue(all('entry' in t for t in ast.literal_eval(metadata.value).values()))

    def test_old_downloads_become_one_python_file_with_visible_code(self):
        old = '# 保留包装代码\nEXTRA = 42\nMULTI_WORKBOOK_SCRIPTS = ' + repr(self.targets) + '\n'
        formatted = readable_script(old)
        self.assertIn('EXTRA = 42', formatted)
        self.assert_targets_equal(script_targets(formatted))
        response = script_download_response(old, '工资.py', {'Cache-Control': 'no-store'})
        self.assertIn('text/x-python', response.media_type)
        self.assertNotIn('.zip', response.headers['content-disposition'])
        self.assertEqual(response.headers['cache-control'], 'no-store')
        self.assert_targets_equal(script_targets(response.body.decode()))
        self.assertEqual(readable_script(formatted), formatted)

    def test_uploaded_single_file_executes_edited_sections_in_separate_namespaces(self):
        import tempfile
        from pathlib import Path
        from openpyxl import load_workbook
        from unittest.mock import patch
        source = ("from pathlib import Path\nfrom openpyxl import Workbook\n"
                  "amount = VALUE\ndef helper():\n    return amount\n"
                  "def main():\n    book = Workbook()\n    book.active['A1'] = helper()\n"
                  "    book.save(Path(output_folder) / 'result.xlsx')\n")
        targets = {name: {'code': source.replace('VALUE', str(value)), 'structure': {}}
                   for name, value in [('明细.xlsx', 11), ('汇总.xlsx', 22)]}
        # This is the same compile/exec/main path as a single uploaded .py.
        code = build_script(targets).replace('amount = 11', 'amount = 33')
        with tempfile.TemporaryDirectory() as folder, patch('backend.utils.multi_workbook.resolve_child_template', return_value=None):
            env = {'output_folder': folder, '__name__': 'uploaded_script'}
            exec(compile(code, '<uploaded.py>', 'exec'), env)
            self.assertEqual(env['main']()['output_files'], ['明细.xlsx', '汇总.xlsx'])
            for name, value in [('明细.xlsx', 33), ('汇总.xlsx', 22)]:
                book = load_workbook(Path(folder) / name)
                try:
                    self.assertEqual(book.active['A1'].value, value)
                finally:
                    book.close()

    def test_future_imports_and_correction_keep_visible_sections(self):
        from backend.utils.multi_workbook import correct_workbooks
        from types import SimpleNamespace
        targets = {'一.xlsx': {'code': 'from __future__ import annotations\ndef main(): return 1\n', 'structure': {}}}
        code = build_script(targets)
        compile(code, 'future.py', 'exec')
        self.assertIn('from __future__ import annotations', script_targets(code)['一.xlsx']['code'])
        generator = SimpleNamespace(generate_correction_code=lambda **kw: kw['original_code'].replace('return 1', 'return 2'))
        corrected = correct_workbooks(generator, {'original_code': code})
        self.assertIn('return 2', script_targets(corrected)['一.xlsx']['code'])
        self.assertNotIn("'code':", corrected)

    def test_legacy_download_remains_executable_as_one_file(self):
        import tempfile
        from pathlib import Path
        from unittest.mock import patch
        code = "from pathlib import Path\ndef main():\n    Path(output_folder, 'output.xlsx').write_bytes(b'fixture')\n"
        targets = {'旧版.xlsx': {'code': code, 'structure': {}}}
        old = ('from backend.utils.multi_workbook import execute_workbooks\n'
               'MULTI_WORKBOOK_SCRIPTS = ' + repr(targets) + '\n'
               'def main():\n    return execute_workbooks(MULTI_WORKBOOK_SCRIPTS, globals())\n')
        with tempfile.TemporaryDirectory() as folder, patch('backend.utils.multi_workbook.resolve_child_template', return_value=None):
            env = {'output_folder': folder, '__name__': 'uploaded_script'}
            exec(compile(readable_script(old), 'downloaded.py', 'exec'), env)
            env['main']()
            self.assertEqual((Path(folder) / '旧版.xlsx').read_bytes(), b'fixture')

    def test_section_template_edits_are_read_from_executed_code(self):
        import tempfile
        from pathlib import Path
        from unittest.mock import patch
        code = ("from pathlib import Path\nTEMPLATE_NAME = '原模板.xlsx'\n"
                "def main():\n    Path(output_folder, 'output.xlsx').write_bytes(TEMPLATE_PATH.encode())\n")
        program = build_script({'结果.xlsx': {'code': code, 'structure': {}}})
        program = program.replace("TEMPLATE_NAME = '原模板.xlsx'", "TEMPLATE_NAME = '新模板.xlsx'")
        calls = []
        def resolve(code, tenant=None):
            calls.append(code)
            return 'resolved-template.xlsx'
        with tempfile.TemporaryDirectory() as folder, patch('backend.utils.multi_workbook.resolve_child_template', side_effect=resolve):
            env = {'output_folder': folder, '__name__': 'uploaded_script'}
            exec(compile(program, 'edited.py', 'exec'), env)
            env['main']()
            self.assertIn('新模板.xlsx', calls[0])
            self.assertEqual((Path(folder) / '结果.xlsx').read_bytes(), b'resolved-template.xlsx')

    def test_single_file_passes_the_training_sandbox(self):
        import tempfile
        from pathlib import Path
        from backend.sandbox.code_sandbox import CodeSandbox, _execute_script_in_proc
        source = ("from pathlib import Path\nfrom openpyxl import Workbook\n"
                  "def main():\n    book = Workbook()\n"
                  "    book.active['A1'] = 7\n    book.save(Path(output_folder) / 'result.xlsx')\n")
        program = build_script({name: {'code': source, 'structure': {}}
                                for name in ['明细.xlsx', '汇总.xlsx']})
        sandbox = CodeSandbox()
        self.assertEqual(sandbox._clean_code_syntax(program), program)
        self.assertTrue(sandbox._is_safe_script(program))
        with tempfile.TemporaryDirectory() as folder:
            inputs = Path(folder) / 'inputs'
            inputs.mkdir()
            result = _execute_script_in_proc(program, {'input_folder': str(inputs), 'output_folder': folder})
            self.assertTrue(result['success'], result.get('error'))
            self.assertTrue((Path(folder) / '明细.xlsx').is_file())
            self.assertTrue((Path(folder) / '汇总.xlsx').is_file())

    def test_multiline_strings_keep_their_values_when_executed(self):
        import tempfile
        from pathlib import Path
        source = ('from pathlib import Path\n'
                  'text = """first\nsecond\n  third"""\n'
                  'formatted = f"""{1 + 1}\nnext"""\n'
                  'def main():\n    Path(output_folder, "result.xlsx").write_text(text + formatted)\n')
        with tempfile.TemporaryDirectory() as folder:
            env = {'output_folder': folder, '__name__': 'uploaded_script'}
            program = build_script({'多行.xlsx': {'code': source, 'structure': {}}})
            self.assertEqual(script_targets(program)['多行.xlsx']['code'], source)
            exec(compile(program, 'strings.py', 'exec'), env)
            env['main']()
            self.assertEqual((Path(folder) / '多行.xlsx').read_text(), 'first\nsecond\n  third2\nnext')

    def test_single_download_stays_python(self):
        code = 'def main():\n    return 1\n'
        response = script_download_response(code, '工资.py')
        self.assertEqual(response.body.decode(), code)
        self.assertIn('text/x-python', response.media_type)

    def test_result_names_hide_legacy_scheme_labels_without_changing_paths(self):
        from backend.utils.result_selection import result_display_name
        self.assertEqual(result_display_name('工资 · 方案 abc012345678_202610_20261028102030_012abc34_原版.zip'),
                         '工资_202610_20261028102030_原版.zip')
        self.assertEqual(result_display_name('工资_202610_20261028102030123456_纯值.zip'),
                         '工资_202610_20261028102030123456_纯值.zip')


if __name__ == '__main__':
    unittest.main()
