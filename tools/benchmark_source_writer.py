"""对比指定 Git 基线与当前公式脚本的源表写入，验证逐格值/格式一致。"""
import argparse
import ast
import gc
import hashlib
import json
from pathlib import Path
import statistics
import subprocess
import sys
import time
import types

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import numpy as np
import openpyxl
import pandas as pd
from backend.ai_engine.formula_code_generator import FormulaCodeGenerator
from backend.utils.source_sheet_writer import dt_to_excel_serial, is_long_digit_text


def functions(code):
    namespace = {'pd': pd, 'PatternFill': openpyxl.styles.PatternFill, 'Font': openpyxl.styles.Font,
                 'dt_to_excel_serial': dt_to_excel_serial, 'is_long_digit_text': is_long_digit_text}
    for node in ast.parse(code).body:
        if isinstance(node, ast.FunctionDef):
            exec(compile(ast.Module(body=[node], type_ignores=[]), '<generated>', 'exec'), namespace)
    return namespace


def digest(ws):
    h = hashlib.sha256()
    for row in ws.iter_rows():
        for cell in row:
            h.update(repr((cell.value, cell.data_type, cell.number_format, cell.font.bold,
                           cell.fill.fgColor.rgb)).encode('utf-8'))
    return h.hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--rows', type=int, default=3000)
    parser.add_argument('--columns', type=int, default=80)
    parser.add_argument('--repeats', type=int, default=3)
    parser.add_argument('--baseline-ref', default='HEAD')
    args = parser.parse_args()
    source = subprocess.check_output(['git', 'show', f'{args.baseline_ref}:backend/ai_engine/formula_code_generator.py'],
                                     cwd=ROOT).decode('utf-8')
    module = types.ModuleType('backend.ai_engine._benchmark_baseline')
    module.__package__ = 'backend.ai_engine'
    exec(compile(source, '<baseline>', 'exec'), module.__dict__)
    fill = 'def fill_result_sheets(*args):\n    pass'
    generators = {'baseline': module.FormulaCodeGenerator(ai_provider=object()),
                  'current': FormulaCodeGenerator(ai_provider=object())}
    writers = {name: functions(gen._build_complete_code(fill))['write_source_sheets']
               for name, gen in generators.items()}
    columns = [f'金额{i}' for i in range(args.columns)]
    df = pd.DataFrame(np.arange(args.rows * args.columns).reshape(args.rows, args.columns) / 100,
                      columns=columns)
    data = {'源表': {'df': df, 'columns': columns, 'column_formats': dict.fromkeys(columns, '0.00')}}
    results, signatures = {}, {}
    for name, writer in writers.items():
        times = []
        for _ in range(args.repeats):
            gc.collect()
            wb = openpyxl.Workbook()
            start = time.perf_counter()
            sheets = writer(wb, data)
            times.append(time.perf_counter() - start)
            signatures[name] = digest(sheets['源表']['ws'])
            wb.close()
            del wb, sheets
        results[name] = {'seconds': times, 'median_seconds': statistics.median(times)}
    assert len(set(signatures.values())) == 1, '输出值/格式不一致，性能结果无效'
    print(json.dumps({'rows': args.rows, 'columns': args.columns, 'stage': 'source_sheet_write_only',
                      'baseline_ref': args.baseline_ref, 'output_identical': True,
                      'results': results,
                      'speedup': results['baseline']['median_seconds'] / results['current']['median_seconds']},
                     ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
